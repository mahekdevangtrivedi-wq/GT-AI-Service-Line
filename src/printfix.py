# -*- coding: utf-8 -*-
"""Conservative, printer-driver-friendly page setup applied to every generated workbook.
Avoids settings that some Windows drivers (incl. 'Microsoft Print to PDF') reject: missing paper size,
header/footer margins larger than page margins, unbounded print ranges, empty workbook-protection nodes."""
from openpyxl.utils import get_column_letter


def harden(wb, keep_a3=()):
    wb.security = None
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        ps = ws.page_setup
        if ws.title not in keep_a3:
            ps.paperSize = ws.PAPERSIZE_A4
        if ps.orientation is None:
            ps.orientation = "portrait"
        m = ws.page_margins
        m.left = min(m.left or 0.5, 0.5); m.right = min(m.right or 0.5, 0.5)
        m.top = max(0.55, min(m.top or 0.6, 0.75)); m.bottom = max(0.55, min(m.bottom or 0.6, 0.75))
        m.header = 0.25; m.footer = 0.25
        if not ws.print_area and ws.max_row > 1:
            ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        ws.print_options.gridLines = False
