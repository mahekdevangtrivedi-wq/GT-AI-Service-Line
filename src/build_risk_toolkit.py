# -*- coding: utf-8 -*-
"""GT AI Risk Assessment Toolkit (Excel) - blank template or worked example."""
import sys, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
sys.path.insert(0, os.path.dirname(__file__))
from risk_data import *

OUT = sys.argv[1]
EXAMPLE = len(sys.argv) > 2 and sys.argv[2] == "example"

GT = "4F2D7F"; LAV = "EDE7F6"; LAV2 = "D9CCEB"; GREY = "F2F2F2"; INPUT = "FFF9E6"
RAG = {"Low": "C6E0B4", "Medium": "FFE699", "High": "F4B183", "Critical": "E06666", "Prohibited": "000000"}
thin = Side(style="thin", color="BFBFBF")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
WT = Alignment(wrap_text=True, vertical="top")
WC = Alignment(wrap_text=True, vertical="center")
CC = Alignment(wrap_text=True, vertical="center", horizontal="center")

wb = openpyxl.Workbook()


def title(ws, text, sub=None, width_cols=8):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width_cols)
    c = ws.cell(1, 1, text)
    c.font = Font(size=16, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=GT)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width_cols)
        s = ws.cell(2, 1, sub); s.font = Font(italic=True, color=GT, size=10); s.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[2].height = 30


def hdr(c, fill=GT):
    c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = PatternFill("solid", fgColor=fill); c.alignment = CC; c.border = B


def box(c, size=9, bold=False, fill=None, align=WT):
    c.font = Font(size=size, bold=bold); c.border = B; c.alignment = align
    if fill: c.fill = PatternFill("solid", fgColor=fill)


def rag_cf(ws, rng):
    for k, v in RAG.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=v),
                                                      font=Font(bold=True, color="FFFFFF" if k in ("Critical", "Prohibited") else "000000")))


# =========================================================== Reference: Scales
sc = wb.active; sc.title = "Scales"
title(sc, "Risk Scales & Criteria", "Likelihood, impact (by dimension), control effectiveness, rating bands and tier requirements used throughout this toolkit. Aligned to ISO 31000 / ISO/IEC 23894 and the GT AI Risk Assessment Framework.", 9)
r = 4
sc.cell(r, 1, "Likelihood").font = Font(bold=True, color=GT, size=12); r += 1
for i, h in enumerate(["Score", "Level", "Description"], 1): hdr(sc.cell(r, i, h))
sc.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
r += 1
L_START = r
for s, n, d in LIKELIHOOD:
    box(sc.cell(r, 1, s), align=CC); box(sc.cell(r, 2, n), bold=True); box(sc.cell(r, 3, d))
    sc.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9); sc.row_dimensions[r].height = 28; r += 1
r += 1
sc.cell(r, 1, "Impact (rate the highest applicable dimension)").font = Font(bold=True, color=GT, size=12); r += 1
for i, h in enumerate(["Score", "Level"] + IMPACT_DIMS, 1): hdr(sc.cell(r, i, h))
r += 1
I_START = r
for s, n, ds in IMPACT:
    box(sc.cell(r, 1, s), align=CC); box(sc.cell(r, 2, n), bold=True)
    for j, d in enumerate(ds, 3): box(sc.cell(r, j, d))
    sc.row_dimensions[r].height = 40; r += 1
r += 1
sc.cell(r, 1, "Control effectiveness (residual = inherent score x factor)").font = Font(bold=True, color=GT, size=12); r += 1
for i, h in enumerate(["Rating", "Factor", "Description"], 1): hdr(sc.cell(r, i, h))
sc.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9); r += 1
CE_START = r
for n, f, d in CONTROL_EFFECTIVENESS:
    box(sc.cell(r, 1, n), bold=True); box(sc.cell(r, 2, f), align=CC); box(sc.cell(r, 3, d))
    sc.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9); sc.row_dimensions[r].height = 28; r += 1
CE_END = r - 1
r += 1
sc.cell(r, 1, "Risk rating bands (score = likelihood x impact, 1-25)").font = Font(bold=True, color=GT, size=12); r += 1
for i, h in enumerate(["Rating", "Min", "Max", "Required response"], 1): hdr(sc.cell(r, i, h))
sc.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9); r += 1
RB_START = r
for n, lo, hi, d in RATING_BANDS:
    box(sc.cell(r, 1, n), bold=True, fill=RAG[n]); box(sc.cell(r, 2, lo), align=CC); box(sc.cell(r, 3, hi), align=CC); box(sc.cell(r, 4, d))
    sc.merge_cells(start_row=r, start_column=4, end_row=r, end_column=9); sc.row_dimensions[r].height = 22; r += 1
r += 1
sc.cell(r, 1, "Inherent risk tier requirements").font = Font(bold=True, color=GT, size=12); r += 1
for i, h in enumerate(["Tier", "Score from", "Assessment & control requirements", "", "", "Approval authority", "", "Review frequency", ""], 1):
    hdr(sc.cell(r, i, h))
sc.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5); sc.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7); sc.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)
r += 1
T_START = r
for n, lo, req, appr, freq in TIERS:
    box(sc.cell(r, 1, n), bold=True, fill=RAG[n]); box(sc.cell(r, 2, lo), align=CC); box(sc.cell(r, 3, req)); box(sc.cell(r, 6, appr)); box(sc.cell(r, 8, freq))
    sc.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5); sc.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7); sc.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)
    sc.row_dimensions[r].height = 42; r += 1
T_END = r - 1
for col, w in zip("ABCDEFGHI", [11, 16, 26, 22, 22, 22, 22, 22, 22]):
    sc.column_dimensions[col].width = w

# =========================================================== Reference: Taxonomy
tx = wb.create_sheet("Taxonomy")
title(tx, "AI Risk Taxonomy", f"{len(TAXONOMY)} risks in {len(CATEGORIES)} categories. Sources: NIST AI RMF & AI 600-1, ISO/IEC 23894, EU AI Act, OWASP LLM / Agentic Top 10, MITRE ATLAS, MIT AI Risk Repository, GT AI Ready7. RCM IDs refer to GT AI RCM v2.0.", 7)
tcols = [("Risk ID", 9), ("Category", 26), ("Risk", 34), ("Description / scenario", 60), ("Key sources", 34), ("RCM v2.0 controls", 18), ("Example KRI", 34)]
for i, (h, w) in enumerate(tcols, 1):
    hdr(tx.cell(4, i, h)); tx.column_dimensions[get_column_letter(i)].width = w
catname = dict(CATEGORIES)
for k, t in enumerate(TAXONOMY):
    rr = 5 + k
    vals = [t[0], catname[t[1]], t[2], t[3], t[4], t[5], t[6]]
    for j, v in enumerate(vals, 1): box(tx.cell(rr, j, v), bold=(j == 3))
    tx.row_dimensions[rr].height = 36
tx.freeze_panes = "D5"; tx.auto_filter.ref = f"A4:G{4+len(TAXONOMY)}"

# =========================================================== Lists (hidden)
ls = wb.create_sheet("Lists")
lists = {
    "A": ["Yes", "No"], "B": [x[1] for x in LIKELIHOOD], "C": [x[1] for x in IMPACT], "D": [x[0] for x in CONTROL_EFFECTIVENESS],
    "E": ["Avoid", "Mitigate", "Transfer", "Accept"], "F": ["Open", "Treatment planned", "In progress", "Closed", "Accepted"],
    "G": ["Concept / design", "Procurement", "Development", "Pilot / validation", "Deployment", "Operation", "Review", "Retirement"],
    "H": ["Provider (develops / places on market)", "Deployer (uses under own authority)", "Provider & deployer", "Importer / distributor"],
    "I": ["Generative AI / LLM assistant", "Agentic AI", "Predictive / classification ML", "Computer vision / biometrics", "Recommendation / personalisation", "Robotic / intelligent process automation", "Embedded AI in commercial software", "Other"],
}
for col, vals in lists.items():
    for i, v in enumerate(vals, 1): ls[f"{col}{i}"] = v
for k, f in enumerate(TIER_FACTORS):
    col = get_column_letter(12 + k)
    for i, v in enumerate(f[3], 1): ls[f"{col}{i}"] = v
ls.sheet_state = "hidden"


def dv_list(ws, rng, src):
    dv = DataValidation(type="list", formula1=src, allow_blank=True); ws.add_data_validation(dv); dv.add(rng); return dv


# =========================================================== 1. Use case profile
up = wb.create_sheet("1. Use Case Profile", 0)
title(up, "1. AI Use Case Profile", "Establish the context (ISO 31000 cl.6.3 / ISO/IEC 23894 cl.6.3). Yellow cells are inputs.", 4)
up.column_dimensions["A"].width = 3; up.column_dimensions["B"].width = 38; up.column_dimensions["C"].width = 80; up.column_dimensions["D"].width = 3
prof = [("Organisation", "Example: Bahrain retail bank (CBB-licensed)"),
        ("AI use case name", "Example: 'Ask Aya' - GenAI customer-service assistant with account-servicing actions"),
        ("Use case ID (AI register ref.)", "AI-2026-014"),
        ("Business owner", "Head of Retail Digital Channels"),
        ("Technical owner", "Head of Digital Engineering"),
        ("Description & intended purpose", "Arabic/English chatbot on mobile app answering product queries and performing low-value servicing actions (card freeze, statement request, standing-order changes up to BHD 500) via core-banking APIs. Built on a third-party LLM with retrieval over product documentation."),
        ("Out-of-scope / prohibited uses", "Credit decisions, investment advice, complaints adjudication, collection of card PINs/OTP"),
        ("AI system type", "Agentic AI"),
        ("Lifecycle stage", "Pilot / validation"),
        ("Organisation's role (EU AI Act terminology)", "Deployer (uses under own authority)"),
        ("Affected stakeholders", "Retail customers (incl. elderly and non-native speakers), contact-centre staff, operations"),
        ("Data used (sources & categories)", "Customer identity & account data (personal data), transaction history, product documents, chat transcripts"),
        ("Third parties / vendors", "LLM API provider (hosted in-region), conversational-AI platform vendor, cloud provider"),
        ("Jurisdictions / regulations in scope", "Bahrain PDPL; CBB Rulebook (Vol.1 - HC, OM, BC modules); iGA AI Policy (reference); EU AI Act Art. 50 transparency (EU-resident customers)"),
        ("Assessment date", "24/09/2026"),
        ("Assessor(s)", "GT Bahrain AI GRC team with bank Risk, Compliance, InfoSec, DPO")]
PROFILE_ROW = {}
for i, (k, v) in enumerate(prof):
    rr = 4 + i
    box(up.cell(rr, 2, k), bold=True, fill=LAV)
    c = up.cell(rr, 3, v if EXAMPLE else None); box(c, fill=INPUT)
    up.row_dimensions[rr].height = 48 if k.startswith("Description") else 22
    PROFILE_ROW[k] = rr
dv_list(up, f"C{PROFILE_ROW['AI system type']}", "=Lists!$I$1:$I$8")
dv_list(up, f"C{PROFILE_ROW['Lifecycle stage']}", "=Lists!$G$1:$G$8")
dv_list(up, f"C{PROFILE_ROW['Organisation' + chr(39) + 's role (EU AI Act terminology)']}", "=Lists!$H$1:$H$4")
NAME_REF = f"'1. Use Case Profile'!$C${PROFILE_ROW['AI use case name']}"

# =========================================================== 2. Prohibited screen
ps = wb.create_sheet("2. Prohibited Screen", 1)
title(ps, "2. Prohibited Practice Screening", "Answer every question. Any 'Yes' means the use case must not proceed without Legal review (EU AI Act Art. 5 - applicable since 2 Feb 2025; aligned with GCC AI Ethics Manual and Bahrain PDPL principles).", 4)
ps.column_dimensions["A"].width = 5; ps.column_dimensions["B"].width = 95; ps.column_dimensions["C"].width = 12; ps.column_dimensions["D"].width = 30
hdr(ps.cell(4, 1, "#")); hdr(ps.cell(4, 2, "Does the AI system...")); hdr(ps.cell(4, 3, "Answer")); hdr(ps.cell(4, 4, "Reference"))
PROH = [("deploy subliminal, manipulative or deceptive techniques that materially distort behaviour and cause significant harm?", "Art. 5(1)(a)"),
        ("exploit vulnerabilities due to age, disability or social/economic situation to distort behaviour causing significant harm?", "Art. 5(1)(b)"),
        ("evaluate or classify people based on social behaviour or personal traits leading to detrimental or disproportionate treatment (social scoring)?", "Art. 5(1)(c)"),
        ("assess the risk of a person committing a criminal offence based solely on profiling or personality traits?", "Art. 5(1)(d)"),
        ("create or expand facial-recognition databases through untargeted scraping of images from the internet or CCTV?", "Art. 5(1)(e)"),
        ("infer emotions of natural persons in the workplace or education institutions (except for medical or safety reasons)?", "Art. 5(1)(f)"),
        ("categorise people by biometric data to infer race, political opinions, trade-union membership, religious beliefs, sex life or sexual orientation?", "Art. 5(1)(g)"),
        ("perform real-time remote biometric identification in publicly accessible spaces for law-enforcement purposes?", "Art. 5(1)(h)")]
for i, (q, ref) in enumerate(PROH):
    rr = 5 + i
    box(ps.cell(rr, 1, i + 1), align=CC); box(ps.cell(rr, 2, q)); c = ps.cell(rr, 3, "No" if EXAMPLE else None); box(c, fill=INPUT, align=CC); box(ps.cell(rr, 4, "EU AI Act " + ref))
    ps.row_dimensions[rr].height = 30
dv_list(ps, "C5:C12", "=Lists!$A$1:$A$2")
box(ps.cell(14, 2, "Screening result"), bold=True, fill=LAV)
ps["C14"] = '=IF(COUNTA(C5:C12)<8,"Incomplete",IF(COUNTIF(C5:C12,"Yes")>0,"PROHIBITED","Pass"))'
box(ps["C14"], bold=True, align=CC)
ps.merge_cells("C14:D14")
ps.conditional_formatting.add("C14", CellIsRule(operator="equal", formula=['"PROHIBITED"'], fill=PatternFill("solid", fgColor="000000"), font=Font(bold=True, color="FFFFFF")))
ps.conditional_formatting.add("C14", CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", fgColor="C6E0B4"), font=Font(bold=True)))

# =========================================================== 3. Tiering
tr = wb.create_sheet("3. Risk Tiering", 2)
title(tr, "3. Inherent Risk Tiering", "Select the statement that best describes the use case for each factor. The weighted score (0-100) and override rules determine the inherent risk tier, which sets assessment depth, approval authority and review frequency (see Scales).", 7)
for col, w in zip("ABCDEFG", [6, 32, 70, 8, 8, 10, 3]):
    tr.column_dimensions[col].width = w
for i, h in enumerate(["#", "Factor", "Selected description (choose from list)", "Level (0-4)", "Weight", "Weighted"], 1):
    hdr(tr.cell(4, i, h))
EX_LEVELS = [2, 3, 4, 3, 3, 3, 3, 4, 3, 2]
for k, (fid, fname, w, levels) in enumerate(TIER_FACTORS):
    rr = 5 + k
    col = get_column_letter(12 + k)
    box(tr.cell(rr, 1, fid), align=CC); box(tr.cell(rr, 2, fname), bold=True)
    c = tr.cell(rr, 3, levels[EX_LEVELS[k]] if EXAMPLE else None); box(c, fill=INPUT)
    dv_list(tr, f"C{rr}", f"=Lists!${col}$1:${col}$5")
    tr.cell(rr, 4, f'=IF(C{rr}="","",MATCH(C{rr},Lists!${col}$1:${col}$5,0)-1)'); box(tr.cell(rr, 4), align=CC)
    box(tr.cell(rr, 5, w), align=CC)
    tr.cell(rr, 6, f'=IF(D{rr}="","",D{rr}*E{rr})'); box(tr.cell(rr, 6), align=CC)
    tr.row_dimensions[rr].height = 32
last = 5 + len(TIER_FACTORS) - 1
S = last + 2
labels = [("Factors answered", f'=COUNT(D5:D{last})&" of {len(TIER_FACTORS)}"'),
          ("Weighted score (0-100)", f'=IF(COUNT(D5:D{last})<{len(TIER_FACTORS)},"",ROUND(SUM(F5:F{last})/(4*SUM(E5:E{last}))*100,0))'),
          ("Score-based tier", f'=IF(C{S+1}="","",IF(C{S+1}>=75,"Critical",IF(C{S+1}>=50,"High",IF(C{S+1}>=25,"Medium","Low"))))'),
          ("Override triggered", f'=IF(C{S+1}="","",IF(AND(D5=4,D6=4),"Critical - significant decisions made autonomously",IF(OR(D5=4,D11=4),"Min. High - significant-effect decisions or high-risk classification",IF(AND(D5>=3,D6>=3),"Min. High - material decisions with limited oversight",IF(D7=4,"Min. Medium - sensitive data","None")))))'),
          ("Prohibited-practice screen", "='2. Prohibited Screen'!C14"),
          ("FINAL INHERENT RISK TIER", f'=IF(C{S+4}="PROHIBITED","Prohibited",IF(C{S+1}="","Incomplete",IF(LEFT(C{S+3},8)="Critical","Critical",IF(AND(LEFT(C{S+3},9)="Min. High",C{S+2}<>"Critical"),"High",IF(AND(LEFT(C{S+3},11)="Min. Medium",C{S+2}="Low"),"Medium",C{S+2})))))'),
          ("Assessment requirements", f'=IFERROR(VLOOKUP(C{S+5},Scales!$A${T_START}:$H${T_END},3,FALSE),IF(C{S+5}="Prohibited","Do not proceed. Refer to Legal and AI Committee.",""))'),
          ("Approval authority", f'=IFERROR(VLOOKUP(C{S+5},Scales!$A${T_START}:$H${T_END},6,FALSE),IF(C{S+5}="Prohibited","Legal + Executive Committee",""))'),
          ("Review frequency", f'=IFERROR(VLOOKUP(C{S+5},Scales!$A${T_START}:$H${T_END},8,FALSE),"")')]
for i, (k, f) in enumerate(labels):
    rr = S + i
    box(tr.cell(rr, 2, k), bold=True, fill=LAV)
    c = tr.cell(rr, 3, f); box(c, bold=(k.startswith("FINAL")), align=WC)
    tr.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
    tr.row_dimensions[rr].height = 30 if i >= 6 else 20
TIER_CELL = f"'3. Risk Tiering'!$C${S+5}"
SCORE_CELL = f"'3. Risk Tiering'!$C${S+1}"
rag_cf(tr, f"C{S+2}"); rag_cf(tr, f"C{S+5}")
tr.cell(S + 5, 3).font = Font(bold=True, size=12)

# =========================================================== 4. Risk register
rg = wb.create_sheet("4. Risk Register", 3)
title(rg, "4. AI Risk Register", "Mark each taxonomy risk as applicable (Yes/No), tailor the scenario, rate inherent likelihood & impact, assess existing control effectiveness and define treatment. Scores, ratings and heat map update automatically. Add bespoke risks in the blank rows at the bottom.", 20)
rcols = [("Applicable?", 9), ("Risk ID", 8), ("Category", 18), ("Risk", 26), ("Scenario (tailor to use case)", 44), ("Likelihood", 12), ("Impact", 12),
         ("Inherent score", 8), ("Inherent rating", 10), ("Existing / required controls (RCM v2.0)", 20), ("Control effectiveness", 14),
         ("Residual score", 8), ("Residual rating", 10), ("Within appetite?", 9), ("Treatment", 10), ("Treatment actions", 40), ("Risk owner", 16), ("Due date", 11), ("Status", 12), ("KRI", 28)]
for i, (h, w) in enumerate(rcols, 1):
    hdr(rg.cell(4, i, h)); rg.column_dimensions[get_column_letter(i)].width = w
rg.row_dimensions[4].height = 42
EX = {  # example ratings: id -> (L, I, CE, treatment, actions, owner, status)
    "GOV-03": (3, 4, "Satisfactory", "Mitigate", "Complete ISO/IEC 42005 impact assessment and AI Committee approval before go-live.", "Head of Retail Digital", "In progress"),
    "LEG-02": (2, 4, "Satisfactory", "Mitigate", "Confirm EU AI Act classification (not Annex III); implement Art. 50 disclosure for EU-resident users.", "Compliance", "Treatment planned"),
    "LEG-03": (3, 4, "Needs improvement", "Mitigate", "Map CBB BC / OM requirements; obtain Compliance sign-off; notify CBB per outsourcing rules if required.", "Compliance", "In progress"),
    "PRI-02": (4, 4, "Needs improvement", "Mitigate", "PII redaction before LLM call; DLP on transcripts; no-training & 30-day retention clause with provider.", "DPO", "In progress"),
    "PRI-03": (2, 3, "Strong", "Accept", "In-region hosting confirmed; document transfer assessment for failover region.", "DPO", "Closed"),
    "PER-02": (4, 3, "Needs improvement", "Mitigate", "Restrict answers to retrieved product documents with citations; refuse when confidence low; weekly QA sampling.", "Product Owner", "In progress"),
    "TRA-01": (2, 3, "Satisfactory", "Mitigate", "Display 'AI assistant' disclosure and handover-to-human option at session start.", "Product Owner", "Treatment planned"),
    "SEC-01": (4, 4, "Needs improvement", "Mitigate", "Prompt-injection filters; separate system/user/retrieved content; OWASP LLM red-team before launch.", "CISO", "In progress"),
    "SEC-04": (3, 5, "Needs improvement", "Mitigate", "Scoped OAuth tokens per action; transaction limits (BHD 500); step-up authentication for changes.", "CISO", "In progress"),
    "HUM-01": (3, 3, "Satisfactory", "Mitigate", "Agent-assist mode for complex queries; train contact-centre staff on verification.", "Head of Contact Centre", "Treatment planned"),
    "HUM-03": (3, 5, "Needs improvement", "Mitigate", "Human approval for standing-order changes; kill-switch; full action logging.", "Head of Digital Engineering", "In progress"),
    "TPR-02": (3, 3, "Satisfactory", "Mitigate", "Obtain model card, SOC 2 and ISO/IEC 42001 evidence; contractual incident notification (24h).", "Procurement", "Treatment planned"),
    "OPS-01": (2, 3, "Satisfactory", "Mitigate", "Fallback to human agents / static FAQ; SLA monitoring.", "IT Operations", "Closed"),
    "SOC-02": (3, 4, "Satisfactory", "Mitigate", "Complaint tagging for AI interactions; monthly review by Conduct Risk.", "Conduct Risk", "Treatment planned"),
    "SOC-03": (3, 4, "Satisfactory", "Mitigate", "Out-of-band verification for high-risk requests; voice-channel deepfake awareness.", "Fraud Risk", "In progress"),
    "FAI-02": (3, 3, "Needs improvement", "Mitigate", "Evaluate Arabic dialect coverage and accessibility for elderly users; expand test set.", "Data Science Lead", "Treatment planned"),
}
rows_n = len(TAXONOMY) + 10
lrow = 4 + rows_n
for k in range(rows_n):
    rr = 5 + k
    if k < len(TAXONOMY):
        t = TAXONOMY[k]
        vals = {2: t[0], 3: catname[t[1]], 4: t[2], 5: t[3], 10: t[5], 20: t[6]}
    else:
        vals = {2: f"BSP-{k - len(TAXONOMY) + 1:02d}"}
    for col in range(1, 21):
        c = rg.cell(rr, col, vals.get(col))
        box(c, fill=INPUT if col in (1, 5, 6, 7, 10, 11, 15, 16, 17, 18, 19, 20) else None, align=WT if col in (4, 5, 10, 16, 20) else WC)
    if EXAMPLE and k < len(TAXONOMY):
        tid = TAXONOMY[k][0]
        if tid in EX:
            Lx, Ix, ce, trt, act, own, st = EX[tid]
            rg.cell(rr, 1, "Yes"); rg.cell(rr, 6, LIKELIHOOD[Lx - 1][1]); rg.cell(rr, 7, IMPACT[Ix - 1][1]); rg.cell(rr, 11, ce)
            rg.cell(rr, 15, trt); rg.cell(rr, 16, act); rg.cell(rr, 17, own); rg.cell(rr, 19, st)
        else:
            rg.cell(rr, 1, "No")
    rg.cell(rr, 8, f'=IF(OR($A{rr}<>"Yes",F{rr}="",G{rr}=""),"",MATCH(F{rr},Lists!$B$1:$B$5,0)*MATCH(G{rr},Lists!$C$1:$C$5,0))')
    rg.cell(rr, 9, f'=IF(H{rr}="","",IF(H{rr}>=17,"Critical",IF(H{rr}>=10,"High",IF(H{rr}>=5,"Medium","Low"))))')
    rg.cell(rr, 12, f'=IF(OR(H{rr}="",K{rr}=""),"",ROUND(H{rr}*VLOOKUP(K{rr},Scales!$A${CE_START}:$B${CE_END},2,FALSE),1))')
    rg.cell(rr, 13, f'=IF(L{rr}="","",IF(L{rr}>=17,"Critical",IF(L{rr}>=10,"High",IF(L{rr}>=5,"Medium","Low"))))')
    rg.cell(rr, 14, f'=IF(M{rr}="","",IF(OR(M{rr}="Low",M{rr}="Medium"),"Yes","No"))')
    for col in (8, 9, 12, 13, 14): rg.cell(rr, col).alignment = CC
    rg.row_dimensions[rr].height = 44
dv_list(rg, f"A5:A{lrow}", "=Lists!$A$1:$A$2"); dv_list(rg, f"F5:F{lrow}", "=Lists!$B$1:$B$5"); dv_list(rg, f"G5:G{lrow}", "=Lists!$C$1:$C$5")
dv_list(rg, f"K5:K{lrow}", "=Lists!$D$1:$D$4"); dv_list(rg, f"O5:O{lrow}", "=Lists!$E$1:$E$4"); dv_list(rg, f"S5:S{lrow}", "=Lists!$F$1:$F$5")
rag_cf(rg, f"I5:I{lrow}"); rag_cf(rg, f"M5:M{lrow}")
rg.conditional_formatting.add(f"N5:N{lrow}", CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="F4B183")))
rg.conditional_formatting.add(f"A5:T{lrow}", FormulaRule(formula=['$A5="No"'], font=Font(color="A6A6A6")))
rg.freeze_panes = "E5"; rg.auto_filter.ref = f"A4:T{lrow}"
REG = "'4. Risk Register'"

# =========================================================== 5. Heat map & summary
hm = wb.create_sheet("5. Heat Map & Summary", 4)
title(hm, "5. Risk Profile - Heat Map & Summary", None, 14)
hm.merge_cells("A2:N2")
hm["A2"] = f'="Use case: "&IF({NAME_REF}="","(not specified)",{NAME_REF})&"   |   Inherent tier: "&{TIER_CELL}&"   |   Tiering score: "&{SCORE_CELL}'
hm["A2"].font = Font(bold=True, color=GT, size=11)
for col in "ABCDEFGHIJKLMN": hm.column_dimensions[col].width = 11
hm.column_dimensions["A"].width = 16
hm["A4"] = "Inherent risk heat map (number of applicable risks)"; hm["A4"].font = Font(bold=True, color=GT, size=12)
# grid: rows likelihood 5..1, cols impact 1..5
for j in range(5):
    c = hm.cell(5, 2 + j, IMPACT[j][1]); hdr(c)
for i in range(5):
    lv = 5 - i
    rr = 6 + i
    hdr(hm.cell(rr, 1, LIKELIHOOD[lv - 1][1]))
    for j in range(5):
        iv = j + 1
        c = hm.cell(rr, 2 + j, f'=COUNTIFS({REG}!$A$5:$A${lrow},"Yes",{REG}!$F$5:$F${lrow},"{LIKELIHOOD[lv-1][1]}",{REG}!$G$5:$G${lrow},"{IMPACT[iv-1][1]}")')
        sc_ = lv * iv
        band = "Critical" if sc_ >= 17 else "High" if sc_ >= 10 else "Medium" if sc_ >= 5 else "Low"
        c.fill = PatternFill("solid", fgColor=RAG[band]); c.alignment = CC; c.border = B
        c.font = Font(bold=True, size=14, color="FFFFFF" if band == "Critical" else "000000")
    hm.row_dimensions[rr].height = 34
hm.cell(11, 1, "Likelihood / Impact").font = Font(italic=True, size=8)
# summary table
hm["H4"] = "Risk count by rating"; hm["H4"].font = Font(bold=True, color=GT, size=12)
for j, h in enumerate(["Rating", "Inherent", "Residual"]): hdr(hm.cell(5, 8 + j, h))
for i, band in enumerate(["Critical", "High", "Medium", "Low"]):
    rr = 6 + i
    box(hm.cell(rr, 8, band), bold=True, fill=RAG[band], align=CC)
    if band == "Critical": hm.cell(rr, 8).font = Font(bold=True, color="FFFFFF")
    hm.cell(rr, 9, f'=COUNTIF({REG}!$I$5:$I${lrow},"{band}")'); box(hm.cell(rr, 9), align=CC)
    hm.cell(rr, 10, f'=COUNTIF({REG}!$M$5:$M${lrow},"{band}")'); box(hm.cell(rr, 10), align=CC)
box(hm.cell(10, 8, "Total rated"), bold=True); hm.cell(10, 9, "=SUM(I6:I9)"); hm.cell(10, 10, "=SUM(J6:J9)")
box(hm.cell(10, 9), bold=True, align=CC); box(hm.cell(10, 10), bold=True, align=CC)
box(hm.cell(11, 8, "Outside appetite (residual High/Critical)"), bold=True)
hm.merge_cells("H11:I11"); hm.cell(11, 10, "=J6+J7"); box(hm.cell(11, 10), bold=True, align=CC)
hm.conditional_formatting.add("J11", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F4B183")))
# by category
hm["A14"] = "Residual risk by category"; hm["A14"].font = Font(bold=True, color=GT, size=12)
for j, h in enumerate(["Category", "", "", "Applicable", "Critical", "High", "Medium", "Low", "Avg residual"]):
    if h: hdr(hm.cell(15, 1 + j, h))
hm.merge_cells("A15:C15")
for i, (cid, cname) in enumerate(CATEGORIES):
    rr = 16 + i
    hm.cell(rr, 1, cname); hm.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3); box(hm.cell(rr, 1))
    hm.cell(rr, 4, f'=COUNTIFS({REG}!$C$5:$C${lrow},"{cname}",{REG}!$A$5:$A${lrow},"Yes")')
    for j, band in enumerate(["Critical", "High", "Medium", "Low"]):
        hm.cell(rr, 5 + j, f'=COUNTIFS({REG}!$C$5:$C${lrow},"{cname}",{REG}!$M$5:$M${lrow},"{band}")')
    hm.cell(rr, 9, f'=IFERROR(ROUND(AVERAGEIFS({REG}!$L$5:$L${lrow},{REG}!$C$5:$C${lrow},"{cname}",{REG}!$A$5:$A${lrow},"Yes"),1),"-")')
    for col in range(4, 10): box(hm.cell(rr, col), align=CC)
cat_last = 16 + len(CATEGORIES) - 1
ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100
ch.title = "Residual risks by category"; ch.y_axis.title = "Number of risks"
data = Reference(hm, min_col=5, max_col=8, min_row=15, max_row=cat_last)
cats = Reference(hm, min_col=1, min_row=16, max_row=cat_last)
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
for s, col in zip(ch.series, ["C00000", "ED7D31", "FFC000", "70AD47"]):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ch.height = 9; ch.width = 18; ch.legend.position = "b"
hm.add_chart(ch, "K14")
# top residual risks
tt = cat_last + 3
hm.cell(tt, 1, "Top 10 residual risks (highest residual score)").font = Font(bold=True, color=GT, size=12)
for j, h in enumerate(["Rank", "Risk ID", "Risk", "", "", "", "Residual score", "Residual rating", "Treatment", "Owner", "Status"]):
    if h: hdr(hm.cell(tt + 1, 1 + j, h))
hm.merge_cells(start_row=tt + 1, start_column=3, end_row=tt + 1, end_column=6)
# helper: unique ranking key in hidden column of register (V)
for k in range(rows_n):
    rr = 5 + k
    rg.cell(rr, 22, f'=IF(L{rr}="","",L{rr}+ROW()/100000)')
rg.column_dimensions["V"].hidden = True
for i in range(10):
    rr = tt + 2 + i
    hm.cell(rr, 1, i + 1)
    key = f'LARGE({REG}!$V$5:$V${lrow},{i+1})'
    m = f'MATCH({key},{REG}!$V$5:$V${lrow},0)'
    hm.cell(rr, 2, f'=IFERROR(INDEX({REG}!$B$5:$B${lrow},{m}),"")')
    hm.cell(rr, 3, f'=IFERROR(INDEX({REG}!$D$5:$D${lrow},{m}),"")')
    hm.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
    hm.cell(rr, 7, f'=IFERROR(INDEX({REG}!$L$5:$L${lrow},{m}),"")')
    hm.cell(rr, 8, f'=IFERROR(INDEX({REG}!$M$5:$M${lrow},{m}),"")')
    hm.cell(rr, 9, f'=IFERROR(INDEX({REG}!$O$5:$O${lrow},{m}),"")')
    hm.cell(rr, 10, f'=IFERROR(INDEX({REG}!$Q$5:$Q${lrow},{m}),"")')
    hm.cell(rr, 11, f'=IFERROR(INDEX({REG}!$S$5:$S${lrow},{m}),"")')
    for col in range(1, 12):
        box(hm.cell(rr, col), align=CC if col in (1, 2, 7, 8, 9, 11) else WC)
    hm.row_dimensions[rr].height = 22
rag_cf(hm, f"H{tt+2}:H{tt+11}")
# sign-off
so = tt + 13
hm.cell(so, 1, "Approval & sign-off").font = Font(bold=True, color=GT, size=12)
hm.cell(so + 1, 1, "Required approval authority:"); hm.cell(so + 1, 1).font = Font(bold=True)
hm.cell(so + 1, 4, "='3. Risk Tiering'!C" + str(S + 7)); hm.merge_cells(start_row=so + 1, start_column=4, end_row=so + 1, end_column=11)
for i, h in enumerate(["Role", "", "", "Name", "", "", "Decision (Approve / Approve with conditions / Reject)", "", "", "Date", ""]):
    if h: hdr(hm.cell(so + 2, 1 + i, h))
for a, b in [(1, 3), (4, 6), (7, 9), (10, 11)]:
    hm.merge_cells(start_row=so + 2, start_column=a, end_row=so + 2, end_column=b)
for i, role in enumerate(["Use-case / business owner", "Risk management (2nd line)", "Compliance / Legal", "Approval authority"]):
    rr = so + 3 + i
    hm.cell(rr, 1, role)
    for a, b in [(1, 3), (4, 6), (7, 9), (10, 11)]:
        hm.merge_cells(start_row=rr, start_column=a, end_row=rr, end_column=b)
        for col in range(a, b + 1): hm.cell(rr, col).border = B
    hm.row_dimensions[rr].height = 26
hm.page_setup.orientation = "landscape"; hm.page_setup.fitToWidth = 1; hm.page_setup.fitToHeight = 0
hm.sheet_properties.pageSetUpPr.fitToPage = True

# =========================================================== Guide
gd = wb.create_sheet("Guide", 0)
title(gd, "GT AI Risk Assessment Toolkit" + (" - Worked Example" if EXAMPLE else ""), "Companion to the GT AI Risk Assessment Framework (v1.0, 2026). Grant Thornton Bahrain - AI Governance, Risk & Compliance.", 3)
gd.column_dimensions["A"].width = 4; gd.column_dimensions["B"].width = 30; gd.column_dimensions["C"].width = 100
steps = [("Step 1 - Use Case Profile", "Document context: purpose, owners, lifecycle stage, organisation's role, stakeholders, data, vendors and applicable regulation (ISO 31000 cl.6.3; ISO/IEC 42005 cl.6)."),
         ("Step 2 - Prohibited Screen", "Screen against prohibited practices. Any 'Yes' stops the assessment and requires Legal review."),
         ("Step 3 - Risk Tiering", "Rate 10 inherent-risk factors. The weighted score plus override rules produce the tier (Low / Medium / High / Critical), which determines assessment depth, approval authority and review frequency."),
         ("Step 4 - Risk Register", f"Work through the {len(TAXONOMY)} taxonomy risks: mark applicability, tailor scenarios, rate likelihood & impact, evaluate existing control effectiveness (mapped to RCM v2.0 controls) and plan treatment. Add bespoke risks in rows BSP-01..10."),
         ("Step 5 - Heat Map & Summary", "Review the inherent heat map, residual profile, top-10 residual risks and obtain sign-off from the approval authority required by the tier."),
         ("Monitor & review", "Re-assess at the review frequency for the tier or on material change (model/provider change, new data, new purpose, expanded user base, incidents). Track KRIs listed per risk."),
         ("Colour legend", "Yellow = input cell | Purple header = calculated / reference | Green/Amber/Orange/Red = Low / Medium / High / Critical"),
         ("Scoring logic", "Inherent score = Likelihood (1-5) x Impact (1-5). Residual score = Inherent score x control-effectiveness factor (Strong 0.4, Satisfactory 0.6, Needs improvement 0.8, Weak/None 1.0). Bands: Low 1-4, Medium 5-9, High 10-16, Critical 17-25. Residual Low/Medium = within appetite (default; calibrate to the organisation's approved AI risk appetite - RCM RM-6).")]
for i, (k, v) in enumerate(steps):
    rr = 4 + i
    box(gd.cell(rr, 2, k), bold=True, fill=LAV); box(gd.cell(rr, 3, v))
    gd.row_dimensions[rr].height = 44
if EXAMPLE:
    gd.cell(13, 2, "About this example").font = Font(bold=True, color="C00000")
    gd.cell(13, 3, "Illustrative worked example for a hypothetical Bahrain retail bank deploying an agentic GenAI customer-service assistant. Ratings are indicative and for training / demonstration only.").alignment = WT
    gd.row_dimensions[13].height = 32
for ws_ in wb.worksheets:
    ws_.sheet_properties.tabColor = GT if ws_.title[0].isdigit() or ws_.title == "Guide" else "A6A6A6"
    if ws_.title not in ("5. Heat Map & Summary",):
        ws_.page_setup.orientation = "landscape"; ws_.page_setup.fitToWidth = 1; ws_.page_setup.fitToHeight = 0
        ws_.sheet_properties.pageSetUpPr.fitToPage = True
order = ["Guide", "1. Use Case Profile", "2. Prohibited Screen", "3. Risk Tiering", "4. Risk Register", "5. Heat Map & Summary", "Taxonomy", "Scales", "Lists"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(OUT)
print("saved", OUT, "register rows", rows_n)
