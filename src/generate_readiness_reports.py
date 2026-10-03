# -*- coding: utf-8 -*-
"""GT AI Readiness - report generator (GT internal).

Turns questionnaire responses into a client PDF report and an Outlook email draft, without the client ever
seeing the scoring engine.

Input : a responses file - either the Microsoft Forms Excel export, or GT_AI_Readiness_Responses_Template.xlsx.
        Column headers must contain the question tag, e.g. "[Q7] ..." (the questionnaire titles include it).
        Answers may be the full option text or the option letter (A-E).
Output: for each response, in OUT_DIR/<organisation>/:
        - <org>_AI_Readiness_Report.pdf      (Report sheet only - send this to the client)
        - <org>_AI_Readiness_Engine.xlsx     (filled engine - GT internal, do not send)
        - <org>_email_draft.eml              (opens in Outlook as an unsent draft with the PDF attached)

Usage : python generate_readiness_reports.py RESPONSES.xlsx ENGINE.xlsx OUT_DIR [--gt-contact "Name <email>"]
Needs : Python 3.9+, openpyxl, pypdf, LibreOffice (soffice) installed.
"""
import sys, os, re, shutil, subprocess, tempfile, datetime, argparse, glob
from email.message import EmailMessage
import openpyxl
from pypdf import PdfReader, PdfWriter

REPORT_MARK = "Readiness & Risk Assessment - Confidential"   # footer text printed on every report page
LETTERS = "ABCDE"


def soffice():
    for c in [shutil.which("soffice"), shutil.which("libreoffice"),
              r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"]:
        if c and os.path.exists(c):
            return c
    sys.exit("LibreOffice (soffice) not found - install it from libreoffice.org")


def convert(src, fmt, outdir):
    prof = "file:///" + os.path.join(tempfile.gettempdir(), "gt_lo_profile").replace("\\", "/")
    subprocess.run([soffice(), f"-env:UserInstallation={prof}", "--headless", "--convert-to", fmt, "--outdir", outdir, src],
                   check=True, capture_output=True, timeout=600)
    ext = fmt.split(":")[0]
    return os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + "." + ext)


def norm(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s).lower()).strip()


def read_responses(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    heads = [c.value for c in ws[1]]
    tags = {}
    for j, h in enumerate(heads):
        m = re.search(r"\[(P\d+|A\d|Q\d+)\]", str(h or ""))
        if m: tags[m.group(1)] = j
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in row): continue
        rows.append({t: row[j] for t, j in tags.items()})
    return rows


def engine_maps(eng):
    prof = {}   # label -> row
    ws = eng["1. Profile"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value: prof[str(ws.cell(r, 2).value)] = r
    asx = eng["2. Assessment"]; qrow = {}
    for r in range(1, asx.max_row + 1):
        v = asx.cell(r, 2).value
        if isinstance(v, str) and re.fullmatch(r"A\d|Q\d+", v): qrow[v] = r
    opts = {}
    qb = eng["QBank"]
    for r in range(2, qb.max_row + 1):
        if qb.cell(r, 1).value: opts.setdefault(qb.cell(r, 1).value, []).append(qb.cell(r, 2).value)
    return prof, qrow, opts


def match_option(ans, options):
    if ans is None or str(ans).strip() == "": return None
    a = str(ans).strip()
    if len(a) == 1 and a.upper() in LETTERS[:len(options)]: return options[LETTERS.index(a.upper())]
    m = re.match(r"^([A-E])[\.\)]\s+", a)
    if m: a = a[m.end():]
    na = norm(a)
    for o in options:
        if norm(o) == na: return o
    for o in options:
        if norm(o).startswith(na[:40]) or na.startswith(norm(o)[:40]): return o
    raise ValueError(f"answer not recognised: {ans!r}")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("responses"); ap.add_argument("engine"); ap.add_argument("out")
    ap.add_argument("--gt-contact", default="Grant Thornton Bahrain - AI Governance, Risk & Compliance")
    a = ap.parse_args()
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from tool_data import PROFILE_Q
    os.makedirs(a.out, exist_ok=True)
    for resp in read_responses(a.responses):
        eng = openpyxl.load_workbook(a.engine)
        prof, qrow, opts = engine_maps(eng)
        org = str(resp.get("P1") or "Organisation").strip()
        safe = re.sub(r"[^A-Za-z0-9]+", "_", org).strip("_")[:60]
        odir = os.path.join(a.out, safe); os.makedirs(odir, exist_ok=True)
        pws = eng["1. Profile"]
        for pid, _t, label, _o in PROFILE_Q:
            if label in prof and resp.get(pid) not in (None, ""): pws.cell(prof[label], 3).value = resp.get(pid)
        if "Assessment date (dd/mm/yyyy)" in prof:
            pws.cell(prof["Assessment date (dd/mm/yyyy)"], 3).value = datetime.date.today().strftime("%d/%m/%Y")
        missing = []
        for qid, r in qrow.items():
            try:
                v = match_option(resp.get(qid), opts[qid])
            except ValueError as e:
                sys.exit(f"{org} {qid}: {e}")
            if v is None: missing.append(qid)
            eng["2. Assessment"].cell(r, 5).value = v
        if missing: print(f"WARNING {org}: unanswered {missing} - report will be marked provisional")
        eng.active = eng.sheetnames.index("3. Report")
        xlsx = os.path.join(odir, f"{safe}_AI_Readiness_Engine.xlsx"); eng.save(xlsx)
        tmp = tempfile.mkdtemp()
        # recalculated copy -> headline values for the email
        recalc = convert(xlsx, "xlsx:Calc MS Excel 2007 XML", tmp)
        vals = openpyxl.load_workbook(recalc, data_only=True)["Calc"]
        overall, band, tier, expo = vals["B45"].value, vals["B46"].value, vals["C32"].value, vals["B57"].value
        # PDF of the whole workbook -> keep report pages only
        full = convert(xlsx, "pdf", tmp)
        rd = PdfReader(full); w = PdfWriter()
        for pg in rd.pages:
            if REPORT_MARK in (pg.extract_text() or ""): w.add_page(pg)
        pdf = os.path.join(odir, f"{safe}_AI_Readiness_Report.pdf")
        with open(pdf, "wb") as f: w.write(f)
        # email draft
        msg = EmailMessage(); msg["To"] = str(resp.get("P8") or ""); msg["Subject"] = f"{org} - AI Readiness & Risk Assessment report"
        msg["X-Unsent"] = "1"
        name = str(resp.get("P6") or "").split(" ")[0] or "Sir / Madam"
        pct = f"{overall:.0%}" if isinstance(overall, (int, float)) else "-"
        msg.set_content(
            f"Dear {name},\n\nThank you for completing the AI Readiness & Risk Assessment. Please find attached your report.\n\n"
            f"Headline results:\n  - Overall AI readiness: {pct} ({band})\n  - Inherent AI risk tier: {tier}\n  - Overall AI risk exposure: {expo}\n\n"
            "The report sets out your results by domain, the key findings, your top priority actions and a suggested roadmap. "
            "We would be pleased to walk you through the results in a short debrief session.\n\n"
            f"Kind regards,\n{a.gt_contact}\n")
        with open(pdf, "rb") as f:
            msg.add_attachment(f.read(), maintype="application", subtype="pdf", filename=os.path.basename(pdf))
        with open(os.path.join(odir, f"{safe}_email_draft.eml"), "wb") as f: f.write(bytes(msg))
        shutil.rmtree(tmp, ignore_errors=True)
        print(f"{org}: readiness {pct} ({band}), tier {tier}, exposure {expo}, {len(w.pages)} report pages -> {odir}")


if __name__ == "__main__":
    main()
