# -*- coding: utf-8 -*-
"""Client-facing GT AI Readiness Questionnaire (Word) + response capture template (Excel).
No scores, weights or levels are shown. The Word file is formatted for Microsoft Forms 'Quick import'
(numbered questions, lettered options); question titles carry an ID tag ([P1], [A1], [Q5] ...) that the
GT report generator uses to map responses back to the scoring engine."""
import sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
sys.path.insert(0, os.path.dirname(__file__))
from tool_data import CONTEXT, READINESS, PILLARS, PROFILE_Q

DOCX, XLSX = sys.argv[1], sys.argv[2]
GT = RGBColor(0x4F, 0x2D, 0x7F)
LET = "ABCDE"


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)


doc = Document()
s = doc.sections[0]; s.page_height = Cm(29.7); s.page_width = Cm(21); s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for lvl in (1, 2):
    h = doc.styles[f"Heading {lvl}"]; h.font.name = "Calibri"; h.font.color.rgb = GT; h.font.size = Pt(16 if lvl == 1 else 12.5)
    rpr = h.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(a), "Calibri")

t = doc.add_table(rows=1, cols=1); c = t.rows[0].cells[0]; shade(c, "4F2D7F"); c.text = ""
p = c.paragraphs[0]; r = p.add_run("AI Readiness & Risk Assessment - Questionnaire"); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(255, 255, 255)
p2 = c.add_paragraph(); r = p2.add_run("Grant Thornton Bahrain  |  AI Governance, Risk & Compliance"); r.font.size = Pt(11); r.font.color.rgb = RGBColor(0xE8, 0xE0, 0xF3)
doc.add_paragraph()
doc.add_heading("Before you start", 1)
for txt in ["This questionnaire takes about 20-25 minutes. It has three parts: your organisation profile (14 questions), how AI is used (4 questions) and how AI is governed and managed (16 questions across 5 areas).",
            "For each question choose ONE answer - the statement that best describes the situation today. If you are unsure between two answers, choose the lower one: Grant Thornton may ask for supporting evidence for the answer selected.",
            "Answers are confidential and used only to prepare your AI Readiness & Risk Assessment report. Grant Thornton analyses the responses and sends the report to the email address you provide.",
            "You can complete this questionnaire online (link provided by Grant Thornton) or fill in this document and return it to your Grant Thornton contact."]:
    doc.add_paragraph(txt)

n = 0


def question(qid, title, guide, options):
    global n
    n += 1
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{n}. [{qid}] {title}"); r.bold = True
    if guide:
        g = doc.add_paragraph(); g.paragraph_format.keep_with_next = True; rr = g.add_run(guide); rr.italic = True; rr.font.size = Pt(9); rr.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
    if options:
        for i, o in enumerate(options):
            op = doc.add_paragraph(); op.paragraph_format.left_indent = Cm(0.8); op.paragraph_format.space_after = Pt(1)
            op.paragraph_format.keep_with_next = i < len(options) - 1
            op.add_run(f"{LET[i] if len(options) <= 5 else str(i + 1)}. {o}")
    else:
        a = doc.add_paragraph("Answer: ______________________________________________"); a.paragraph_format.left_indent = Cm(0.8)


doc.add_heading("Part 1 - Organisation profile", 1)
for pid, title, _lab, opts in PROFILE_Q:
    question(pid, title, None, opts)
doc.add_page_break()
doc.add_heading("Part 2 - How AI is used in your organisation", 1)
doc.add_paragraph("These questions describe your AI use and its potential impact. There are no right or wrong answers.")
for q in CONTEXT:
    question(q["id"], q["q"], q["guide"], [o[0] for o in q["opts"]])
doc.add_page_break()
doc.add_heading("Part 3 - How AI is governed and managed", 1)
for d in PILLARS:
    doc.add_heading(d, 2)
    for q in [x for x in READINESS if x["pillar"] == d]:
        question(q["id"], q["q"], q["guide"].replace("Evidence:", "Examples of evidence:"), q["opts"])
doc.add_heading("Thank you", 1)
doc.add_paragraph("Please submit the online form, or return this document to your Grant Thornton contact. Your report will normally be issued within 5 working days of receipt.")
for sec in doc.sections:
    fp = sec.footer.paragraphs[0]; fp.text = "Grant Thornton Bahrain - AI Readiness & Risk Assessment Questionnaire - Confidential"; fp.runs[0].font.size = Pt(8)
doc.save(DOCX)

# ---------------------------------------------------------------- response capture template
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Responses"
lists = wb.create_sheet("Lists"); col = 1
heads = [("Response ID", None)] + [(f"[{pid}] {title}", opts) for pid, title, _l, opts in PROFILE_Q] + \
        [(f"[{q['id']}] {q['topic']}", [o[0] for o in q["opts"]]) for q in CONTEXT] + [(f"[{q['id']}] {q['topic']}", q["opts"]) for q in READINESS]
thin = Side(style="thin", color="C9C9C9")
for j, (h, opts) in enumerate(heads, 1):
    c = ws.cell(1, j, h); c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = PatternFill("solid", fgColor="4F2D7F")
    c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 28 if opts else 22
    if opts:
        for i, o in enumerate(opts, 1): lists.cell(i, col, o)
        L = openpyxl.utils.get_column_letter(col)
        dv = DataValidation(type="list", formula1=f"=Lists!${L}$1:${L}${len(opts)}", allow_blank=True); ws.add_data_validation(dv); dv.add(f"{openpyxl.utils.get_column_letter(j)}2:{openpyxl.utils.get_column_letter(j)}200")
        col += 1
ws.row_dimensions[1].height = 60; ws.freeze_panes = "B2"
lists.sheet_state = "hidden"
wb.save(XLSX)
print("saved", DOCX, XLSX, "questions:", n)
