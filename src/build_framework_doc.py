# -*- coding: utf-8 -*-
"""GT AI Risk Assessment Framework - Word document."""
import sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.path.insert(0, os.path.dirname(__file__))
from risk_data import *
import diagrams

OUT = sys.argv[1]
IMG = "/home/user/work/img"; os.makedirs(IMG, exist_ok=True)
diagrams.process_diagram(f"{IMG}/process.png")
diagrams.heatmap(f"{IMG}/heatmap.png")
diagrams.lines_of_defence(f"{IMG}/lod.png")

GT = RGBColor(0x4F, 0x2D, 0x7F); GTHEX = "4F2D7F"; LAVHEX = "EDE7F6"
doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8)

st = doc.styles
st["Normal"].font.name = "Calibri"; st["Normal"].font.size = Pt(10.5)
st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st["Normal"].paragraph_format.space_after = Pt(6); st["Normal"].paragraph_format.line_spacing = 1.12
for lvl, size in [(1, 17), (2, 13), (3, 11.5)]:
    h = st[f"Heading {lvl}"]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = GT
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
    rpr = h.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(a), "Calibri")
    h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10); h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)


def set_cell_margins(table, m=60):
    tblPr = table._tbl.tblPr; mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(m if side in ("top", "bottom") else 90)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)


def table(headers, rows, widths=None, size=8.5, header_fill=GTHEX, first_col_bold=False, zebra=True, fills=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_margins(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        p = c.paragraphs[0]; r = p.add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, header_fill); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = str(v).split("\n")
            for li, line in enumerate(lines):
                p = cells[i].paragraphs[0] if li == 0 else cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(line); run.font.size = Pt(size)
                if first_col_bold and i == 0: run.bold = True
            if zebra and ri % 2 == 1: shade(cells[i], "F7F4FB")
            if fills and (ri, i) in fills: shade(cells[i], fills[(ri, i)])
    if widths:
        t.autofit = False
        tot = sum(widths); scale = 17.0 / tot if tot > 17.0 else 1.0
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(Cm(widths[i] * scale).twips)))
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w * scale)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def para(text, bold=False, italic=False, size=None, color=None, align=None, after=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p


def rich(parts):
    p = doc.add_paragraph()
    for txt, b in parts:
        r = p.add_run(txt); r.bold = b
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style); p.paragraph_format.space_after = Pt(2)
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else:
            p.add_run(it)


def callout(title, text, fill="EDE7F6"):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; shade(c, fill); c.text = ""
    p = c.paragraphs[0]; r = p.add_run(title); r.bold = True; r.font.color.rgb = GT; r.font.size = Pt(10)
    p2 = c.add_paragraph(); r2 = p2.add_run(text); r2.font.size = Pt(9.5)
    set_cell_margins(t, 120)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def field(paragraph, instr):
    r = paragraph.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r2 = paragraph.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2._r.append(it)
    r3 = paragraph.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate"); r3._r.append(f2)
    r4 = paragraph.add_run("Right-click and select 'Update Field' to build the table of contents." if "TOC" in instr else "1")
    r5 = paragraph.add_run(); f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end"); r5._r.append(f3)


# ------------------------------------------------------------------ cover
for _ in range(3): doc.add_paragraph()
t = doc.add_table(rows=1, cols=1); c = t.rows[0].cells[0]; shade(c, GTHEX); c.text = ""
p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(18)
r = p.add_run("GT AI Risk Assessment Framework"); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor(255, 255, 255)
p2 = c.add_paragraph(); r = p2.add_run("Methodology for identifying, assessing, treating and monitoring risks across the AI lifecycle"); r.font.size = Pt(13); r.font.color.rgb = RGBColor(0xE8, 0xE0, 0xF3)
p3 = c.add_paragraph(); p3.paragraph_format.space_after = Pt(18)
r = p3.add_run("Aligned to ISO/IEC 23894  |  ISO/IEC 42001  |  ISO/IEC 42005  |  NIST AI RMF & AI 600-1  |  EU AI Act  |  Bahrain PDPL & iGA AI Policy  |  NSW AIAF"); r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(255, 255, 255)
set_cell_margins(t, 200)
for _ in range(2): doc.add_paragraph()
para("Grant Thornton Bahrain  |  AI Governance, Risk & Compliance Service Line", bold=True, color=GT, size=12)
para("Version 1.0  |  September 2026", size=10.5)
para("Classification: Internal / Client-ready methodology", italic=True, size=9.5)
for _ in range(10): doc.add_paragraph()
para("This framework is a Grant Thornton Bahrain methodology document. It is intended to be tailored to each client's context, risk appetite and regulatory obligations. It does not constitute legal advice.", italic=True, size=8.5, color=RGBColor(0x59, 0x59, 0x59))
doc.add_page_break()

# ------------------------------------------------------------------ doc control + TOC
doc.add_heading("Document control", 1)
table(["Item", "Detail"], [["Document", "GT AI Risk Assessment Framework"], ["Version", "1.0 (initial release)"], ["Date", "24 September 2026"],
                            ["Owner", "Grant Thornton Bahrain - AI GRC Service Line"], ["Companion artefacts", "GT AI Risk Assessment Toolkit (Excel) - template and worked example\nGT AI RCM v2.0 (2026) - control library referenced by this framework\nGT AI Readiness Assessment Tool - organisation-level readiness"],
                            ["Review cycle", "Annually, or earlier on material regulatory change (e.g., EU AI Act Digital Omnibus outcome, Bahrain AI law enactment)"]],
      widths=[4, 13], first_col_bold=True)
doc.add_heading("Contents", 1)
field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
doc.add_page_break()

# ------------------------------------------------------------------ 1 exec summary
doc.add_heading("1. Executive summary", 1)
para("Artificial intelligence introduces risks that conventional IT, model or operational risk frameworks do not fully capture: probabilistic and opaque behaviour, dependence on data and foundation-model providers, new attack surfaces such as prompt injection, impacts on individuals' rights, and fast-moving regulation such as the EU AI Act and the region's emerging AI policies. Organisations in Bahrain are adopting AI quickly - the national General Policy for the Use of AI (iGA, 2025) and Bahrain Economic Vision 2030 actively encourage it - but most do not yet have a structured, repeatable way to decide which AI uses are acceptable and under what conditions.")
para("The GT AI Risk Assessment Framework (GT-AIRAF) provides that method. It combines the ISO 31000 / ISO/IEC 23894 risk-management process with AI-specific practices from NIST AI RMF, ISO/IEC 42005 impact assessment, the EU AI Act risk-based classification and the tiered self-assessment model of the NSW AI Assessment Framework. Its key features are:")
bullets([("Proportionality by design - ", "a 10-factor inherent-risk tiering (Low / Medium / High / Critical) determines assessment depth, approval authority and review frequency, so low-risk productivity uses are not slowed down while high-impact systems receive rigorous scrutiny."),
         ("A comprehensive AI risk taxonomy - ", f"{len(TAXONOMY)} risks in {len(CATEGORIES)} categories, spanning governance, legal, fairness, privacy, data, performance, transparency, security, human oversight, third-party, operational and societal risks, including emerging agentic-AI and AI-enabled threat risks."),
         ("Multi-dimensional impact assessment - ", "impacts on individuals and rights are assessed alongside financial, regulatory, operational, reputational and safety/societal impacts, on a common 5x5 scale."),
         ("Direct link to controls - ", "every risk maps to controls in the GT AI Risk & Control Matrix (RCM v2.0, 72 controls), which in turn map to ISO/IEC 42001, EU AI Act, NIST AI RMF, Bahrain PDPL and CBB requirements."),
         ("Lifecycle coverage - ", "assessments are triggered at defined lifecycle gates and on material change, with KRIs for continuous monitoring."),
         ("Practical tooling - ", "an Excel toolkit automates tiering, scoring, heat maps, top-risk reporting and sign-off.")])
callout("How the framework fits the GT AI GRC service line", "Organisation level: the GT AI Readiness Assessment Tool (and AI Ready7) measures capability maturity.  Control level: the GT AI RCM v2.0 defines what 'good' looks like.  Use-case level: this framework assesses and treats risk for each AI system and feeds the AI register, risk register and management reporting.")

# ------------------------------------------------------------------ 2 purpose scope
doc.add_heading("2. Purpose, scope and application", 1)
doc.add_heading("2.1 Purpose", 2)
bullets(["Provide a consistent, defensible and auditable method to identify, analyse, evaluate, treat and monitor risks from the development, procurement and use of AI systems.",
         "Enable informed go / no-go and conditional-approval decisions proportionate to the risk of each AI use case.",
         "Demonstrate conformity with ISO/IEC 42001 clauses 6.1.2 (AI risk assessment), 6.1.3 (AI risk treatment) and 6.1.4 (AI system impact assessment), EU AI Act Art. 9 / Art. 26-27, and supervisory expectations of regulators such as the CBB.",
         "Generate the evidence (risk register, impact assessment, approvals) required by auditors, certification bodies and regulators."])
doc.add_heading("2.2 Scope", 2)
para("The framework applies to any 'AI system' as defined by the OECD / EU AI Act: a machine-based system that, for explicit or implicit objectives, infers from the input it receives how to generate outputs such as predictions, content, recommendations or decisions that can influence physical or virtual environments. In scope are:")
bullets(["AI developed in-house, customised by vendors, or procured as commercial off-the-shelf or SaaS services;",
         "AI features embedded in existing enterprise software (e.g., CRM, ERP, productivity suites);",
         "generative AI and large language model (LLM) applications, retrieval-augmented generation (RAG) and AI agents;",
         "traditional machine-learning models (e.g., credit scoring, fraud detection, forecasting) and intelligent automation;",
         "third-party AI used by suppliers to deliver services to the organisation."])
para("Publicly available general-purpose AI tools used for low-risk productivity tasks may be governed by the AI Acceptable Use Policy (RCM GL-5) rather than a full assessment, provided the policy explicitly authorises that use - consistent with the NSW AIAF approach.")
doc.add_heading("2.3 When to apply the framework", 2)
table(["Trigger type", "Assessment points"],
      [["Lifecycle gates", "Concept / design (initial tiering)  -  Procurement (vendor due diligence)  -  Pilot / validation  -  Pre-deployment (full assessment & approval)  -  Operation (periodic review per tier)  -  Retirement"],
       ["Material change", "New or changed purpose, user population or geography  -  change of model, provider or model version  -  new data sources or categories  -  reduced human oversight or increased autonomy (e.g., adding agent actions)  -  integration with new systems"],
       ["Risk signals", "KRI threshold breach  -  incident or near miss  -  complaint or contest trend  -  new regulation or guidance  -  new threat intelligence (e.g., new jailbreak technique)  -  discovery of unassessed (shadow) AI"]],
      widths=[3.5, 13.5], first_col_bold=True)

# ------------------------------------------------------------------ 3 standards basis
doc.add_heading("3. Standards basis and alignment", 1)
para("The framework is deliberately standards-based so that outputs are recognisable to auditors, certification bodies and regulators. The table below shows how each reference is used.")
table(["Reference", "How it is used in GT-AIRAF"],
      [["ISO 31000:2018 / ISO/IEC 23894:2023", "Core process (context, identification, analysis, evaluation, treatment, monitoring, communication) and AI-specific risk sources."],
       ["ISO/IEC 42001:2023", "Management-system requirements for AI risk assessment (6.1.2), treatment (6.1.3), impact assessment (6.1.4), Annex A controls and objectives."],
       ["ISO/IEC 42005:2025", "Structure and content of AI system impact assessments (individuals, groups, society)."],
       ["NIST AI RMF 1.0 & AI 600-1 (GenAI Profile)", "Govern-Map-Measure-Manage functions; trustworthiness characteristics; 12 generative-AI risk categories used in the taxonomy."],
       ["EU AI Act (Reg. 2024/1689)", "Prohibited practices (Art. 5), high-risk classification (Art. 6, Annex III), risk management (Art. 9), deployer duties and FRIA (Art. 26-27), transparency (Art. 50)."],
       ["NSW AI Assessment Framework (AIAF)", "Tiered self-assessment model: risk triggers, Low-Critical risk bands, escalation to an AI review committee, contestability principle."],
       ["OWASP Top 10 for LLM / Agentic Applications; MITRE ATLAS; NIST AI 100-2", "Security risk identification and threat modelling for AI systems."],
       ["Bahrain PDPL (Law 30/2018); iGA General Policy for the Use of AI (2025); GCC AI Ethics Manual; CBB Rulebook", "Local legal and regulatory criteria in impact scales, taxonomy and treatment."],
       ["MIT AI Risk Repository; OECD AI Incidents Monitor", "Completeness check of the taxonomy; incident-based likelihood calibration."],
       ["GT AI Ready7 & GT AI RCM v2.0", "Organisation-level maturity inputs and the control library used for treatment."]],
      widths=[5.5, 11.5], first_col_bold=True)

# ------------------------------------------------------------------ 4 principles
doc.add_heading("4. Guiding principles", 1)
table(["Principle", "What it means in practice"],
      [["Proportionate", "Effort and approval levels scale with inherent risk tier; Low-tier uses follow a light path."],
       ["Lifecycle-based", "Risk is assessed at every gate and re-assessed on material change - not a one-off exercise."],
       ["Socio-technical", "Risks to people, rights and society are considered alongside technical, financial and legal risks."],
       ["Evidence-based", "Ratings are supported by testing, monitoring data, incident history and documentation."],
       ["Human accountability", "Every AI system has named business and technical owners; AI never 'owns' a decision."],
       ["Integrated", "AI risk feeds the enterprise risk register, uses the enterprise risk scales and escalation routes, and links to existing privacy, security and model-risk processes."],
       ["Transparent & contestable", "Assessments are documented; affected persons have a route to challenge AI-supported decisions."]],
      widths=[4, 13], first_col_bold=True)

# ------------------------------------------------------------------ 5 governance
doc.add_heading("5. Governance, roles and responsibilities", 1)
para("The framework operates within a three-lines model with board-level oversight of AI risk appetite.")
doc.add_picture(f"{IMG}/lod.png", width=Cm(16.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
table(["Activity", "Use-case owner", "Tech / data team", "Risk (2nd line)", "Privacy / Legal / Compliance", "InfoSec", "AI Committee", "Internal Audit"],
      [["Use-case profile & tiering", "A/R", "C", "C", "C", "C", "I", "-"],
       ["Prohibited-practice screening", "R", "C", "C", "A", "-", "I", "-"],
       ["Risk identification & analysis", "A", "R", "C", "C", "C", "I", "-"],
       ["Impact assessment (High/Critical)", "A", "R", "R", "R", "C", "I", "-"],
       ["Independent validation / challenge", "I", "C", "A/R", "C", "R", "I", "-"],
       ["Treatment plan & controls", "A", "R", "C", "C", "R", "I", "-"],
       ["Approval - Low / Medium", "A", "-", "C", "C", "C", "I", "-"],
       ["Approval - High / Critical", "R", "C", "C", "C", "C", "A", "-"],
       ["KRI monitoring & re-assessment", "A", "R", "R", "C", "C", "I", "-"],
       ["Independent assurance", "I", "I", "I", "I", "I", "I", "A/R"]],
      widths=[4.2, 1.8, 1.8, 1.8, 2.2, 1.5, 1.8, 1.8], size=7.5, first_col_bold=True)
para("R = Responsible, A = Accountable, C = Consulted, I = Informed.", italic=True, size=8.5)

# ------------------------------------------------------------------ 6 process
doc.add_heading("6. The AI risk assessment process", 1)
para("The process has seven stages. Stages 0-1 are completed for every AI use case; stages 2-6 are completed at a depth determined by the tier.")
doc.add_picture(f"{IMG}/process.png", width=Cm(17))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Stage 0 - Intake and context establishment", 2)
para("The use-case owner registers the AI use case in the AI register (RCM GL-8) and completes the use-case profile. The objective is a shared understanding of what the system is intended to do, for whom and under which constraints.")
bullets(["Intended purpose, out-of-scope uses and success measures;", "AI system type (GenAI, agentic, predictive ML, computer vision, embedded AI, etc.) and lifecycle stage;",
         "the organisation's role - provider, deployer, importer or distributor - which determines EU AI Act obligations;", "affected stakeholders, including vulnerable groups;",
         "data sources and categories (personal, sensitive, confidential, client data);", "third parties and hosting locations (relevant for PDPL cross-border rules);", "applicable laws, regulations, contracts and internal policies."])

doc.add_heading("Stage 1 - Screening and inherent risk tiering", 2)
rich([("Step 1a - Prohibited-practice screen. ", True), ("Eight questions based on EU AI Act Art. 5 (manipulation, exploitation of vulnerabilities, social scoring, predictive policing on profiling alone, untargeted facial scraping, workplace/education emotion recognition, sensitive biometric categorisation, real-time remote biometric identification). Any 'Yes' stops the assessment pending Legal review.", False)])
rich([("Step 1b - Inherent risk tiering. ", True), ("Ten factors are rated 0-4 before considering controls. Each factor has a weight reflecting its contribution to potential harm. The weighted score is normalised to 0-100.", False)])
table(["#", "Factor", "Weight", "Level 0 (lowest)", "Level 4 (highest)"],
      [[f[0], f[1], f[2], f[3][0], f[3][4]] for f in TIER_FACTORS], widths=[1, 3.6, 1.3, 5.3, 5.8], size=7.5, first_col_bold=False)
para("Tier thresholds and override rules:", bold=True)
table(["Tier", "Score", "Assessment requirements", "Approval authority", "Review"],
      [[n, f">= {lo}" if lo else "0-24", req, appr, freq] for n, lo, req, appr, freq in TIERS],
      widths=[1.8, 1.6, 6.8, 3.6, 3.2], size=8, first_col_bold=True,
      fills={(0, 0): "C6E0B4", (1, 0): "FFE699", (2, 0): "F4B183", (3, 0): "E06666"})
bullets(["Decisions with legal or similarly significant effects (T1 = 4) or high-risk regulatory classification such as EU AI Act Annex III (T7 = 4) - minimum tier High.",
         "Material decisions about individuals (T1 >= 3) with limited human oversight (T2 >= 3) - minimum tier High.",
         "Significant-effect decisions made autonomously (T1 = 4 and T2 = 4) - tier Critical.",
         "Processing of sensitive personal or regulated data (T3 = 4) - minimum tier Medium."])
para("Override rules mirror the AIAF 'pattern risk' concept: certain combinations are high-risk regardless of the average score.", italic=True, size=9)

doc.add_heading("Stage 2 - Risk identification", 2)
para("Risks are identified by working systematically through the GT AI risk taxonomy (Section 7), supplemented by:")
bullets(["structured workshops with the use-case owner, data scientists, IT, risk, privacy, legal and business users;",
         "an AI system impact assessment per ISO/IEC 42005 for High and Critical tiers - covering intended and unintended impacts on individuals, groups and society, and a Fundamental Rights Impact Assessment where EU AI Act Art. 27 applies;",
         "threat modelling using OWASP LLM / Agentic Top 10 and MITRE ATLAS for systems with external exposure or agent capabilities;",
         "review of incidents in comparable systems (AI Incident Database, OECD AIM) and vendor documentation (model cards, system cards);",
         "a data-protection impact assessment where personal data is processed (Bahrain PDPL)."])
para("Each identified risk is recorded as a scenario: cause -> event -> consequence (e.g., 'Indirect prompt injection in a retrieved document causes the assistant to disclose another customer's balance, resulting in a PDPL breach and customer harm').")

doc.add_heading("Stage 3 - Risk analysis", 2)
para("Each risk is rated on inherent likelihood and inherent impact, assuming no controls (or only controls inherent to the design). The inherent score is Likelihood x Impact (1-25).")
table(["Score", "Likelihood", "Definition"], [[s, n, d] for s, n, d in LIKELIHOOD], widths=[1.3, 2.7, 13], size=8.5, first_col_bold=True)
para("Impact is rated on the highest applicable of six dimensions. Financial thresholds are indicative for a mid-sized Bahraini organisation and must be calibrated to the client's enterprise risk scales.", size=9.5)
table(["Score"] + IMPACT_DIMS, [[f"{s} {n}"] + ds for s, n, ds in IMPACT], widths=[2.1, 2.7, 2.2, 2.5, 2.4, 2.6, 2.5], size=7.2, first_col_bold=True)
doc.add_picture(f"{IMG}/heatmap.png", width=Cm(10.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Stage 4 - Risk evaluation", 2)
para("Existing controls are identified (using RCM v2.0 control IDs) and their effectiveness rated. The residual score is the inherent score multiplied by the control-effectiveness factor. Residual risk is then compared with the organisation's AI risk appetite (RCM RM-6).")
table(["Control effectiveness", "Factor", "Definition"], [[n, f, d] for n, f, d in CONTROL_EFFECTIVENESS], widths=[3.5, 1.5, 12], size=8.5, first_col_bold=True)
table(["Residual rating", "Score", "Required response"], [[n, f"{lo}-{hi}", d] for n, lo, hi, d in RATING_BANDS], widths=[3, 1.8, 12.2], size=8.5, first_col_bold=True,
      fills={(0, 0): "C6E0B4", (1, 0): "FFE699", (2, 0): "F4B183", (3, 0): "E06666"})
callout("Default risk appetite", "Unless the board has approved a different AI risk appetite: residual Low and Medium risks are within appetite; residual High risks require a time-bound treatment plan approved by the AI Committee; residual Critical risks are outside appetite - the system must not be deployed (or must be suspended) until the risk is reduced, unless formally accepted by the Executive Committee with board notification.")

doc.add_heading("Stage 5 - Risk treatment", 2)
table(["Option", "When to use", "Examples for AI systems"],
      [["Avoid", "Risk is unacceptable and cannot be reduced cost-effectively; prohibited practice", "Do not deploy; remove a high-risk feature (e.g., autonomous payment execution); restrict scope"],
       ["Mitigate", "Controls can reduce likelihood or impact to within appetite", "Guardrails, human-in-the-loop, retrieval grounding, bias mitigation, PII redaction, least-privilege agent permissions, monitoring"],
       ["Transfer / share", "Part of the risk can be borne by a third party", "Contractual indemnities and SLAs with AI vendors; cyber / professional-indemnity insurance"],
       ["Accept", "Residual risk is within appetite, or formally accepted by the correct authority", "Documented acceptance with rationale, conditions, owner and review date"]],
      widths=[2.4, 5.4, 9.2], size=8.5, first_col_bold=True)
para("Treatment plans specify the actions, RCM controls to implement, owner, due date and expected residual rating. For High and Critical tiers, treatment effectiveness must be verified by testing (e.g., bias testing, red teaming, independent validation - RCM RS-5, AA-3, LC-9) before approval.")

doc.add_heading("Stage 6 - Monitoring, review and reporting", 2)
bullets([("KRIs: ", "each material risk has at least one key risk indicator with thresholds and a mandatory response (e.g., drift score, hallucination rate, override rate, blocked injection attempts, complaints)."),
         ("Re-assessment: ", "per tier frequency (Critical quarterly, High six-monthly, Medium annually, Low every 24 months) and on any trigger in Section 2.3."),
         ("Reporting: ", "portfolio-level AI risk dashboard to the AI Committee quarterly and to the board risk committee at least semi-annually - tier distribution, residual High/Critical risks, overdue actions, incidents and KRI breaches."),
         ("Incidents: ", "AI incidents are managed through RCM IM-1 to IM-3, including regulatory notification (EU AI Act Art. 73, PDPL, CBB) where required, and feed back into risk ratings.")])

# ------------------------------------------------------------------ 7 taxonomy
doc.add_heading("7. GT AI risk taxonomy", 1)
para(f"The taxonomy contains {len(TAXONOMY)} risks in {len(CATEGORIES)} categories. It consolidates the NIST AI 600-1 generative-AI risks, ISO/IEC 23894 risk sources, EU AI Act requirements, OWASP and MITRE ATLAS security threats and the GT AI Ready7 pillars. Each risk is linked to RCM v2.0 controls and an example KRI in the toolkit.")
catname = dict(CATEGORIES)
for cid, cname in CATEGORIES:
    doc.add_heading(f"7.{[c[0] for c in CATEGORIES].index(cid)+1} {cname}", 3)
    table(["ID", "Risk", "Description / scenario", "RCM controls"],
          [[t[0], t[2], t[3], t[5]] for t in TAXONOMY if t[1] == cid], widths=[1.5, 4, 8.8, 2.7], size=7.8, first_col_bold=True)

# ------------------------------------------------------------------ 8 tier requirements
doc.add_heading("8. Minimum requirements by tier", 1)
table(["Requirement", "Low", "Medium", "High", "Critical"],
      [["AI register entry & owner", "Yes", "Yes", "Yes", "Yes"],
       ["Prohibited-practice screen & tiering", "Yes", "Yes", "Yes", "Yes"],
       ["Risk register (taxonomy walkthrough)", "Key risks only", "Yes", "Yes", "Yes"],
       ["DPIA (if personal data)", "Screening", "Yes", "Yes", "Yes"],
       ["AI system impact assessment (ISO/IEC 42005) / FRIA", "-", "Light", "Full", "Full + external input"],
       ["Security threat model (OWASP / ATLAS)", "-", "If external exposure", "Yes", "Yes"],
       ["Bias / fairness testing", "-", "If affects individuals", "Yes", "Yes + independent"],
       ["Independent validation (LC-9) / red teaming (AA-3)", "-", "-", "Yes", "Yes (pre-deployment)"],
       ["Human oversight design (RS-1)", "Policy", "Documented", "Documented & tested", "Documented, tested, staffed"],
       ["Legal review", "-", "If flagged", "Yes", "Yes"],
       ["Approval", "Owner", "Function head + Risk", "AI Committee", "ExCo / Board Risk Committee"],
       ["Staged rollout & kill-switch", "-", "-", "Recommended", "Mandatory"],
       ["KRI monitoring", "-", "Periodic", "Continuous", "Continuous + real-time alerting"],
       ["Re-assessment", "24 months", "12 months", "6 months", "Quarterly"]],
      widths=[5.5, 2.3, 2.9, 3, 3.3], size=8, first_col_bold=True)

# ------------------------------------------------------------------ 9 integration
doc.add_heading("9. Integration with the GT AI GRC toolset", 1)
table(["Artefact", "Relationship to this framework"],
      [["GT AI Readiness Assessment Tool", "Organisation-level: identifies capability gaps (e.g., missing inventory, weak monitoring) that increase likelihood ratings across use cases; its inherent-risk questions reuse the tiering factors T1-T3."],
       ["GT AI Ready7", "Pillar maturity scores inform control-effectiveness ratings (e.g., Protect AI Systems maturity informs security controls)."],
       ["GT AI RCM v2.0", "Control library for treatment; the taxonomy maps each risk to control IDs; control test results inform effectiveness ratings."],
       ["AI register (RCM GL-8)", "Records tier, residual rating, approval and next review date for every AI system."],
       ["Enterprise risk register", "High and Critical residual AI risks are escalated into the enterprise risk register using common scales."],
       ["ISO/IEC 42001 AIMS", "Framework outputs constitute the documented information required by clauses 6.1.2-6.1.4 and 8.2-8.4."]],
      widths=[5, 12], first_col_bold=True)

# ------------------------------------------------------------------ 10 worked example
doc.add_heading("10. Worked example (summary)", 1)
para("Use case: agentic GenAI customer-service assistant ('Ask Aya') for a CBB-licensed retail bank, answering product questions in Arabic and English and performing low-value servicing actions (card freeze, statement requests, standing-order changes up to BHD 500) through core-banking APIs. The full example is provided in the toolkit (Worked Example file).")
table(["Stage", "Outcome"],
      [["Screening", "No prohibited practice identified."],
       ["Tiering", "Weighted score 74 / 100 - tier HIGH (drivers: autonomy of actions, sensitive financial personal data, customer-facing agentic GenAI, vulnerable customers). Approval: AI Committee; six-monthly review."],
       ["Identification", "16 applicable taxonomy risks; ISO/IEC 42005 impact assessment and DPIA required."],
       ["Analysis", "Highest inherent risks: prompt injection (16), data leakage (16), unbounded agent actions (15), insecure integration / excessive privilege (15)."],
       ["Evaluation", "4 residual High risks outside appetite before treatment (SEC-01, PRI-02, HUM-03, SEC-04)."],
       ["Treatment", "PII redaction and no-training vendor clauses; prompt-injection filters and content separation; scoped OAuth tokens and BHD 500 limit; human approval for standing-order changes; kill-switch; OWASP red-team before launch; Art. 50 disclosure."],
       ["Monitoring", "KRIs: blocked injection attempts, DLP events, hallucination rate from weekly QA sampling, agent actions without approval, AI-related complaints."]],
      widths=[3, 14], first_col_bold=True)

# ------------------------------------------------------------------ appendices
doc.add_heading("Appendix A - Risk register fields", 1)
table(["Field", "Description"],
      [["Risk ID / category / title", "From the taxonomy or bespoke (BSP-nn)"], ["Scenario", "Cause -> event -> consequence, specific to the use case"],
       ["Inherent likelihood / impact / score / rating", "Before controls; 1-5 scales; score 1-25"], ["Existing / required controls", "RCM v2.0 control IDs"],
       ["Control effectiveness", "Strong / Satisfactory / Needs improvement / Weak-None"], ["Residual score / rating / within appetite", "Inherent x effectiveness factor"],
       ["Treatment option & actions", "Avoid / Mitigate / Transfer / Accept; specific actions"], ["Owner, due date, status", "Accountable individual and tracking"],
       ["KRI", "Indicator, threshold and response"]], widths=[6, 11], first_col_bold=True)
doc.add_heading("Appendix B - Glossary", 1)
table(["Term", "Definition"],
      [["AI system", "Machine-based system that infers from inputs how to generate outputs such as predictions, content, recommendations or decisions (OECD / EU AI Act Art. 3(1))."],
       ["Provider / deployer", "Provider develops an AI system or places it on the market under its name; deployer uses an AI system under its authority (EU AI Act Art. 3)."],
       ["General-purpose AI (GPAI) model", "Model displaying significant generality, capable of a wide range of tasks (e.g., large language models)."],
       ["Agentic AI", "AI systems that plan and take actions via tools, APIs or other systems with some autonomy."],
       ["Inherent / residual risk", "Risk before / after considering the effect of existing controls."],
       ["FRIA", "Fundamental Rights Impact Assessment (EU AI Act Art. 27)."],
       ["KRI", "Key risk indicator - metric providing early warning of increasing risk exposure."],
       ["Hallucination / confabulation", "Generation of plausible but false or unsupported content."],
       ["Prompt injection", "Manipulation of an LLM through crafted input, directly or via content it processes."]], widths=[4.5, 12.5], first_col_bold=True)
doc.add_heading("Appendix C - Key references", 1)
bullets(["ISO/IEC 23894:2023 Information technology - AI - Guidance on risk management", "ISO/IEC 42001:2023 AI management system; ISO/IEC 42005:2025 AI system impact assessment", "ISO 31000:2018 Risk management - Guidelines",
         "NIST AI 100-1 AI Risk Management Framework (2023); NIST AI 600-1 Generative AI Profile (2024); NIST AI 100-2 E2025 Adversarial Machine Learning", "Regulation (EU) 2024/1689 (Artificial Intelligence Act)",
         "OWASP Top 10 for LLM Applications 2025; OWASP Top 10 for Agentic Applications; MITRE ATLAS", "Kingdom of Bahrain Law No. 30 of 2018 (Personal Data Protection Law); iGA General Policy for the Use of AI v1.0 (2025); GCC Guiding Manual on the Ethics of AI Use",
         "Central Bank of Bahrain Rulebook (HC, RM, OM, BC modules)", "Digital NSW - NSW AI Assessment Framework (AIAF)", "MIT AI Risk Repository; OECD AI Incidents Monitor"])

# ------------------------------------------------------------------ header / footer
for s in doc.sections:
    hp = s.header.paragraphs[0]; hp.text = ""; r = hp.add_run("Grant Thornton Bahrain  |  GT AI Risk Assessment Framework v1.0"); r.font.size = Pt(8); r.font.color.rgb = GT
    fp = s.footer.paragraphs[0]; fp.text = ""; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("Page "); r.font.size = Pt(8)
    field(fp, "PAGE")
    s.different_first_page_header_footer = True
# update fields on open
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)
doc.save(OUT)
print("saved", OUT)
