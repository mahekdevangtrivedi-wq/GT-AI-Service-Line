# -*- coding: utf-8 -*-
"""GT AI Opportunity Discovery Toolkit (Excel) - where exactly does AI fit in an organisation?

Usage: build_opportunity_toolkit.py OUT.xlsx [gt|bank]     (no example key = blank template)
"""
import sys, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.pagebreak import Break
sys.path.insert(0, os.path.dirname(__file__))
import printfix
from opportunity_data import *
from opportunity_library import LIBRARY, FUNCTIONS as LFUNC, FUNCTION_ORDER, SECTORS as LSECT, SECTOR_ORDER, SECTOR_PROFILES

OUT = sys.argv[1]
EXKEY = sys.argv[2] if len(sys.argv) > 2 else None
EX = EXAMPLES.get(EXKEY) if EXKEY else None
EXAMPLE = EX is not None

GT = "4F2D7F"; GT2 = "7B5BA6"; LAV = "EDE7F6"; LAV2 = "D9CCEB"; INPUT = "FFF9E6"; GREY = "F2F2F2"; DARK = "3B2160"
TEAL = "00A7B5"; CORAL = "E8705F"
VFILL = {"AI fit now": "A9D08E", "Automate - AI not needed": "9BC2E6", "AI later - close data gap first": "FFE699",
         "Stretch - partner or later stage": "F4B183", "Redesign / standardise first": "D9D9D9"}
RFILL = {"Low": "C6E0B4", "Medium": "FFE699", "High": "F4B183"}
thin = Side(style="thin", color="C9C9C9"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
WT = Alignment(wrap_text=True, vertical="top"); WC = Alignment(wrap_text=True, vertical="center")
CC = Alignment(wrap_text=True, vertical="center", horizontal="center")
wb = openpyxl.Workbook()
NF = len(FAMILIES); NL = len(LIBRARY)
FUNCTIONS = [LFUNC[k][0] for k in FUNCTION_ORDER]
SECTOR_LIST = [LSECT[k] for k in SECTOR_ORDER if k != "XI"] + ["Other"]


def title(ws, t, sub, span):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(1, 1, t); c.font = Font(size=16, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=GT)
    c.alignment = Alignment(vertical="center", indent=1); ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    s = ws.cell(2, 1, sub); s.font = Font(italic=True, size=9.5, color=GT); s.alignment = WC; ws.row_dimensions[2].height = 32


def cell(ws, r, c, v, bold=False, fill=None, align=WC, size=9.5, color="000000", fmt=None, border=True, italic=False):
    x = ws.cell(r, c, v); x.font = Font(size=size, bold=bold, color=color, italic=italic); x.alignment = align
    if border: x.border = B
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    if fmt: x.number_format = fmt
    return x


def band(ws, r, text, span, fill=GT, c0=1):
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + span - 1)
    x = ws.cell(r, c0, text); x.font = Font(bold=True, color="FFFFFF", size=10.5); x.fill = PatternFill("solid", fgColor=fill)
    x.alignment = Alignment(vertical="center", indent=1); ws.row_dimensions[r].height = 22


def merged(ws, r, c1, c2, v, **kw):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    for c in range(c1, c2 + 1): ws.cell(r, c).border = B
    if kw.get("fill"):
        for c in range(c1, c2 + 1): ws.cell(r, c).fill = PatternFill("solid", fgColor=kw["fill"])
    return cell(ws, r, c1, v, **kw)


def dv_list(ws, src, rng, prompt=None, strict=True):
    dv = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=strict)
    if prompt: dv.prompt = prompt; dv.showInputMessage = True
    ws.add_data_validation(dv); dv.add(rng)


def verdict_cf(ws, rng):
    for v, f in VFILL.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'], fill=PatternFill("solid", fgColor=f)))


def risk_cf(ws, rng):
    for v, f in RFILL.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'], fill=PatternFill("solid", fgColor=f)))


def portrait(ws, landscape=False, fit_h=0):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = fit_h; ws.sheet_properties.pageSetUpPr.fitToPage = True


# ================================================================== Lists (hidden)
L = wb.active; L.title = "Lists"
L["A1"] = "Task nature"; L["B1"] = "Solution family"; L["C1"] = "Level"; L["D1"] = "AI?"; L["E1"] = "Data req"; L["F1"] = "Potential"
L["G1"] = "Typical solutions"; L["H1"] = "KPIs"; L["I1"] = "Design notes"
for i, f in enumerate(FAMILIES, 2):
    for j, v in enumerate(f, 1): L.cell(i, j, v)
LA = f"Lists!$A$2:$A${NF + 1}"; LB = f"Lists!$B$2:$B${NF + 1}"


def fam(col):
    return f"Lists!${col}$2:${col}${NF + 1}"


def put_pairs(col, rows, start=2):
    for i, (a, b) in enumerate(rows, start):
        L.cell(i, col, a); L.cell(i, col + 1, b)
    return f"Lists!${get_column_letter(col)}${start}:${get_column_letter(col)}${start + len(rows) - 1}", \
           f"Lists!${get_column_letter(col + 1)}${start}:${get_column_letter(col + 1)}${start + len(rows) - 1}"


DA_O, DA_V = put_pairs(13, DATA_AVAIL, 2)
VO_O, VO_V = put_pairs(13, VOLUME, 8)
ST_O, ST_V = put_pairs(13, STANDARD, 13)
AU_O, AU_V = put_pairs(13, AUTONOMY, 18)
for i, v in enumerate(YESNO, 23): L.cell(i, 13, v)
YN = "Lists!$M$23:$M$24"
for i, v in enumerate(FUNCTIONS, 2): L.cell(i, 16, v)
FN = f"Lists!$P$2:$P${len(FUNCTIONS) + 1}"
for i, v in enumerate(SECTOR_LIST, 2): L.cell(i, 17, v)
SEC = f"Lists!$Q$2:$Q${len(SECTOR_LIST) + 1}"
POS_RANGE = {}
for k, (qid, axis, q, g, opts) in enumerate(POSITION_Q):
    col = 19 + k
    L.cell(1, col, qid)
    for i, o in enumerate(opts, 2): L.cell(i, col, o)
    cl = get_column_letter(col); POS_RANGE[qid] = f"Lists!${cl}$2:${cl}$5"
for i, (code, name, ceil, head, desc) in enumerate(ARCHETYPES, 2):
    for j, v in enumerate((code, name, ceil, head, desc), 33): L.cell(i, j, v)
ARC = lambda col: f"Lists!${col}$2:${col}$5"
for i, (lvl, name) in enumerate(LEVELS, 2): L.cell(i, 38, lvl); L.cell(i, 39, name)
for i, (t, c) in enumerate(CONTROLS, 2): L.cell(i, 41, t); L.cell(i, 42, c)
for i, (v, m, n) in enumerate(VERDICTS, 2): L.cell(i, 44, v); L.cell(i, 45, m); L.cell(i, 46, n)
for i, h in enumerate(HORIZONS, 2): L.cell(i, 48, h)
VER = "Lists!$AR$2:$AR$6"
for i in range(1, 51): L.cell(i + 1, 50, i)
L.sheet_state = "hidden"

# ================================================================== Guide
g = wb.create_sheet("Guide", 0)
title(g, "GT AI Opportunity Discovery Toolkit", "Where exactly does AI fit in this organisation - and where is automation enough? "
      "A facilitated workshop tool of the GT AI Strategy & Roadmap service, built on a 240-entry, standards-based AI use-case library.", 3)
g.column_dimensions["A"].width = 3; g.column_dimensions["B"].width = 30; g.column_dimensions["C"].width = 108
r = 4
band(g, r, "WHAT THIS TOOLKIT ANSWERS", 2, c0=2); r += 1
for k, v in [("1. Is the organisation positioned for AI?", "Do you have the DATA (digital, joined-up, reliable, with history and outcomes) and the SCALE / CAPACITY (volumes, people, budget, "
                                                     "technology team) for AI - or is automation enough for now? Result: one of four archetypes and a ceiling on solution complexity."),
             ("2. Where exactly does AI fit?", "Process by process, selected from the GT AI Use-Case Library for the client's sector: which solution family fits (automate, off-the-shelf AI, "
                                               "configured AI, custom / agentic AI), is the data there, what is the risk tier - and what exactly GT recommends."),
             ("3. What should be done first?", "Value x feasibility, risk-adjusted priority, horizon (H1 / H2 / H3), opportunity report, roadmap lanes and one-page use-case canvases."),
             ("Not a readiness test", "Governance and capability maturity are measured by the GT AI Readiness & Risk Assessment (AI Ready7-based). This toolkit locates the "
                                      "opportunities and the right type of solution for each.")]:
    cell(g, r, 2, k, bold=True, fill=LAV); cell(g, r, 3, v, align=WT); g.row_dimensions[r].height = 44; r += 1
r += 1; band(g, r, "HOW TO USE IT (WORKSHOP SEQUENCE)", 2, c0=2); r += 1
steps = [("1. AI Position", "Select the sector, then answer 12 questions with leadership, IT and finance (about 30 minutes). Data Position and Scale & Capacity scores place the organisation in an archetype."),
         ("2. Use-Case Library", f"{NL} use cases: {sum(1 for x in LIBRARY if x[1] == 'Cross-industry')} cross-industry plus {sum(1 for x in LIBRARY if x[1] != 'Cross-industry')} sector-specific entries "
                                 "(banking, insurance, capital markets, healthcare, telecom, government, retail, hospitality, manufacturing, energy, logistics, real estate, education, professional services). "
                                 "Each entry carries the GT recommendation, data needed, KPIs, regulatory notes and a minimum risk tier."),
         ("3. AI Fit Diagnostic", "Select up to 50 use cases from the library (the list shows cross-industry + the selected sector) and answer 9 inputs each (yellow). Custom processes can be typed in - "
                                  "then select the task nature and function. Everything else calculates, including the library recommendation."),
         ("4. Opportunity Report", "The client-ready output: AI position, verdict mix, where AI fits by function, the top 10 opportunities with GT recommendations and the sector lens."),
         ("5. Roadmap", "Opportunities placed automatically in H1 / H2 / H3 lanes plus the enabling foundations they depend on."),
         ("6. Use-Case Canvas", "Pick an opportunity number; a one-page canvas is produced (recommendation, data, KPIs, risk tier, regulatory notes, minimum controls, next step)."),
         ("7. Sector Profiles", "Value at stake, automation & AI potential, value pools and key regulators for each sector."),
         ("Reference", "Solution families, archetypes, verdicts, minimum controls, function taxonomy (APQC) and the standards behind each step."),
         ("Colour legend", "Yellow = input  |  Purple / lavender = header or calculated  |  Verdict colours: green = AI fit now, blue = automate, yellow = data gap, "
                           "orange = stretch, grey = redesign first")]
for k, v in steps:
    cell(g, r, 2, k, bold=True, fill=LAV); cell(g, r, 3, v, align=WT); g.row_dimensions[r].height = 40; r += 1
r += 1; band(g, r, "DECISION LOGIC (SUMMARY)", 2, c0=2); r += 1
logic = [("Solution family", "From the nature of the task (library default, or override). 10 families, ISO/IEC 22989 / OECD-based, each with a complexity level 1-5 and a minimum data requirement."),
         ("Data gate", "Data available for the process >= the family's minimum requirement -> 'Data ready', otherwise 'Data gap'."),
         ("Organisational fit", "Family level <= the archetype's ceiling -> 'Within current position', otherwise 'Beyond current position'."),
         ("Verdict (in order)", "Process not standardised -> Redesign first;  family needs no AI -> Automate;  data gap -> AI later;  beyond position -> Stretch;  otherwise -> AI fit now."),
         ("Value (1-5)", "30% volume + 25% effort (hours / month) + 25% strategic importance + 20% customer / quality impact."),
         ("Feasibility (1-5)", "30% data fit + 25% process standardisation + 25% solution simplicity + 20% organisational fit."),
         ("Risk tier", "AIAF-style inherent risk: affects individuals' rights / access (+2), personal or confidential data (+1), autonomy (+0 to +2), impact of errors (+1 if 4, +2 if 5): "
                       "4+ High, 2-3 Medium, 0-1 Low; never below the library's minimum tier (EU AI Act Annex III / sector rules)."),
         ("Priority", "Value x Feasibility, reduced 15% for High risk and 20% where the process must be redesigned first. Ranked 1..n."),
         ("Horizon", "H1: quick wins / fill-ins that fit now or need only automation.  H2: strategic bets, stretch items, redesigned processes - and any High-risk or agentic (level 5) "
                     "use case, which waits for its impact assessment and controls.  H3: items blocked by data gaps.  Backlog: deprioritised.")]
for k, v in logic:
    cell(g, r, 2, k, bold=True, fill=LAV); cell(g, r, 3, v, align=WT); g.row_dimensions[r].height = 32; r += 1
portrait(g)

# ================================================================== 1. AI Position
P = wb.create_sheet("1. AI Position", 1)
title(P, "1. Organisational AI Position", "Does the organisation have the data and the scale / capacity for AI - or is automation enough for now? "
      "Select the sector and one answer per question (yellow). This screen does not measure governance maturity.", 6)
for c, w in zip("ABCDEF", [6, 46, 44, 34, 8, 4]): P.column_dimensions[c].width = w
prof = ["Organisation", "Sector (filters the use-case library)", "Scope of discovery", "Assessment date", "Facilitated by"]
for i, k in enumerate(prof):
    cell(P, 4 + i, 1, "", fill=LAV); cell(P, 4 + i, 2, k, bold=True, fill=LAV)
    merged(P, 4 + i, 3, 4, "", fill=INPUT)
dv_list(P, f"={SEC}", "C5")
SEL = "'1. AI Position'!$C$5"
if EXAMPLE:
    for i, k in enumerate(["Organisation", "Sector", "Scope", "Assessment date", "Facilitated by"]):
        P.cell(4 + i, 3).value = EX["profile"][k]
r = 10
for j, h in enumerate(["#", "Question", "Answer (select)", "Why it matters", "Points"], 1): cell(P, r, j, h, bold=True, fill=GT, color="FFFFFF", align=CC)
P.row_dimensions[r].height = 24
POS_ROW = {}
r += 1
for axis, label in [("Data", "A. DATA POSITION - is the data there for AI?"), ("Scale", "B. SCALE & CAPACITY - is there enough volume and capacity to justify and run AI?")]:
    band(P, r, label, 5, fill=GT2); r += 1
    for qid, ax, q, gd, opts in POSITION_Q:
        if ax != axis: continue
        cell(P, r, 1, qid, bold=True, fill=LAV, align=CC); cell(P, r, 2, q, align=WT)
        a = cell(P, r, 3, "", fill=INPUT, align=WT)
        if EXAMPLE: a.value = opts[EX["position"][qid]]
        dv_list(P, f"={POS_RANGE[qid]}", f"C{r}")
        cell(P, r, 4, gd, align=WT, size=8.5, color="595959")
        cell(P, r, 5, f'=IFERROR(MATCH(C{r},{POS_RANGE[qid]},0),"")', align=CC, fill=LAV)
        P.row_dimensions[r].height = 42; POS_ROW[qid] = r; r += 1
d_rows = [POS_ROW[q[0]] for q in POSITION_Q if q[1] == "Data"]; s_rows = [POS_ROW[q[0]] for q in POSITION_Q if q[1] == "Scale"]
r += 1; band(P, r, "RESULT - AI POSITION", 5); r += 1
RES = r
d_sum = "+".join(f"N(E{x})" for x in d_rows); s_sum = "+".join(f"N(E{x})" for x in s_rows)
d_cnt = ",".join(f"E{x}" for x in d_rows); s_cnt = ",".join(f"E{x}" for x in s_rows)
rows = [("Answered", f'=COUNT({d_cnt})+COUNT({s_cnt})&" of 12"', None),
        ("Data Position score", f'=IF(COUNT({d_cnt})=6,({d_sum}-6)/18,"")', "0%"),
        ("Scale & Capacity score", f'=IF(COUNT({s_cnt})=6,({s_sum}-6)/18,"")', "0%"),
        ("Archetype", f'=IF(OR(C{RES + 1}="",C{RES + 2}=""),"",IF(C{RES + 1}<0.5,IF(C{RES + 2}<0.5,1,2),IF(C{RES + 2}<0.5,3,4)))', None),
        ("AI position", f'=IF(C{RES + 3}="","Complete all 12 questions",INDEX({ARC("AH")},C{RES + 3}))', None),
        ("Solution ceiling today", f'=IF(C{RES + 3}="","",INDEX({ARC("AI")},C{RES + 3})&" - "&INDEX(Lists!$AM$2:$AM$6,INDEX({ARC("AI")},C{RES + 3})))', None),
        ("What it means", f'=IF(C{RES + 3}="","",INDEX({ARC("AJ")},C{RES + 3}))', None),
        ("Detail", f'=IF(C{RES + 3}="","",INDEX({ARC("AK")},C{RES + 3}))', None),
        ("Foundation flag", f'=IF(E{POS_ROW["D1"]}=1,"Core processes are largely manual / paper-based: digitise them first - AI and automation need digital inputs.",'
                            f'IF(AND(N(E{POS_ROW["D2"]})>0,N(E{POS_ROW["D2"]})<=2,N(E{POS_ROW["D6"]})<=1),"Little digital history and no recorded outcomes: predictive / ML use cases will fail the data gate until data is captured.",'
                            f'"No foundation red flags from the position screen."))', None)]
for i, (k, f, fmt) in enumerate(rows):
    cell(P, RES + i, 2, k, bold=True, fill=LAV)
    x = merged(P, RES + i, 3, 5, f, fill=LAV2 if i in (4, 5) else None, bold=i in (1, 2, 4, 5), size=11 if i == 4 else 9.5, align=WT)
    if fmt: x.number_format = fmt
    P.row_dimensions[RES + i].height = 20 if i < 6 else 44
P.row_dimensions[RES + 7].height = 58
POS_SCORE_D, POS_SCORE_S, POS_ARCH = f"'1. AI Position'!$C${RES + 1}", f"'1. AI Position'!$C${RES + 2}", f"'1. AI Position'!$C${RES + 3}"
POS_NAME = f"'1. AI Position'!$C${RES + 4}"
CEIL = f"IF({POS_ARCH}=\"\",5,INDEX({ARC('AI')},{POS_ARCH}))"
POS_FLAG = f"'1. AI Position'!$C${RES + 8}"
r = RES + 10; band(P, r, "POSITION MATRIX - the highlighted box is where the organisation sits", 5); r += 1
MX = r
cell(P, MX, 2, "Data Position HIGH  (>= 50%)", bold=True, align=CC, fill=GREY)
cell(P, MX + 1, 2, "Data Position LOW  (< 50%)", bold=True, align=CC, fill=GREY)
cell(P, MX + 2, 3, "Scale & Capacity LOW  (< 50%)", bold=True, align=CC, fill=GREY)
cell(P, MX + 2, 4, "Scale & Capacity HIGH  (>= 50%)", bold=True, align=CC, fill=GREY)
for (rr, cc), ai in {(MX, 3): 2, (MX, 4): 3, (MX + 1, 3): 0, (MX + 1, 4): 1}.items():
    code, name, ceil, head, desc = ARCHETYPES[ai]
    cell(P, rr, cc, f"{code}. {name}\n{head}", align=CC, size=9.5)
    P.conditional_formatting.add(f"{get_column_letter(cc)}{rr}", FormulaRule(formula=[f"{POS_ARCH}={ai + 1}"], fill=PatternFill("solid", fgColor="A9D08E"), font=Font(bold=True)))
P.row_dimensions[MX].height = 66; P.row_dimensions[MX + 1].height = 66; P.row_dimensions[MX + 2].height = 22
P.print_area = f"A1:E{MX + 2}"
portrait(P, fit_h=1)

# ================================================================== 2. Use-Case Library
LS = "'2. Use-Case Library'!"
LB0, LB1 = 5, 5 + NL - 1
LIBR = lambda col: f"{LS}${col}${LB0}:${col}${LB1}"
Lib = wb.create_sheet("2. Use-Case Library", 2)
title(Lib, f"2. GT AI Use-Case Library ({NL} entries)", "Cross-industry (APQC functions) and sector-specific AI and automation use cases with GT recommendations, data needs, KPIs, "
      "regulatory notes and a minimum risk tier. Select entries in the AI Fit Diagnostic; the dropdown there shows cross-industry + the sector chosen on sheet 1.", 12)
for c, w in zip("ABCDEFGHIJKL", [9, 16, 20, 30, 28, 30, 22, 48, 26, 26, 40, 9]): Lib.column_dimensions[c].width = w
heads = ["ID", "Sector", "Function (APQC)", "Process / use case", "Typical pain point", "Nature of the task", "Solution family", "GT recommendation",
         "Data needed", "KPIs", "Regulatory & risk notes", "Min. risk"]
for j, h in enumerate(heads, 1): cell(Lib, 4, j, h, bold=True, fill=GT, color="FFFFFF", align=CC, size=9)
for k, (lid, sec, fn, nat, proc, pain, rec, data, kpi, notes, floor) in enumerate(LIBRARY):
    r = LB0 + k
    vals = [lid, sec, fn, proc, pain, FAMILIES[nat][0], FAMILIES[nat][1], rec, data, kpi, notes, floor]
    for j, v in enumerate(vals, 1):
        cell(Lib, r, j, v, align=WT, size=8.5, bold=j in (1, 4), fill=LAV if sec == "Cross-industry" and j in (1, 2) else None)
    Lib.cell(r, 13, f"{lid} | {proc}")                                           # M label (dropdown value)
    Lib.cell(r, 14, f'=OR({SEL}="",{SEL}="Other",B{r}="Cross-industry",B{r}={SEL})')   # N include
    Lib.cell(r, 15, f'=IF(N{r},COUNTIF($N${LB0}:N{r},TRUE),"")')                 # O running count
    Lib.row_dimensions[r].height = 46
risk_cf(Lib, f"L{LB0}:L{LB1}")
for c in "MNO": Lib.column_dimensions[c].hidden = True
Lib.freeze_panes = "E5"; Lib.print_title_rows = "4:4"; Lib.auto_filter.ref = f"A4:L{LB1}"
Lib.print_area = f"A1:L{LB1}"
Lib.page_setup.paperSize = Lib.PAPERSIZE_A3; portrait(Lib, landscape=True)
# filtered dropdown list (Lists!BA)
for k in range(1, NL + 1):
    L.cell(k + 1, 53, f'=IFERROR(INDEX({LIBR("M")},MATCH({k},{LIBR("O")},0)),"")')
DDL = f'OFFSET(Lists!$BA$2,0,0,MAX(1,COUNTIF(Lists!$BA$2:$BA${NL + 1},"?*")),1)'

# ================================================================== 3. AI Fit Diagnostic
D = wb.create_sheet("3. AI Fit Diagnostic", 3)
NROW = 50; R0 = 8; R1 = R0 + NROW - 1
title(D, "3. AI Fit Diagnostic - process by process", "Select a use case from the library (column B; type a custom process if it is not in the library, then select its task nature "
      "and function in D-E) and answer the yellow inputs F-N. Effort = staff hours spent on the process per month (all people involved).", 36)
widths = [4, 40, 24, 26, 18, 16, 9, 18, 18, 8, 8, 9, 9, 18,
          8, 18, 22, 24, 6, 6, 10, 13, 22, 7, 7, 8, 12, 8, 6, 17, 10, 48, 36, 5, 5, 5]
for i, w in enumerate(widths, 1): D.column_dimensions[get_column_letter(i)].width = w
merged(D, 3, 2, 2, "Organisation's AI position (from sheet 1):", bold=True, fill=LAV)
merged(D, 3, 3, 5, f"={POS_NAME}", bold=True, fill=LAV2)
merged(D, 3, 6, 8, "Solution ceiling level:", bold=True, fill=LAV)
cell(D, 3, 9, f"={CEIL}", bold=True, fill=LAV2, align=CC)
merged(D, 4, 2, 2, "Quadrant threshold (value and feasibility, 1-5):", bold=True, fill=LAV)
cell(D, 4, 3, 3.25, fill=INPUT, align=CC)
merged(D, 4, 4, 5, '="Library entries offered: "&COUNTIF(Lists!$BA$2:$BA$' + str(NL + 1) + ',"?*")', fill=LAV, italic=True)
THR = "$C$4"
merged(D, 6, 1, 14, "INPUTS (yellow)", bold=True, fill=GT, color="FFFFFF", align=CC)
merged(D, 6, 15, 33, "AUTO-CALCULATED (recommendation and notes come from the GT AI Use-Case Library)", bold=True, fill=DARK, color="FFFFFF", align=CC)
heads = ["#", "Use case (select from library, or type a custom process)", "Client-specific note (optional)", "Task nature - custom entries / override", "Function - custom entries only",
         "Volume", "Effort (hrs / month)", "Process standardised?", "Data available for this process", "Strategic importance (1-5)", "Customer / quality impact (1-5)",
         "Affects individuals' rights or access?", "Personal or confidential data?", "Intended autonomy",
         "Library ID", "Function", "Nature of the task", "Solution family", "Level", "AI?", "Data gate", "Organisational fit", "VERDICT", "Value", "Feasib-ility",
         "Risk tier", "Quadrant", "Priority score", "Rank", "Horizon", "Est. hours released / yr", "GT recommendation (library)", "Regulatory & risk notes (library)", "h", "key", "name"]
for j, h in enumerate(heads, 1):
    cell(D, 7, j, h, bold=True, fill=GT if j <= 14 else DARK, color="FFFFFF", align=CC, size=8.5)
D.row_dimensions[7].height = 52
LBK = lambda col, r: f"INDEX({LIBR(col)},MATCH($O{r},{LIBR('A')},0))"
for r in range(R0, R1 + 1):
    cell(D, r, 1, r - R0 + 1, align=CC, fill=LAV)
    for c in range(2, 15):
        cell(D, r, c, None, fill=INPUT, align=WT, size=8.5)
    full = f'OR(COUNTA(F{r}:N{r})<9,Q{r}="")'
    score = f'IF(L{r}="Yes",2,0)+IF(M{r}="Yes",1,0)+INDEX({AU_V},MATCH(N{r},{AU_O},0))+IF(K{r}=5,2,IF(K{r}>=4,1,0))'
    calc_n = f'IF(T{r}="No",IF(L{r}="Yes",2,1),IF({score}>=4,3,IF({score}>=2,2,1)))'
    floor_n = f'IF($O{r}="Custom",1,MATCH({LBK("L", r)},{{"Low","Medium","High"}},0))'
    f = {
        15: f'=IF(B{r}="","",IFERROR(INDEX({LIBR("A")},MATCH(B{r},{LIBR("M")},0)),"Custom"))',
        16: f'=IF(B{r}="","",IF(O{r}="Custom",IF(E{r}="","(select function in E)",E{r}),{LBK("C", r)}))',
        17: f'=IF(B{r}="","",IF(D{r}<>"",D{r},IF(O{r}="Custom","",{LBK("F", r)})))',
        18: f'=IF(Q{r}="","",INDEX({LB},MATCH(Q{r},{LA},0)))',
        19: f'=IF(R{r}="","",INDEX({fam("C")},MATCH(R{r},{LB},0)))',
        20: f'=IF(R{r}="","",INDEX({fam("D")},MATCH(R{r},{LB},0)))',
        21: f'=IF(OR(R{r}="",I{r}=""),"",IF(INDEX({DA_V},MATCH(I{r},{DA_O},0))>=INDEX({fam("E")},MATCH(R{r},{LB},0)),"Data ready","Data gap"))',
        22: f'=IF(R{r}="","",IF(S{r}<={CEIL},"Within current position","Beyond current position"))',
        23: f'=IF(B{r}="","",IF({full},"Complete inputs",IF(H{r}="{STANDARD[2][0]}","{VERDICTS[4][0]}",IF(T{r}="No","{VERDICTS[1][0]}",'
            f'IF(U{r}="Data gap","{VERDICTS[2][0]}",IF(V{r}="Beyond current position","{VERDICTS[3][0]}","{VERDICTS[0][0]}"))))))',
        24: f'=IF(OR(W{r}="",W{r}="Complete inputs"),"",ROUND(0.3*INDEX({VO_V},MATCH(F{r},{VO_O},0))+0.25*IF(G{r}<20,1,IF(G{r}<80,2,IF(G{r}<160,3,IF(G{r}<480,4,5))))'
            f'+0.25*J{r}+0.2*K{r},2))',
        25: f'=IF(X{r}="","",ROUND(0.3*IF(U{r}="Data ready",5,IF(INDEX({DA_V},MATCH(I{r},{DA_O},0))=INDEX({fam("E")},MATCH(R{r},{LB},0))-1,3,1))'
            f'+0.25*INDEX({ST_V},MATCH(H{r},{ST_O},0))+0.25*(6-S{r})+0.2*IF(V{r}="Within current position",5,2),2))',
        26: f'=IF(X{r}="","",CHOOSE(MAX({calc_n},{floor_n}),"Low","Medium","High"))',
        27: f'=IF(X{r}="","",IF(AND(X{r}>={THR},Y{r}>={THR}),"Quick win",IF(X{r}>={THR},"Strategic bet",IF(Y{r}>={THR},"Fill-in","Deprioritise"))))',
        28: f'=IF(X{r}="","",ROUND(X{r}*Y{r}*IF(Z{r}="High",0.85,1)*IF(W{r}="{VERDICTS[4][0]}",0.8,1),2))',
        29: f'=IF(AB{r}="","",RANK(AB{r},$AB${R0}:$AB${R1})+COUNTIF($AB${R0 - 1}:AB{r - 1},AB{r}))',
        30: f'=IF(AH{r}="","",INDEX(Lists!$AV$2:$AV$5,AH{r}))',
        31: f'=IF(AB{r}="","",ROUND(G{r}*12*INDEX({fam("F")},MATCH(R{r},{LB},0)),0))',
        32: f'=IF(OR(B{r}="",R{r}=""),"",IF(O{r}="Custom",INDEX({fam("I")},MATCH(R{r},{LB},0)),{LBK("H", r)}))',
        33: f'=IF(OR(B{r}="",O{r}="Custom"),"",{LBK("K", r)})',
        34: f'=IF(AB{r}="","",IF(AA{r}="Deprioritise",4,IF(W{r}="{VERDICTS[2][0]}",3,IF(OR(W{r}="{VERDICTS[4][0]}",W{r}="{VERDICTS[3][0]}",AA{r}="Strategic bet",Z{r}="High",S{r}=5),2,1))))',
        35: f'=IF(AC{r}="","",AH{r}*100+AC{r})',
        36: f'=IF(B{r}="","",IF(O{r}="Custom",B{r},{LBK("D", r)}))',
    }
    for c, fx in f.items():
        x = cell(D, r, c, fx, align=CC if c not in (16, 17, 18, 22, 23, 30, 32, 33) else WC, size=8 if c in (32, 33) else 8.5, bold=c in (23, 29))
        if c in (24, 25, 28): x.number_format = "0.00"
        if c == 31: x.number_format = "#,##0"
    D.row_dimensions[r].height = 52
rng = lambda col: f"{col}{R0}:{col}{R1}"
dv_list(D, f"={DDL}", rng("B"), "Select a library use case (cross-industry + your sector), or type a custom process.", strict=False)
dv_list(D, f"={LA}", rng("D"), "Only for custom processes, or to override the library's task nature.")
dv_list(D, f"={FN}", rng("E"), "Only for custom processes.")
dv_list(D, f"={VO_O}", rng("F")); dv_list(D, f"={ST_O}", rng("H")); dv_list(D, f"={DA_O}", rng("I"))
dv_list(D, '"1,2,3,4,5"', f"J{R0}:K{R1}"); dv_list(D, f"={YN}", f"L{R0}:M{R1}"); dv_list(D, f"={AU_O}", rng("N"))
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100000", allow_blank=True); D.add_data_validation(dv); dv.add(rng("G"))
verdict_cf(D, rng("W")); risk_cf(D, rng("Z"))
D.conditional_formatting.add(rng("U"), CellIsRule(operator="equal", formula=['"Data gap"'], font=Font(color="C00000", bold=True)))
D.conditional_formatting.add(rng("V"), CellIsRule(operator="equal", formula=['"Beyond current position"'], font=Font(color="C55A11", bold=True)))
D.conditional_formatting.add(rng("AA"), CellIsRule(operator="equal", formula=['"Quick win"'], fill=PatternFill("solid", fgColor="A9D08E")))
D.conditional_formatting.add(rng("O"), CellIsRule(operator="equal", formula=['"Custom"'], font=Font(color="C55A11", bold=True)))
for c in ("AH", "AI", "AJ"): D.column_dimensions[c].hidden = True
D.freeze_panes = "C8"; D.print_title_rows = "6:7"
D.print_area = f"A1:AG{R1}"
D.page_setup.paperSize = D.PAPERSIZE_A3; portrait(D, landscape=True)
if EXAMPLE:
    label = {x[4]: f"{x[0]} | {x[4]}" for x in LIBRARY}
    for k, (proc, vi, hrs, si, di, sti, ci, rights, pers, ai) in enumerate(EX["processes"]):
        r = R0 + k
        vals = {2: label[proc], 6: VOLUME[vi][0], 7: hrs, 8: STANDARD[si][0], 9: DATA_AVAIL[di][0], 10: sti, 11: ci, 12: rights, 13: pers, 14: AUTONOMY[ai][0]}
        for c, v in vals.items(): D.cell(r, c).value = v
DG = "'3. AI Fit Diagnostic'!"
DR = lambda col: f"{DG}${col}${R0}:${col}${R1}"

# ================================================================== 7. Sector Profiles (built before the report, which looks it up)
SP = wb.create_sheet("7. Sector Profiles")
title(SP, "7. Sector Profiles", "Value at stake, automation & AI potential, where value sits and the regulators / rules that shape AI use - by sector. "
      "Library entries are counted live from the GT AI Use-Case Library.", 7)
for c, w in zip("ABCDEFG", [26, 22, 13, 44, 52, 10, 10]): SP.column_dimensions[c].width = w
for j, h in enumerate(["Sector", "Value at stake by 2030 (GDP uplift)", "Automation & AI potential", "Where the value is (typical value pools)", "Regulators & key rules",
                       "Library entries", "High-risk entries"], 1):
    cell(SP, 4, j, h, bold=True, fill=GT, color="FFFFFF", align=CC, size=9)
SP0 = 5
for k, (sec, gdp, pot, pools, regs) in enumerate(SECTOR_PROFILES):
    r = SP0 + k
    for j, v in enumerate([sec, gdp, pot, pools, regs], 1):
        cell(SP, r, j, v, align=WT, size=9, bold=j == 1, fill=LAV if j == 1 else None)
    cell(SP, r, 6, f'=COUNTIF({LIBR("B")},A{r})', align=CC)
    cell(SP, r, 7, f'=COUNTIFS({LIBR("B")},A{r},{LIBR("L")},"High")', align=CC)
    SP.row_dimensions[r].height = 58
SP1 = SP0 + len(SECTOR_PROFILES) - 1
r = SP1 + 1
cell(SP, r, 1, "Cross-industry", bold=True, fill=LAV); merged(SP, r, 2, 5, "Finance, HR, customer service, marketing, procurement, risk & compliance, IT, facilities, management and knowledge processes common to all sectors.", align=WT, size=9)
cell(SP, r, 6, f'=COUNTIF({LIBR("B")},"Cross-industry")', align=CC); cell(SP, r, 7, f'=COUNTIFS({LIBR("B")},"Cross-industry",{LIBR("L")},"High")', align=CC)
SP.row_dimensions[r].height = 30; r += 1
cell(SP, r, 1, "Total", bold=True, fill=LAV2); merged(SP, r, 2, 5, "", fill=LAV2)
cell(SP, r, 6, f"=SUM(F{SP0}:F{r - 1})", bold=True, align=CC, fill=LAV2); cell(SP, r, 7, f"=SUM(G{SP0}:G{r - 1})", bold=True, align=CC, fill=LAV2); r += 2
merged(SP, r, 1, 7, "Value at stake: QuantumBlack & PwC analysis - predictions for 2030 (as presented in GT Bahrain 'Introduction to Artificial Intelligence', 2026); "
                    "insurance and capital markets share the financial-services figure; hospitality shares retail & consumer. Sectors not sized separately in that analysis are described qualitatively. "
                    "Potential is GT's qualitative assessment of automation and AI opportunity density in the library.", italic=True, size=8, color="595959", align=WT, border=False)
SP.row_dimensions[r].height = 40
SP.print_area = f"A1:G{r}"; portrait(SP, landscape=True)

# ================================================================== 4. Opportunity Report
M = wb.create_sheet("4. Opportunity Report", 4)
title(M, "4. AI Opportunity Report", "Where AI fits - and where automation is enough - by function, with the top-ranked opportunities and GT recommendations from the use-case library.", 11)
for c, w in zip("ABCDEFGHIJK", [5, 30, 11, 11, 11, 11, 11, 11, 13, 13, 22]): M.column_dimensions[c].width = w
r = 4
band(M, r, "A. ORGANISATIONAL AI POSITION", 11); r += 1
for k, f in [("Organisation", "='1. AI Position'!C4"), ("Sector", f"={SEL}&\"\""), ("AI position (archetype)", f"={POS_NAME}"),
             ("Data Position / Scale & Capacity", f'=IF({POS_SCORE_D}="","",TEXT({POS_SCORE_D},"0%")&"  /  "&TEXT({POS_SCORE_S},"0%"))'),
             ("What it means", f"='1. AI Position'!C{RES + 6}"), ("Foundation flag", f"={POS_FLAG}")]:
    merged(M, r, 1, 2, k, bold=True, fill=LAV)
    merged(M, r, 3, 11, f, align=WT, bold=k.startswith("AI position"), fill=LAV2 if k.startswith("AI position") else None)
    M.row_dimensions[r].height = 30 if k in ("What it means", "Foundation flag") else 18; r += 1
r += 1; band(M, r, "B. VERDICT SUMMARY - what kind of solution fits the processes assessed", 11); r += 1
for j, h in enumerate(["", "Verdict", "Processes", "% of total", "Hours released / yr", "", "Meaning"], 1):
    if h: cell(M, r, j, h, bold=True, fill=GT, color="FFFFFF", align=CC)
M.merge_cells(start_row=r, start_column=7, end_row=r, end_column=11)
r += 1; VS = r
for v, mean, nxt in VERDICTS:
    cell(M, r, 2, v, bold=True, fill=VFILL[v])
    cell(M, r, 3, f'=COUNTIF({DR("W")},B{r})', align=CC)
    cell(M, r, 4, f'=IF(SUM($C${VS}:$C${VS + 4})=0,"",C{r}/SUM($C${VS}:$C${VS + 4}))', align=CC, fmt="0%")
    cell(M, r, 5, f'=SUMIF({DR("W")},B{r},{DR("AE")})', align=CC, fmt="#,##0")
    merged(M, r, 6, 11, mean, align=WT, size=8.5)
    M.row_dimensions[r].height = 30; r += 1
cell(M, r, 2, "Total assessed", bold=True, fill=LAV); cell(M, r, 3, f"=SUM(C{VS}:C{VS + 4})", bold=True, align=CC, fill=LAV)
cell(M, r, 4, "", fill=LAV); cell(M, r, 5, f"=SUM(E{VS}:E{VS + 4})", bold=True, align=CC, fill=LAV, fmt="#,##0")
r += 1
merged(M, r, 2, 11, f'=IF(C{r - 1}=0,"",TEXT((C{VS}+C{VS + 3})/C{r - 1},"0%")&" of assessed processes are AI opportunities (now or with a partner); "&TEXT(C{VS + 1}/C{r - 1},"0%")&'
                    f'" are better served by automation without AI; "&TEXT((C{VS + 2}+C{VS + 4})/C{r - 1},"0%")&" need data or process foundations first.")',
       italic=True, align=WT, color=GT, bold=True)
M.row_dimensions[r].height = 30; r += 2
band(M, r, "C. WHERE AI FITS BY FUNCTION (number of processes)", 11); r += 1
fh = ["", "Function", "AI fit now", "Stretch", "AI later (data)", "Automate (no AI)", "Redesign first", "Total", "Hours / yr", "High-risk items", "Headline"]
for j, h in enumerate(fh, 1):
    if h: cell(M, r, j, h, bold=True, fill=GT, color="FFFFFF", align=CC, size=8.5)
M.row_dimensions[r].height = 30; r += 1; FS = r
vorder = [VERDICTS[0][0], VERDICTS[3][0], VERDICTS[2][0], VERDICTS[1][0], VERDICTS[4][0]]
for fn in FUNCTIONS:
    cell(M, r, 2, fn, bold=True, fill=LAV, size=8.5)
    for j, v in enumerate(vorder):
        cell(M, r, 3 + j, f'=COUNTIFS({DR("P")},$B{r},{DR("W")},"{v}")', align=CC)
    cell(M, r, 8, f"=SUM(C{r}:G{r})", align=CC, bold=True)
    cell(M, r, 9, f'=SUMIF({DR("P")},$B{r},{DR("AE")})', align=CC, fmt="#,##0")
    cell(M, r, 10, f'=COUNTIFS({DR("P")},$B{r},{DR("Z")},"High")', align=CC)
    cell(M, r, 11, f'=IF(H{r}=0,"",IF(C{r}+D{r}>=MAX(E{r},F{r},G{r}),"AI opportunity",IF(F{r}>=MAX(E{r},G{r}),"Automation is enough",IF(E{r}>=G{r},"Fix data first","Redesign first"))))',
         align=CC, size=8.5, bold=True)
    r += 1
FE = r - 1
for j, col in enumerate("CDEFGHIJ"):
    cell(M, r, 3 + j, f"=SUM({col}{FS}:{col}{FE})", bold=True, fill=LAV, align=CC, fmt="#,##0" if col == "I" else None)
cell(M, r, 2, "Total", bold=True, fill=LAV); cell(M, r, 11, "", fill=LAV)
for col, v in zip("CDEFG", vorder):
    M.conditional_formatting.add(f"{col}{FS}:{col}{FE}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=VFILL[v])))
M.conditional_formatting.add(f"K{FS}:K{FE}", CellIsRule(operator="equal", formula=['"AI opportunity"'], fill=PatternFill("solid", fgColor="A9D08E")))
M.conditional_formatting.add(f"K{FS}:K{FE}", CellIsRule(operator="equal", formula=['"Automation is enough"'], fill=PatternFill("solid", fgColor="9BC2E6")))
r += 2
ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100; ch.title = "Processes by function and verdict"
ch.y_axis.title = None; ch.x_axis.title = None
ch.add_data(Reference(M, min_col=3, max_col=7, min_row=FS - 1, max_row=FE), titles_from_data=True); ch.set_categories(Reference(M, min_col=2, min_row=FS, max_row=FE))
for s_, v in zip(ch.series, vorder):
    s_.graphicalProperties.solidFill = VFILL[v]; s_.graphicalProperties.line.solidFill = VFILL[v]
ch.height = 9.5; ch.width = 25; ch.legend.position = "b"
ch.x_axis.scaling.orientation = "maxMin"; ch.x_axis.delete = False; ch.y_axis.delete = False; ch.y_axis.majorGridlines = None
M.add_chart(ch, f"B{r}")
r += 20
M.row_breaks.append(Break(id=r - 1))
band(M, r, "D. TOP 10 OPPORTUNITIES (risk-adjusted priority) - with GT recommendations", 11); r += 1
for (c1, c2), h in [((1, 1), "Rank"), ((2, 2), "Process / use case"), ((3, 4), "Solution family"), ((5, 6), "Verdict"), ((7, 7), "Risk"), ((8, 8), "Horizon"), ((9, 11), "GT recommendation")]:
    merged(M, r, c1, c2, h, bold=True, fill=GT, color="FFFFFF", align=CC, size=8.5)
M.row_dimensions[r].height = 24; r += 1; TS = r
for k in range(1, 11):
    mrow = f"MATCH({k},{DR('AC')},0)"
    g_ = lambda col: f'=IFERROR(INDEX({DR(col)},{mrow}),"")'
    cell(M, r, 1, k, bold=True, align=CC, fill=LAV)
    cell(M, r, 2, g_("AJ"), align=WT, bold=True, size=8.5)
    merged(M, r, 3, 4, g_("R"), align=WT, size=8); merged(M, r, 5, 6, g_("W"), align=WT, size=8, bold=True)
    cell(M, r, 7, g_("Z"), align=CC, size=8.5); cell(M, r, 8, g_("AD"), align=WT, size=8)
    merged(M, r, 9, 11, g_("AF"), align=WT, size=7.5)
    M.row_dimensions[r].height = 46; r += 1
verdict_cf(M, f"E{TS}:F{r - 1}"); risk_cf(M, f"G{TS}:G{r - 1}")
r += 1
band(M, r, "E. SECTOR LENS", 11); r += 1
spm = f"MATCH({SEL},'7. Sector Profiles'!$A${SP0}:$A${SP1},0)"
for k, col in [("Value at stake by 2030", "B"), ("Automation & AI potential", "C"), ("Where the value is", "D"), ("Regulators & key rules", "E"), ("Library entries for this sector", "F")]:
    merged(M, r, 1, 2, k, bold=True, fill=LAV)
    merged(M, r, 3, 11, f"=IFERROR(INDEX('7. Sector Profiles'!${col}${SP0}:${col}${SP1},{spm}),\"Select a sector on sheet 1\")", align=WT, size=9)
    M.row_dimensions[r].height = 30 if col in ("B", "D", "E") else 18; r += 1
r += 1
merged(M, r, 1, 11, "Value and feasibility scored 1-5 (sheet 3). Hours released are indicative (effort x typical reduction for the solution family) and must be validated in "
                    "the business case. Risk tiers are inherent (before controls) and never below the library minimum. Recommendations come from the GT AI Use-Case Library.",
       italic=True, size=8, color="595959", align=WT, border=False)
M.row_dimensions[r].height = 28
M.print_area = f"A1:K{r}"; portrait(M)

# ================================================================== 5. Roadmap
RM = wb.create_sheet("5. Roadmap", 5)
title(RM, "5. Opportunity Roadmap", "Opportunities placed automatically by horizon (ordered by priority rank), with the foundations they depend on.", 7)
for c, w in zip("ABCDEFG", [6, 40, 24, 28, 24, 10, 10]): RM.column_dimensions[c].width = w
r = 4
KEY = DR("AI")
for h, hname in enumerate(HORIZONS, 1):
    band(RM, r, hname, 7, fill=[GT, GT2, TEAL, "7F7F7F"][h - 1]); r += 1
    for j, hd in enumerate(["Rank", "Process / use case", "Function", "Solution family", "Verdict", "Risk tier", "Hours / yr"], 1):
        cell(RM, r, j, hd, bold=True, fill=LAV, align=CC, size=8.5)
    r += 1; s0 = r
    for k in range(1, [20, 12, 10, 8][h - 1] + 1):
        key = f'IFERROR(SMALL({KEY},COUNTIF({KEY},"<"&{h * 100})+{k}),"")'
        RM.cell(r, 8, f'=IF({key}="","",IF({key}<{(h + 1) * 100},{key},""))')
        m = f"MATCH($H{r},{KEY},0)"
        for j, col in enumerate(["AC", "AJ", "P", "R", "W", "Z", "AE"], 1):
            cell(RM, r, j, f'=IF($H{r}="","",INDEX({DR(col)},{m}))', align=CC if j in (1, 6, 7) else WT, size=8.5, bold=j in (1, 2), fmt="#,##0" if j == 7 else None)
        RM.row_dimensions[r].height = 26; r += 1
    verdict_cf(RM, f"E{s0}:E{r - 1}"); risk_cf(RM, f"F{s0}:F{r - 1}")
    r += 1
RM.column_dimensions["H"].hidden = True
band(RM, r, "ENABLING FOUNDATIONS (generated from the diagnostic and position screen)", 7, fill=DARK); r += 1
T = DR("W"); W = DR("Z")
enablers = [
    ("Governance (always)", f'="Before H1 go-live: AI acceptable-use policy, AI register and approved-tools list; tier-based use-case risk assessment (GT AI Risk Assessment Framework). '
                            f'"&COUNTIF({W},"High")&" High-risk and "&COUNTIF({W},"Medium")&" Medium-risk items need documented assessments."'),
    ("Data foundations", f'=IF(COUNTIF({T},"{VERDICTS[2][0]}")=0,"No processes are blocked by data gaps.",COUNTIF({T},"{VERDICTS[2][0]}")&" process(es) are blocked by data gaps: assign data owners, '
                         f'digitise inputs and start recording outcomes now so they can move to AI in H3.")'),
    ("Process redesign", f'=IF(COUNTIF({T},"{VERDICTS[4][0]}")=0,"No processes need redesign before automation.",COUNTIF({T},"{VERDICTS[4][0]}")&" process(es) vary by person or team: '
                         f'standardise and document them first, then re-diagnose.")'),
    ("Automation platform", f'=IF(COUNTIF({T},"{VERDICTS[1][0]}")=0,"","Automation is enough for "&COUNTIF({T},"{VERDICTS[1][0]}")&" process(es): select one workflow / RPA / BI platform and '
                            f'automate them as a programme rather than one-offs.")'),
    ("Capability & partners", f'=IF(COUNTIF({T},"{VERDICTS[3][0]}")=0,"All AI opportunities are within the current position.",COUNTIF({T},"{VERDICTS[3][0]}")&" opportunity(ies) exceed the '
                              f'current position: use vendor products or a delivery partner and build skills before scaling.")'),
    ("People", '="AI literacy for all users of AI tools and role-based training for process owners (EU AI Act Art. 4-aligned); change management for each roll-out."'),
    ("Readiness check", '="For governance and capability maturity, run the GT AI Readiness & Risk Assessment (AI Ready7-based) alongside this roadmap."'),
]
for k, f in enablers:
    cell(RM, r, 1, "", fill=LAV); cell(RM, r, 2, k, bold=True, fill=LAV)
    merged(RM, r, 3, 7, f, align=WT, size=8.5); RM.row_dimensions[r].height = 32; r += 1
RM.print_area = f"A1:G{r - 1}"; portrait(RM)

# ================================================================== 6. Use-Case Canvas
CV = wb.create_sheet("6. Use-Case Canvas", 6)
title(CV, "6. AI Use-Case Canvas", "Select an opportunity number from the AI Fit Diagnostic (yellow cell). Content comes from the diagnostic and the GT AI Use-Case Library; "
      "complete the yellow fields with the owner.", 4)
for c, w in zip("ABCD", [28, 46, 28, 46]): CV.column_dimensions[c].width = w
cell(CV, 4, 1, "Opportunity # (from sheet 3)", bold=True, fill=LAV); cell(CV, 4, 2, 1, fill=INPUT, bold=True, align=CC, size=12)
dv_list(CV, "=Lists!$AX$2:$AX$51", "B4")
cell(CV, 4, 3, "Priority rank", bold=True, fill=LAV)
m = f"MATCH($B$4,{DR('A')},0)"
gx = lambda col: f'=IFERROR(IF(INDEX({DR(col)},{m})="","",INDEX({DR(col)},{m})),"")'
lid = f"INDEX({DR('O')},{m})"
lb = lambda col: f'=IFERROR(IF({lid}="Custom","-",INDEX({LIBR(col)},MATCH({lid},{LIBR("A")},0))),"")'
famx = lambda col: f'=IFERROR(INDEX({fam(col)},MATCH(INDEX({DR("R")},{m}),{LB},0)),"")'
cell(CV, 4, 4, gx("AC"), bold=True, align=CC, size=12)
r = 6
band(CV, r, "1. THE OPPORTUNITY", 4); r += 1
for a, fa, b, fb in [("Process / use case", gx("AJ"), "Library ID / sector", f'=IFERROR({lid}&IF({lid}="Custom",""," - "&INDEX({LIBR("B")},MATCH({lid},{LIBR("A")},0))),"")'),
                     ("Typical pain point", lb("E"), "Function", gx("P")),
                     ("Client-specific note", gx("C"), "Nature of the task", gx("Q")),
                     ("Volume", gx("F"), "Effort today (hrs / month)", gx("G"))]:
    cell(CV, r, 1, a, bold=True, fill=LAV); cell(CV, r, 2, fa, align=WC); cell(CV, r, 3, b, bold=True, fill=LAV); cell(CV, r, 4, fb, align=WC)
    CV.row_dimensions[r].height = 34; r += 1
band(CV, r, "2. THE FIT AND GT RECOMMENDATION", 4); r += 1
VR = r
for a, fa, b, fb in [("VERDICT", gx("W"), "Solution family", gx("R")),
                     ("Data gate", gx("U"), "Organisational fit", gx("V")),
                     ("Value / Feasibility", f'=IFERROR(TEXT(INDEX({DR("X")},{m}),"0.00")&" / "&TEXT(INDEX({DR("Y")},{m}),"0.00")&"  ("&INDEX({DR("AA")},{m})&")","")', "Horizon", gx("AD"))]:
    cell(CV, r, 1, a, bold=True, fill=LAV); cell(CV, r, 2, fa, align=WC, bold=a == "VERDICT"); cell(CV, r, 3, b, bold=True, fill=LAV); cell(CV, r, 4, fb, align=WC)
    CV.row_dimensions[r].height = 30; r += 1
verdict_cf(CV, f"B{VR}")
for k, f, h in [("GT recommendation", gx("AF"), 58), ("Data needed", lb("I"), 30), ("Typical solutions", famx("G"), 30), ("Design notes", famx("I"), 44),
                ("What this verdict means", f'=IFERROR(INDEX(Lists!$AS$2:$AS$6,MATCH(B{VR},{VER},0)),"")', 30)]:
    cell(CV, r, 1, k, bold=True, fill=LAV); merged(CV, r, 2, 4, f, align=WC, bold=k == "GT recommendation"); CV.row_dimensions[r].height = h; r += 1
band(CV, r, "3. RISK, REGULATION & MINIMUM CONTROLS", 4); r += 1
cell(CV, r, 1, "Inherent risk tier", bold=True, fill=LAV); cell(CV, r, 2, gx("Z"), bold=True, align=CC); risk_cf(CV, f"B{r}")
cell(CV, r, 3, "Library minimum tier", bold=True, fill=LAV); cell(CV, r, 4, lb("L"), align=CC); RT = r; r += 1
cell(CV, r, 1, "Regulatory & risk notes", bold=True, fill=LAV); merged(CV, r, 2, 4, gx("AG"), align=WC); CV.row_dimensions[r].height = 44; r += 1
cell(CV, r, 1, "Minimum controls before go-live", bold=True, fill=LAV)
merged(CV, r, 2, 4, f'=IFERROR(INDEX(Lists!$AP$2:$AP$4,MATCH(B{RT},Lists!$AO$2:$AO$4,0)),"")', align=WC); CV.row_dimensions[r].height = 66; r += 1
band(CV, r, "4. MEASURING SUCCESS", 4); r += 1
cell(CV, r, 1, "KPIs", bold=True, fill=LAV); merged(CV, r, 2, 4, f'=IFERROR(IF({lid}="Custom",INDEX({fam("H")},MATCH(INDEX({DR("R")},{m}),{LB},0)),INDEX({LIBR("J")},MATCH({lid},{LIBR("A")},0))),"")', align=WC)
CV.row_dimensions[r].height = 30; r += 1
cell(CV, r, 1, "Indicative hours released / yr", bold=True, fill=LAV); cell(CV, r, 2, gx("AE"), align=CC, fmt="#,##0")
cell(CV, r, 3, "Target / success criterion", bold=True, fill=LAV); cell(CV, r, 4, "", fill=INPUT); r += 1
band(CV, r, "5. NEXT STEP & OWNERSHIP", 4); r += 1
cell(CV, r, 1, "Recommended next step", bold=True, fill=LAV)
merged(CV, r, 2, 4, f'=IFERROR(INDEX(Lists!$AT$2:$AT$6,MATCH(B{VR},{VER},0)),"")', align=WC); CV.row_dimensions[r].height = 34; r += 1
for a, b in [("Business owner", "Executive sponsor"), ("Budget / effort estimate", "Decision date")]:
    cell(CV, r, 1, a, bold=True, fill=LAV); cell(CV, r, 2, "", fill=INPUT); cell(CV, r, 3, b, bold=True, fill=LAV); cell(CV, r, 4, "", fill=INPUT); r += 1
cell(CV, r, 1, "Decision", bold=True, fill=LAV); merged(CV, r, 2, 4, "", fill=INPUT); dv_list(CV, '"Proceed to pilot,Proceed to automation,Fix foundations first,Defer,Reject"', f"B{r}"); r += 1
CV.print_area = f"A1:D{r - 1}"; portrait(CV, fit_h=1)

# ================================================================== Reference
RF = wb.create_sheet("Reference")
title(RF, "Reference - solution families, archetypes, verdicts, functions and standards", "The logic behind the toolkit, for facilitators and for the client's appendix.", 6)
for c, w in zip("ABCDEF", [40, 30, 8, 9, 14, 60]): RF.column_dimensions[c].width = w
r = 4; band(RF, r, "SOLUTION FAMILIES (task nature -> solution)", 6); r += 1
for j, h in enumerate(["Nature of the task", "Solution family", "Level", "AI?", "Min. data", "Typical solutions / design notes"], 1):
    cell(RF, r, j, h, bold=True, fill=GT, color="FFFFFF", align=CC)
r += 1
for f in FAMILIES:
    nat, fam_, lvl, ai, dreq, pot, tools, kpi, note = f
    cell(RF, r, 1, nat, align=WT); cell(RF, r, 2, fam_, align=WT, bold=True); cell(RF, r, 3, f"{lvl}", align=CC); cell(RF, r, 4, ai, align=CC)
    cell(RF, r, 5, DATA_AVAIL[dreq][0], align=WT, size=8); cell(RF, r, 6, f"{tools}.  {note}", align=WT, size=8.5)
    RF.row_dimensions[r].height = 52; r += 1
r += 1; band(RF, r, "SOLUTION LEVELS", 6); r += 1
for lvl, name in LEVELS:
    cell(RF, r, 1, f"Level {lvl}", bold=True, fill=LAV); merged(RF, r, 2, 6, name); r += 1
r += 1; band(RF, r, "AI POSITION ARCHETYPES", 6); r += 1
for code, name, ceil, head, desc in ARCHETYPES:
    cell(RF, r, 1, f"{code}. {name}", bold=True, fill=LAV, align=WT); cell(RF, r, 2, f"Ceiling: level {ceil}", align=CC)
    merged(RF, r, 3, 6, f"{head} {desc}", align=WT, size=8.5); RF.row_dimensions[r].height = 58; r += 1
r += 1; band(RF, r, "VERDICTS", 6); r += 1
for v, mean, nxt in VERDICTS:
    cell(RF, r, 1, v, bold=True, fill=VFILL[v]); merged(RF, r, 2, 6, f"{mean}  Next step: {nxt}", align=WT, size=8.5); RF.row_dimensions[r].height = 40; r += 1
r += 1; band(RF, r, "MINIMUM CONTROLS BY RISK TIER", 6); r += 1
for t, c in CONTROLS:
    cell(RF, r, 1, t, bold=True, fill=RFILL[t]); merged(RF, r, 2, 6, c, align=WT, size=8.5); RF.row_dimensions[r].height = 44; r += 1
r += 1; band(RF, r, "FUNCTION TAXONOMY (APQC Process Classification Framework, cross-industry)", 6); r += 1
for k in FUNCTION_ORDER:
    name, apqc = LFUNC[k]
    cell(RF, r, 1, name, bold=True, fill=LAV); merged(RF, r, 2, 5, apqc); cell(RF, r, 6, f'=COUNTIF({LIBR("C")},A{r})&" library entries"', align=WC); r += 1
r += 1; band(RF, r, "STANDARDS AND FRAMEWORKS BEHIND EACH STEP", 6); r += 1
for j, h in enumerate(["Standard / framework", "", "Step", "", "", "How it is used"], 1):
    if h: cell(RF, r, j, h, bold=True, fill=GT, color="FFFFFF", align=CC)
RF.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2); RF.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5); r += 1
for step, std, use in STANDARDS:
    merged(RF, r, 1, 2, std, bold=True, align=WT, size=8.5); merged(RF, r, 3, 5, step, align=CC, size=8.5); cell(RF, r, 6, use, align=WT, size=8.5)
    RF.row_dimensions[r].height = 40; r += 1
RF.print_area = f"A1:F{r - 1}"; portrait(RF)

# ------------------------------------------------------------------ finish
order = ["Guide", "1. AI Position", "2. Use-Case Library", "3. AI Fit Diagnostic", "4. Opportunity Report", "5. Roadmap", "6. Use-Case Canvas", "7. Sector Profiles", "Reference", "Lists"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    if ws.title != "Lists":
        ws.oddFooter.left.text = "GT AI Opportunity Discovery Toolkit"; ws.oddFooter.left.size = 8
        ws.oddFooter.right.text = "Page &P of &N"; ws.oddFooter.right.size = 8
wb.active = 4 if EXAMPLE else 1
printfix.harden(wb, keep_a3=("3. AI Fit Diagnostic", "2. Use-Case Library"))
wb.save(OUT)
print("saved", OUT)
