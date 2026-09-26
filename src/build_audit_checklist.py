# -*- coding: utf-8 -*-
"""GT AI Security & Agent Audit Checklist v2.0 (Excel, auto-scored).
Consolidates: AI Security Audit Checklist v1.0 (85 items, retained numbering), ISACA Cybersecurity
Recommendations for Securing AI Agents (2026) and ISO/IEC 42001 AIMS checks; every item mapped to GT AI RCM v2.1."""
import sys
import openpyxl, os
sys.path.insert(0, os.path.dirname(__file__))
import printfix
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
V1 = "Audit Checklist v1.0"; IS = "ISACA 2026"; ISO = "ISO/IEC 42001"
SECTIONS = [
    ("3", "Infrastructure Security", [
        ("3.1.1", "AI systems are segmented from the general network", V1, "SE-9, SE-5"),
        ("3.1.2", "Firewall rules restrict AI system access", V1, "SE-5"),
        ("3.1.3", "API endpoints are behind a WAF / API gateway", V1, "SE-10"),
        ("3.1.4", "Internal AI services are not exposed to the internet", V1, "SE-5, SE-9"),
        ("3.1.5", "Network traffic is monitored and logged", V1, "SE-5, OM-2"),
        ("3.1.6", "DDoS protection is in place for AI endpoints", V1, "SE-5, OM-4"),
        ("3.2.1", "GPU/TPU instances are hardened", V1, "SE-1, SE-9"),
        ("3.2.2", "Container images are scanned for vulnerabilities", V1, "SE-3, SE-7"),
        ("3.2.3", "Model-serving infrastructure is patched", V1, "SE-3"),
        ("3.2.4", "Kubernetes / orchestration security is configured", V1, "SE-9"),
        ("3.2.5", "Secrets management in place (no hard-coded credentials)", V1, "SE-8"),
        ("3.2.6", "Resource limits prevent abuse (CPU, memory, GPU)", V1, "OM-4, SE-9"),
        ("3.3.1", "Cloud IAM policies follow least privilege", V1, "SE-2"),
        ("3.3.2", "Storage buckets holding training data are private", V1, "SE-4"),
        ("3.3.3", "Cloud audit logging is enabled", V1, "OM-2"),
        ("3.3.4", "Region / data-residency requirements are met (e.g., Bahrain PDPL, CBB)", V1, "PR-6"),
        ("3.3.5", "Cloud security posture management is active", V1, "SE-1")]),
    ("4", "Data Security", [
        ("4.1.1", "Training data encrypted at rest (AES-256)", V1, "SE-4"),
        ("4.1.2", "Model weights / artefacts encrypted at rest", V1, "SE-4, LC-13"),
        ("4.1.3", "Inference logs encrypted at rest", V1, "SE-4, OM-2"),
        ("4.1.4", "Encryption keys properly managed (KMS / HSM)", V1, "SE-4, SE-8"),
        ("4.1.5", "Backup data encrypted", V1, "SE-4"),
        ("4.2.1", "TLS 1.2+ for all AI API communications", V1, "SE-5"),
        ("4.2.2", "Certificate management automated", V1, "SE-5"),
        ("4.2.3", "Internal service-to-service encryption (mTLS)", V1, "SE-5"),
        ("4.2.4", "Data transfers to / from vendors encrypted", V1, "SE-4, TP-3"),
        ("4.3.1", "Data classification applied to AI datasets", V1, "LC-1, SE-4"),
        ("4.3.2", "Data retention policies enforced (incl. prompts, outputs, AI logs)", V1, "PR-7"),
        ("4.3.3", "Secure data deletion procedures followed", V1, "PR-2, LC-10"),
        ("4.3.4", "Data lineage tracked for training datasets", V1, "LC-1"),
        ("4.3.5", "PII minimisation applied before training", V1, "PR-4"),
        ("4.3.6", "Data anonymisation / pseudonymisation verified", V1, "PR-4")]),
    ("5", "Model Security", [
        ("5.1.1", "Model artefacts stored in a secure registry", V1, "LC-13"),
        ("5.1.2", "Model versioning with integrity checksums", V1, "LC-13"),
        ("5.1.3", "Model signing / verification implemented", V1, "LC-13, SE-7"),
        ("5.1.4", "Unauthorised model modifications are detectable", V1, "LC-13, LC-5"),
        ("5.1.5", "Model rollback capability tested", V1, "LC-13, OM-4"),
        ("5.2.1", "Prompt-injection testing conducted", V1, "SE-6, AA-3"),
        ("5.2.2", "Jailbreak resistance verified", V1, "SE-6"),
        ("5.2.3", "Data-poisoning defences in place", V1, "SE-6, LC-1"),
        ("5.2.4", "Model-extraction attack resistance tested", V1, "SE-6"),
        ("5.2.5", "Adversarial input detection implemented", V1, "SE-6"),
        ("5.2.6", "Evasion attack testing conducted", V1, "SE-6, RS-3"),
        ("5.2.7", "Output-manipulation resistance verified", V1, "SE-6"),
        ("5.3.1", "Input validation and sanitisation", V1, "SE-6"),
        ("5.3.2", "Input size limits enforced", V1, "SE-6, OM-4"),
        ("5.3.3", "Output filtering for sensitive data", V1, "SE-6, PR-4"),
        ("5.3.4", "Output content-safety filters active", V1, "SE-6, RS-2"),
        ("5.3.5", "Rate limiting per user / session", V1, "OM-4"),
        ("5.3.6", "Token / cost limits per request", V1, "OM-4")]),
    ("6", "Access Control", [
        ("6.1", "RBAC implemented for AI system access", V1, "SE-2"),
        ("6.2", "MFA required for administrative access", V1, "SE-2"),
        ("6.3", "API-key rotation policy enforced", V1, "SE-8"),
        ("6.4", "Service-account permissions minimised", V1, "SE-2"),
        ("6.5", "Access reviews conducted quarterly", V1, "SE-2"),
        ("6.6", "Privileged access to training pipelines restricted", V1, "SE-2"),
        ("6.7", "Model deployment requires an approval workflow", V1, "LC-5, LC-12"),
        ("6.8", "SSO integration for AI platforms", V1, "SE-2")]),
    ("7", "Logging, Monitoring & Incident Response", [
        ("7.1.1", "All AI API calls logged", V1, "OM-2"),
        ("7.1.2", "Authentication events logged", V1, "OM-2, SE-2"),
        ("7.1.3", "Model deployments / changes logged", V1, "OM-2, LC-5"),
        ("7.1.4", "Data-access events logged", V1, "OM-2"),
        ("7.1.5", "Log integrity protected (tamper-evident)", V1, "OM-2"),
        ("7.1.6", "Log retention meets regulatory requirements", V1, "OM-2, RO-3"),
        ("7.2.1", "Anomalous usage patterns detected", V1, "OM-1, AT-2"),
        ("7.2.2", "Model performance-degradation alerts", V1, "OM-1"),
        ("7.2.3", "Data-drift detection in place", V1, "OM-1"),
        ("7.2.4", "Cost-anomaly detection active", V1, "OM-4, GL-6"),
        ("7.2.5", "Security alert escalation procedures defined", V1, "IM-1"),
        ("7.3.1", "AI-specific incident response plan exists", V1, "IM-1"),
        ("7.3.2", "Incident response tested (tabletop exercises)", V1, "IM-1, IM-3"),
        ("7.3.3", "Model kill switch / emergency shutdown tested", V1, "OM-4"),
        ("7.3.4", "Communication plan for AI security incidents", V1, "IM-2")]),
    ("8", "Supply Chain Security", [
        ("8.1", "Third-party model provenance verified", V1, "SE-7, LC-6"),
        ("8.2", "Open-source model vulnerability scanning", V1, "SE-7"),
        ("8.3", "ML library dependencies tracked (SBOM / AI-BOM)", V1, "SE-7"),
        ("8.4", "Pre-trained model integrity verified", V1, "SE-7, LC-13"),
        ("8.5", "Vendor security assessments current", V1, "TP-2, TP-3"),
        ("8.6", "Dataset provenance and licensing verified", V1, "LC-11")]),
    ("9", "Compliance & Governance", [
        ("9.1", "AI inventory / registry maintained", V1, "GL-8"),
        ("9.2", "Risk classification assigned (EU AI Act where applicable; GT risk tier)", V1, "RM-5"),
        ("9.3", "Data protection impact assessment (DPIA) completed", V1, "PR-1"),
        ("9.4", "AI ethics review conducted", V1, "GL-7"),
        ("9.5", "Documentation meets regulatory requirements", V1, "LC-4, RO-3"),
        ("9.6", "Audit trail for AI decisions maintained", V1, "OM-2")]),
    ("10", "AI Agent Security (ISACA 2026)", [
        ("10.1", "Inventory covers agents, tools, plugins, models, memory / vector stores, data sources and providers", IS, "GL-8"),
        ("10.2", "Trust boundaries defined (user, runtime, model provider, orchestration, tool layer, approvals) with named control owners", IS, "LC-8"),
        ("10.3", "Each agent has its own workload identity; no shared accounts or long-lived tokens", IS, "LC-8, SE-2"),
        ("10.4", "Short-lived, auto-rotated credentials or federated workload identity used", IS, "SE-8"),
        ("10.5", "Least privilege on every tool, API, knowledge source and data store; authorisation enforced on memory retrieval", IS, "SE-2, SE-8"),
        ("10.6", "Tool execution, code interpretation, browser automation and file parsing run in sandboxes", IS, "SE-9"),
        ("10.7", "Outbound access denied by default; egress allow-listed; cloud metadata and admin interfaces blocked", IS, "SE-9"),
        ("10.8", "All retrieved content and tool output treated as untrusted; instruction hierarchy and content boundaries enforced", IS, "SE-6"),
        ("10.9", "Retrieved content cannot trigger actions without a separate policy decision; source trust levels tracked", IS, "SE-10"),
        ("10.10", "Memory isolated by tenant / user / session with TTL retention and integrity validation", IS, "SE-8"),
        ("10.11", "No secrets in prompts; secrets manager and scoped tokens used", IS, "SE-8"),
        ("10.12", "Tools behind an API gateway / action broker with schema, parameter constraints, quotas; dangerous primitives and SSRF blocked", IS, "SE-10"),
        ("10.13", "Deterministic policy enforcement point validates action, target, identity, business rules and approvals", IS, "SE-10"),
        ("10.14", "Human approval for destructive, financial, legal or irreversible actions, with clear execution summaries", IS, "LC-8, RS-1"),
        ("10.15", "Read-only / recommendation-only mode available for higher-risk use cases", IS, "OM-4"),
        ("10.16", "Prompts, tool calls, decisions and approvals logged centrally, redacted and tamper-resistant", IS, "OM-2"),
        ("10.17", "Agent-specific incident playbooks (prompt injection, tool compromise, memory poisoning, cross-tenant exposure)", IS, "IM-1"),
        ("10.18", "Change and version control with rollback for prompts, policies, tool permissions and memory configuration", IS, "LC-12"),
        ("10.19", "LLM regression tests for known attack patterns run before each release", IS, "LC-12"),
        ("10.20", "Continuous red teaming for prompt injection and tool misuse", IS, "AT-4, AA-3"),
        ("10.21", "Rate limits, token budgets, loop limits and circuit breakers enforced", IS, "OM-4"),
        ("10.22", "Global and per-capability kill switches tested; critical workflows can revert to manual", IS, "OM-4"),
        ("10.23", "Behavioural, performance and governance drift of agents monitored", IS, "OM-1")]),
    ("11", "AI Management System (ISO/IEC 42001)", [
        ("11.1", "AIMS scope documented and approved by top management (cl. 4.3)", ISO, "GL-1"),
        ("11.2", "Signed AI policy with purpose, reach and limits, accessible to staff (cl. 5.2)", ISO, "GL-1, GL-5"),
        ("11.3", "Roles assigned incl. owners for AI operations, AI risk and ethics (cl. 5.3)", ISO, "GL-2"),
        ("11.4", "AI risk criteria and acceptance thresholds approved (cl. 6.1.2)", ISO, "RM-6"),
        ("11.5", "Risk assessment covers bias, explainability, robustness and security (cl. 6.1.2)", ISO, "RM-2"),
        ("11.6", "Risk treatment reconciled with Annex A; Statement of Applicability maintained (cl. 6.1.3)", ISO, "GL-9"),
        ("11.7", "Impacts on individuals, groups and society assessed before deployment (cl. 6.1.4)", ISO, "RM-2"),
        ("11.8", "Assessments repeated at planned intervals and on material change; triggers documented (cl. 8.2)", ISO, "RM-5"),
        ("11.9", "Data provenance, collection and preparation criteria documented per dataset (A.7)", ISO, "LC-1, LC-11"),
        ("11.10", "Models tested for memorisation / regurgitation of personal data (practice)", ISO, "PR-7"),
        ("11.11", "Data-subject rights honoured against training data and models (Bahrain PDPL)", ISO, "PR-5"),
        ("11.12", "AI procurement uses fixed assessment questions; terms of use analysed before use (A.10)", ISO, "TP-3"),
        ("11.13", "Sub-contractors of AI suppliers known and assessed (A.10)", ISO, "TP-2"),
        ("11.14", "Affected persons have an accessible route to complain about AI-assisted decisions (A.8)", ISO, "CO-3"),
        ("11.15", "Confidential speak-up channel for AI concerns feeds management review (practice)", ISO, "CO-4"),
        ("11.16", "Internal audit programme for the AIMS; findings tracked to closure (cl. 9.2)", ISO, "AA-1"),
        ("11.17", "Management review of the AIMS performed (cl. 9.3)", ISO, "GL-1")]),
]

GT = "4F2D7F"; GT2 = "7B5BA6"; LAV = "EDE7F6"; INPUT = "FFF9E6"
thin = Side(style="thin", color="C9C9C9"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
WC = Alignment(wrap_text=True, vertical="center"); CC = Alignment(wrap_text=True, vertical="center", horizontal="center")


def cell(ws, r, c, v, bold=False, fill=None, align=WC, size=9.5, color="000000", fmt=None):
    x = ws.cell(r, c, v); x.font = Font(size=size, bold=bold, color=color); x.alignment = align; x.border = B
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    if fmt: x.number_format = fmt
    return x


def banner(ws, r, text, span, fill=GT, size=14, h=30):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
    c = ws.cell(r, 1, text); c.font = Font(size=size, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=fill); c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = h


wb = openpyxl.Workbook()
# ------------------------------------------------ Audit info & summary
sm = wb.active; sm.title = "Summary"
sm.sheet_view.showGridLines = False
for c, w in zip("ABCDEFGH", [34, 12, 10, 10, 10, 12, 12, 16]): sm.column_dimensions[c].width = w
banner(sm, 1, "GT AI SECURITY & AGENT AUDIT CHECKLIST  v2.0", 8, size=16, h=36)
sm.merge_cells("A2:H2")
sm["A2"] = ("Consolidates the AI Security Audit Checklist v1.0 (85 items, numbering retained), ISACA 'Cybersecurity Recommendations for Securing AI Agents' (2026) "
            "and ISO/IEC 42001 AIMS checks. Every item maps to GT AI RCM v2.1 controls. One pass per AI system.")
sm["A2"].font = Font(italic=True, size=9.5, color=GT); sm["A2"].alignment = WC; sm.row_dimensions[2].height = 42
info = ["AI system name", "System owner", "Audit lead", "Audit date", "Audit type (Initial / Annual / Triggered / Pre-deployment)", "Agentic system? (Yes / No)", "Previous audit date", "Previous audit score"]
for i, k in enumerate(info):
    cell(sm, 4 + i, 1, k, bold=True, fill=LAV); sm.merge_cells(start_row=4 + i, start_column=2, end_row=4 + i, end_column=8)
    cell(sm, 4 + i, 2, None, fill=INPUT)
    for c in range(3, 9): sm.cell(4 + i, c).border = B
dv_type = DataValidation(type="list", formula1='"Initial,Annual,Triggered,Pre-deployment"'); sm.add_data_validation(dv_type); dv_type.add("B8")
dv_ag = DataValidation(type="list", formula1='"Yes,No"'); sm.add_data_validation(dv_ag); dv_ag.add("B9")
r0 = 14
for i, h in enumerate(["Section", "Items", "Pass", "Fail", "N/A", "Not assessed", "Score", "Rating"], 1):
    cell(sm, r0, i, h, bold=True, fill=GT, color="FFFFFF", align=CC)
sm.row_dimensions[r0].height = 28

# ------------------------------------------------ Checklist sheet
ck = wb.create_sheet("Checklist")
ck.sheet_view.showGridLines = False
cols = [("#", 7), ("Control / check", 60), ("Source", 16), ("GT AI RCM v2.1", 14), ("Status", 12), ("Risk if failed", 11), ("Evidence reviewed", 34), ("Notes / finding", 34)]
banner(ck, 1, "AUDIT CHECKLIST  -  select Status (Pass / Fail / N/A) and Risk for every item", len(cols), size=13)
for i, (h, w) in enumerate(cols, 1):
    cell(ck, 3, i, h, bold=True, fill=GT, color="FFFFFF", align=CC); ck.column_dimensions[get_column_letter(i)].width = w
ck.row_dimensions[3].height = 26
r = 4
sec_ranges = []
for num, name, items in SECTIONS:
    ck.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
    c = ck.cell(r, 1, f"{num}. {name}"); c.font = Font(bold=True, color="FFFFFF", size=10.5); c.fill = PatternFill("solid", fgColor=GT2); c.alignment = Alignment(vertical="center", indent=1)
    ck.row_dimensions[r].height = 20
    r += 1; start = r
    for iid, text, src, rcm in items:
        cell(ck, r, 1, iid, bold=True, align=CC); cell(ck, r, 2, text); cell(ck, r, 3, src, size=8.5, align=CC); cell(ck, r, 4, rcm, size=8.5, align=CC)
        for c in (5, 6, 7, 8): cell(ck, r, c, None, fill=INPUT)
        ck.row_dimensions[r].height = 30 if len(text) > 70 else 20
        r += 1
    sec_ranges.append((num, name, start, r - 1))
last = r - 1
dvs = DataValidation(type="list", formula1='"Pass,Fail,N/A"', showErrorMessage=True); ck.add_data_validation(dvs); dvs.add(f"E5:E{last}")
dvr = DataValidation(type="list", formula1='"Critical,High,Medium,Low"'); ck.add_data_validation(dvr); dvr.add(f"F5:F{last}")
for lab, col in [("Pass", "C6E0B4"), ("Fail", "F4B183"), ("N/A", "E7E6E6")]:
    ck.conditional_formatting.add(f"E5:E{last}", CellIsRule(operator="equal", formula=[f'"{lab}"'], fill=PatternFill("solid", fgColor=col)))
for lab, col in [("Critical", "E06666"), ("High", "F4B183"), ("Medium", "FFE699"), ("Low", "C6E0B4")]:
    ck.conditional_formatting.add(f"F5:F{last}", CellIsRule(operator="equal", formula=[f'"{lab}"'], fill=PatternFill("solid", fgColor=col)))
ck.freeze_panes = "C4"; ck.auto_filter.ref = f"A3:H{last}"
CK = "Checklist"

# ------------------------------------------------ summary formulas
for i, (num, name, a, b) in enumerate(sec_ranges):
    rr = r0 + 1 + i
    cell(sm, rr, 1, f"{num}. {name}", bold=True)
    cell(sm, rr, 2, f"=COUNTA({CK}!A{a}:A{b})", align=CC)
    cell(sm, rr, 3, f'=COUNTIF({CK}!E{a}:E{b},"Pass")', align=CC)
    cell(sm, rr, 4, f'=COUNTIF({CK}!E{a}:E{b},"Fail")', align=CC)
    cell(sm, rr, 5, f'=COUNTIF({CK}!E{a}:E{b},"N/A")', align=CC)
    cell(sm, rr, 6, f"=B{rr}-C{rr}-D{rr}-E{rr}", align=CC)
    cell(sm, rr, 7, f'=IF(C{rr}+D{rr}=0,"-",C{rr}/(C{rr}+D{rr}))', align=CC, fmt="0%", bold=True)
    cell(sm, rr, 8, f'=IF(G{rr}="-","-",IF(G{rr}>=0.9,"Low Risk",IF(G{rr}>=0.75,"Moderate Risk",IF(G{rr}>=0.6,"High Risk","Critical Risk"))))', align=CC, bold=True)
tot = r0 + 1 + len(sec_ranges)
cell(sm, tot, 1, "OVERALL", bold=True, fill=LAV)
for c, f in [(2, f"=SUM(B{r0+1}:B{tot-1})"), (3, f"=SUM(C{r0+1}:C{tot-1})"), (4, f"=SUM(D{r0+1}:D{tot-1})"), (5, f"=SUM(E{r0+1}:E{tot-1})"), (6, f"=SUM(F{r0+1}:F{tot-1})")]:
    cell(sm, tot, c, f, bold=True, fill=LAV, align=CC)
cell(sm, tot, 7, f'=IF(C{tot}+D{tot}=0,"-",C{tot}/(C{tot}+D{tot}))', bold=True, fill=LAV, align=CC, fmt="0%")
cell(sm, tot, 8, f'=IF(G{tot}="-","-",IF(G{tot}>=0.9,"Low Risk",IF(G{tot}>=0.75,"Moderate Risk",IF(G{tot}>=0.6,"High Risk","Critical Risk"))))', bold=True, fill=LAV, align=CC)
cell(sm, tot + 1, 1, "Critical-risk items failed", bold=True)
cell(sm, tot + 1, 2, f'=COUNTIFS({CK}!E5:E{last},"Fail",{CK}!F5:F{last},"Critical")', bold=True, align=CC)
sm.merge_cells(start_row=tot + 1, start_column=3, end_row=tot + 1, end_column=8)
cell(sm, tot + 1, 3, f'=IF(B{tot+1}>0,"Any failed Critical item escalates the overall rating to Critical Risk regardless of score.","")', size=8.5, color="C00000")
cell(sm, tot + 2, 1, "FINAL RATING", bold=True, fill=GT, color="FFFFFF")
sm.merge_cells(start_row=tot + 2, start_column=2, end_row=tot + 2, end_column=8)
cell(sm, tot + 2, 2, f'=IF(B{tot+1}>0,"Critical Risk",H{tot})', bold=True, size=12)
for lab, col in [("Low Risk", "C6E0B4"), ("Moderate Risk", "FFE699"), ("High Risk", "F4B183"), ("Critical Risk", "E06666")]:
    sm.conditional_formatting.add(f"H{r0+1}:H{tot}", CellIsRule(operator="equal", formula=[f'"{lab}"'], fill=PatternFill("solid", fgColor=col)))
    sm.conditional_formatting.add(f"B{tot+2}", CellIsRule(operator="equal", formula=[f'"{lab}"'], fill=PatternFill("solid", fgColor=col)))
_ag = r0 + 1 + [x[0] for x in sec_ranges].index("10")
sm.conditional_formatting.add(f"A{_ag}:H{_ag}", FormulaRule(formula=['UPPER($B$9)="NO"'], font=Font(color="A6A6A6")))
rb = tot + 4
cell(sm, rb, 1, "Score", bold=True, fill=GT, color="FFFFFF"); cell(sm, rb, 2, "Rating", bold=True, fill=GT, color="FFFFFF")
sm.merge_cells(start_row=rb, start_column=3, end_row=rb, end_column=8); cell(sm, rb, 3, "Required action", bold=True, fill=GT, color="FFFFFF")
for i, (sc, rt, act, col) in enumerate([("90-100%", "Low Risk", "Annual re-audit", "C6E0B4"), ("75-89%", "Moderate Risk", "Remediation within 90 days", "FFE699"),
                                         ("60-74%", "High Risk", "Remediation within 30 days", "F4B183"), ("Below 60%", "Critical Risk", "Immediate remediation; consider suspending the AI system", "E06666")]):
    cell(sm, rb + 1 + i, 1, sc, align=CC); cell(sm, rb + 1 + i, 2, rt, fill=col, bold=True, align=CC)
    sm.merge_cells(start_row=rb + 1 + i, start_column=3, end_row=rb + 1 + i, end_column=8); cell(sm, rb + 1 + i, 3, act)
note = rb + 6
sm.merge_cells(start_row=note, start_column=1, end_row=note, end_column=8)
sm.cell(note, 1, "Score = Pass / (Pass + Fail); N/A excluded. Section 10 applies to agentic AI systems (mark N/A otherwise). Section 11 applies where the organisation operates or plans an ISO/IEC 42001 AI management system. Rating bands per AI Security Audit Checklist v1.0.").font = Font(italic=True, size=8.5, color="595959")
sm.cell(note, 1).alignment = WC; sm.row_dimensions[note].height = 32

# ------------------------------------------------ Findings (auto list of failed items)
fd = wb.create_sheet("Findings & Remediation")
fd.sheet_view.showGridLines = False
fcols = [("#", 5), ("Item", 8), ("Finding (failed check)", 52), ("Severity", 10), ("Section", 22), ("GT AI RCM v2.1", 14), ("Remediation", 40), ("Owner", 16), ("Due date", 12), ("Status", 12)]
banner(fd, 1, "FINDINGS & REMEDIATION PLAN  -  failed items are listed automatically; complete remediation, owner, due date, status", len(fcols), size=12)
for i, (h, w) in enumerate(fcols, 1):
    cell(fd, 3, i, h, bold=True, fill=GT, color="FFFFFF", align=CC); fd.column_dimensions[get_column_letter(i)].width = w
# helper col K on checklist: running count of failed items
for rr in range(5, last + 1):
    ck.cell(rr, 11, f'=IF(E{rr}="Fail",COUNTIF($E$5:E{rr},"Fail"),"")')
    # section name helper
for num, name, a, b in sec_ranges:
    for rr in range(a, b + 1): ck.cell(rr, 12, f"{num}. {name}")
ck.column_dimensions["K"].hidden = True; ck.column_dimensions["L"].hidden = True
N = 125
for i in range(1, N + 1):
    rr = 3 + i
    m = f'MATCH({i},{CK}!$K$5:$K${last},0)'
    cell(fd, rr, 1, i, align=CC)
    cell(fd, rr, 2, f'=IFERROR(INDEX({CK}!$A$5:$A${last},{m}),"")', align=CC, bold=True)
    cell(fd, rr, 3, f'=IFERROR(INDEX({CK}!$B$5:$B${last},{m}),"")')
    cell(fd, rr, 4, f'=IFERROR(INDEX({CK}!$F$5:$F${last},{m})&"","")', align=CC, bold=True)
    cell(fd, rr, 5, f'=IFERROR(INDEX({CK}!$L$5:$L${last},{m}),"")', size=8.5)
    cell(fd, rr, 6, f'=IFERROR(INDEX({CK}!$D$5:$D${last},{m}),"")', size=8.5, align=CC)
    for c in (7, 8, 9, 10): cell(fd, rr, c, None, fill=INPUT)
    fd.row_dimensions[rr].height = 28
for lab, col in [("Critical", "E06666"), ("High", "F4B183"), ("Medium", "FFE699"), ("Low", "C6E0B4")]:
    fd.conditional_formatting.add(f"D4:D{3+N}", CellIsRule(operator="equal", formula=[f'"{lab}"'], fill=PatternFill("solid", fgColor=col)))
dvf = DataValidation(type="list", formula1='"Open,In progress,Closed,Risk accepted"'); fd.add_data_validation(dvf); dvf.add(f"J4:J{3+N}")
fd.freeze_panes = "C4"

# ------------------------------------------------ Sign-off
so = wb.create_sheet("Sign-Off")
so.sheet_view.showGridLines = False
for c, w in zip("ABCD", [28, 30, 30, 16]): so.column_dimensions[c].width = w
banner(so, 1, "SIGN-OFF", 4, size=14)
for i, h in enumerate(["Role", "Name", "Signature", "Date"], 1): cell(so, 3, i, h, bold=True, fill=GT, color="FFFFFF", align=CC)
for i, role in enumerate(["Audit Lead", "System Owner", "CISO", "AI Governance Lead"]):
    cell(so, 4 + i, 1, role, bold=True)
    for c in (2, 3, 4): cell(so, 4 + i, c, None, fill=INPUT)
    so.row_dimensions[4 + i].height = 30

# ------------------------------------------------ print setup (A4, plain)
for ws in wb.worksheets:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if ws.title in ("Checklist", "Findings & Remediation") else "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.page_margins.header = ws.page_margins.footer = 0.3
    ws.sheet_properties.tabColor = GT
ck.print_area = f"A1:H{last}"; ck.print_title_rows = "3:3"
fd.print_area = f"A1:J{3+N}"; fd.print_title_rows = "3:3"
sm.print_area = f"A1:H{note}"; so.print_area = "A1:D7"
printfix.harden(wb)
wb.save(OUT)
print("saved", OUT, "items", sum(len(s[2]) for s in SECTIONS))
