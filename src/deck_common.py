# -*- coding: utf-8 -*-
"""Shared slide helpers for GT proposal decks built on the GT AI Service Lines template."""
import sys, os, copy
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData, BubbleChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
sys.path.insert(0, os.path.dirname(__file__))

prs = S_COVER = S_AGENDA = S_DIVIDER = S_CONTENT = S_THANKS = None
NEW = []


def init(template):
    global prs, S_COVER, S_AGENDA, S_DIVIDER, S_CONTENT, S_THANKS, NEW, SW, SH
    prs = Presentation(template)
    orig = list(prs.slides)
    S_COVER, S_AGENDA, S_DIVIDER, S_CONTENT, S_THANKS = orig[0], orig[1], orig[5], orig[8], orig[18]
    SW, SH = prs.slide_width, prs.slide_height
    NEW = []
    return prs


GT = RGBColor(0x4F, 0x2D, 0x7F); GT2 = RGBColor(0x7B, 0x5B, 0xA6); LAV = RGBColor(0xED, 0xE7, 0xF6); LAV2 = RGBColor(0xD9, 0xCC, 0xEB)
TEAL = RGBColor(0x00, 0xA7, 0xB5); CORAL = RGBColor(0xE8, 0x70, 0x5F); DARK = RGBColor(0x2E, 0x1A, 0x47); GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(255, 255, 255); BLACK = RGBColor(0x22, 0x22, 0x22); AMBER = RGBColor(0xF2, 0xB1, 0x3C); GREEN = RGBColor(0x5B, 0xA8, 0x5B)
FONT = "Arial"


# ------------------------------------------------------------------ slide duplication helpers
def clone_shapes(src, dst, names=None):
    rid_map = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        rid_map[rId] = dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external else dst.part.rels.get_or_add(rel.reltype, rel.target_part)
    for el in src.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[1]
        if tag in ("nvGrpSpPr", "grpSpPr", "extLst"):
            continue
        if names is not None:
            nm = el.find(".//" + qn("p:cNvPr"))
            if nm is None or nm.get("name") not in names:
                continue
        new = copy.deepcopy(el)
        for node in new.iter():
            for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
                v = node.get(attr)
                if v in rid_map: node.set(attr, rid_map[v])
        dst.shapes._spTree.append(new)
    bg = src._element.find(qn("p:cSld")).find(qn("p:bg"))
    if bg is not None and names is None:
        dst._element.find(qn("p:cSld")).insert(0, copy.deepcopy(bg))


def duplicate(src):
    s = prs.slides.add_slide(src.slide_layout)
    for shp in list(s.shapes): shp._element.getparent().remove(shp._element)
    clone_shapes(src, s)
    return s


def set_text(shape, text, size=None, bold=None, color=None):
    tf = shape.text_frame
    runs = [r for p in tf.paragraphs for r in p.runs]
    if not runs: return
    f0 = runs[0].font
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    for r in list(p.runs)[1:]: r._r.getparent().remove(r._r)
    lines = text.split("\n")
    p.runs[0].text = lines[0]
    for ln in lines[1:]:
        np_ = copy.deepcopy(p._p); tf._txBody.append(np_)
        tf.paragraphs[-1].runs[0].text = ln
    for p in tf.paragraphs:
        for r in p.runs:
            if size: r.font.size = Pt(size)
            if bold is not None: r.font.bold = bold
            if color: r.font.color.rgb = color


def find(slide, contains):
    for sh in slide.shapes:
        if sh.has_text_frame and contains in sh.text_frame.text: return sh
        if sh.shape_type == 6:
            for g in sh.shapes:
                if g.has_text_frame and contains in g.text_frame.text: return g
    return None


# ------------------------------------------------------------------ drawing helpers
def txt(slide, x, y, w, h, paras, size=11, color=BLACK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullet=False, spacing=2):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(45720); tf.margin_top = tf.margin_bottom = Emu(22860)
    if isinstance(paras, str): paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.space_after = Pt(spacing)
        segs = p if isinstance(p, list) else [(p, bold)]
        for seg, b in segs:
            r = para.add_run(); r.text = (("•  " if bullet and seg is segs[0][0] else "") + seg)
            r.font.name = FONT; r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = color
    return tb


def box(slide, x, y, w, h, fill=LAV, line=None, text=None, size=11, color=BLACK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is None: s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    if text is not None:
        tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Emu(64008); tf.margin_top = tf.margin_bottom = Emu(36576)
        paras = text if isinstance(text, list) else [text]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            segs = p if isinstance(p, list) else [(p, bold)]
            for seg, b in segs:
                r = para.add_run(); r.text = seg; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = color
    return s


def table(slide, x, y, w, rows, colw, size=9, header_fill=GT, rowh=None, first_bold=True, fills=None):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, Emu(x), Emu(y), Emu(w), Emu(rowh * nr if rowh else 300000 * nr))
    tbl = shp.table
    tot = sum(colw)
    for i, cw in enumerate(colw): tbl.columns[i].width = Emu(int(w * cw / tot))
    for ri, row in enumerate(rows):
        if rowh: tbl.rows[ri].height = Emu(rowh)
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci); c.text = ""
            c.margin_left = c.margin_right = Emu(54864); c.margin_top = c.margin_bottom = Emu(27432)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(val)
            r.font.name = FONT; r.font.size = Pt(size)
            if ri == 0:
                r.font.bold = True; r.font.color.rgb = WHITE; c.fill.solid(); c.fill.fore_color.rgb = header_fill
            else:
                r.font.color.rgb = BLACK; r.font.bold = first_bold and ci == 0
                c.fill.solid(); c.fill.fore_color.rgb = LAV if ri % 2 == 0 else WHITE
                if fills and (ri, ci) in fills:
                    c.fill.fore_color.rgb = fills[(ri, ci)]
    return tbl


L = 524396; R = 11787482; CW = R - L; TOP = 1480000


def content(title, subtitle=None):
    s = prs.slides.add_slide(S_CONTENT.slide_layout)
    for shp in list(s.shapes): shp._element.getparent().remove(shp._element)
    clone_shapes(S_CONTENT, s, names={"Copyright", "Straight Connector 4", "Slide Number Placeholder 1", "Picture 6"})
    cp = find(s, "Grant Thornton Bahrain")
    if cp: set_text(cp, "© 2026 Grant Thornton Bahrain. All rights reserved")
    txt(s, L, 400000, CW, 560000, [[(title, True)]], size=26, color=GT)
    if subtitle:
        txt(s, L - 20000, 930000, CW, 480000, subtitle, size=11.5, color=GREY)
    return s


def divider(title, number):
    s = duplicate(S_DIVIDER)
    set_text(find(s, "AI Service Lines"), title)
    t = find(s, "Grant Thornton Bahrain")
    if t: set_text(t, "© 2026 Grant Thornton Bahrain. All rights reserved")
    return s




def finish(out):
    # ------------------------------------------------------------------ keep only new slides, in order
    sldIdLst = prs.slides._sldIdLst
    keep = {id(sl._element) for sl in NEW}
    idmap = {}
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        part = prs.part.related_part(rId)
        idmap[id(part._element)] = sldId
    for sldId in list(sldIdLst):
        sldIdLst.remove(sldId)
    for sl in NEW:
        sldIdLst.append(idmap[id(sl._element)])
    for rel in list(prs.part.rels.values()):
        if rel.reltype.endswith("/slide") and id(rel.target_part._element) not in keep:
            prs.part.drop_rel(rel.rId)
    def fix_year(shapes):
        for sh in shapes:
            if sh.shape_type == 6: fix_year(sh.shapes); continue
            if sh.has_text_frame:
                for p_ in sh.text_frame.paragraphs:
                    for r_ in p_.runs:
                        if "2025" in r_.text and "Grant Thornton" in sh.text_frame.text: r_.text = r_.text.replace("2025", "2026")
    for sl in NEW: fix_year(sl.shapes)
    for master in prs.slide_masters:
        for el in [master._element] + [lay._element for lay in master.slide_layouts]:
            for para in el.iter(qn("a:p")):
                ts = list(para.iter(qn("a:t")))
                if "Grant Thornton Bahrain" in "".join(t.text or "" for t in ts):
                    for t in ts:
                        if t.text:
                            for yr in ("2024", "2025"): t.text = t.text.replace(yr, "2026")
    prs.save(out)
    print("saved", out, len(NEW), "slides")
