# -*- coding: utf-8 -*-
"""GT INTERNAL operating guide: running the AI Readiness & Risk Assessment without sharing the engine.
Usage: build_ops_guide.py OUT.docx"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
GT = RGBColor(0x4F, 0x2D, 0x7F)
doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexfill); tcPr.append(sh)


def h(text, lvl=1):
    p = doc.add_heading(text, level=lvl)
    for r in p.runs: r.font.color.rgb = GT; r.font.name = "Arial"
    return p


def para(text, bold_lead=None, size=10):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True; r.font.size = Pt(size)
    r = p.add_run(text); r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(4)
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(2)


def tbl(rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.autofit = False
    for j, w in enumerate(widths): t.columns[j].width = Cm(w)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""; r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(9)
            if i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "4F2D7F")
            c.width = Cm(widths[j])
    doc.add_paragraph()
    return t


title = doc.add_paragraph(); r = title.add_run("GT AI Readiness & Risk Assessment"); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = GT
sub = doc.add_paragraph(); r = sub.add_run("Operating guide - collecting answers and issuing reports without sharing the engine"); r.font.size = Pt(12); r.font.color.rgb = GT
warn = doc.add_table(rows=1, cols=1); warn.style = "Table Grid"; c = warn.cell(0, 0); shade(c, "FCE4D6")
c.text = ""; rr = c.paragraphs[0].add_run("GT INTERNAL - DO NOT SEND TO CLIENTS. "); rr.bold = True
c.paragraphs[0].add_run("The Excel engine (GT_AI_Readiness_Assessment_Tool.xlsx) contains the scoring logic, weightings and content mapping. It never leaves GT. "
                        "Clients receive only the questionnaire link and the PDF report.")
doc.add_paragraph()

h("1. How the process works")
para("Essential - report within 5 working days (Day 1 kick-off; questionnaire by Day 2; GT analysis Days 3-4; report Day 5). "
     "Standard - 2 weeks (multi-respondent questionnaire, evidence review and validation interviews in week 1-2; report and leadership debrief end of week 2). "
     "Plus - 6 weeks, extendable to 8 (readiness report week 2; AI inventory, risk & impact assessments, control review, ISO/IEC 42001 gap and AI security review; "
     "final report and board presentation week 6). If responses arrive late, the report is issued within 3 working days of complete responses.", "Service commitment by option: ")
tbl([["Step", "Who", "What", "Day"],
     ["1. Set up the form", "GT engagement team", "Create the client's Microsoft Form from the Word questionnaire (Quick import)", "Day 1"],
     ["2. Client completes", "Client sponsor / respondents", "34 questions: profile (14), AI use profile (4), readiness (16)", "Day 1-2 (~45 min)"],
     ["3. Export responses", "GT", "Forms > Responses > Open results in Excel > download", "Day 3"],
     ["4. Generate report", "GT", "Run generate_readiness_reports.py: engine is filled, recalculated and exported", "Day 3"],
     ["5. Quality review", "GT AI GRC reviewer", "Check answers, headline results, narrative and priorities", "Day 3-4"],
     ["6. Send", "GT engagement lead", "Open the email draft (.eml) in Outlook, review and send", "Day 5"],
     ["7. Debrief", "GT + client", "Walk through results; agree priorities (Standard / Plus)", "Day 5"]], [3.2, 3.4, 8, 1.8])

h("2. Files")
tbl([["File", "Location (repo)", "Share with client?"],
     ["GT_AI_Readiness_Questionnaire.docx", "1.3_.../client_questionnaire/", "Only via Microsoft Forms (or as a fallback document)"],
     ["GT_AI_Readiness_Responses_Template.xlsx", "1.3_.../client_questionnaire/", "Fallback only - if Forms cannot be used"],
     ["GT_AI_Readiness_Assessment_Tool.xlsx (engine)", "1.3_AI_Readiness_Assessment_Tool/", "NO - GT internal"],
     ["generate_readiness_reports.py", "src/", "NO - GT internal"],
     ["<Client>_AI_Readiness_Report.pdf", "generated", "YES - the deliverable"],
     ["<Client>_AI_Readiness_Engine.xlsx", "generated", "NO - keep in the engagement file"]], [6.5, 5.5, 4.5])

h("3. Step-by-step")
h("3.1 Create the Microsoft Form", 2)
bullets(["Go to forms.office.com with your GT account > New Form > Quick import (Word / PDF) > upload GT_AI_Readiness_Questionnaire.docx.",
         "Check the imported form: every question must be a single-choice question with the options A-E (or the listed options), except the free-text profile fields (organisation, names, email).",
         "Do not edit the question tags in square brackets - [P1] ... [P14], [A1] ... [A4], [Q5] ... [Q20]. The report generator uses them to read the answers.",
         "Settings: Who can fill out this form = Anyone can respond; turn off 'Record name'; set a closing date; add the GT confidentiality note from the questionnaire to the form description.",
         "Use one form per client (Duplicate the master form and rename it '<Client> - AI Readiness Assessment'). This keeps responses separated.",
         "Send the link to the client sponsor with the agreed deadline. For a team response, ask the sponsor to submit one consolidated response."], "List Number")
para("If Quick import is not available on your tenant, build the form once manually from the Word document (copy titles exactly, including the tags), then duplicate it for each client.", "Note: ")
h("3.2 Export the responses", 2)
bullets(["Forms > Responses > Open results in Excel. Save the file to the client's engagement folder (not to personal OneDrive).",
         "Fallback: if the client completed the Word document or the Responses Template instead, transcribe into GT_AI_Readiness_Responses_Template.xlsx (one row per organisation). Letters A-E or the full option text are both accepted."], "List Number")
h("3.3 Generate the report", 2)
para("One-off set-up on the GT laptop: install Python 3.9+, LibreOffice (libreoffice.org) and the Python packages:", None)
p = doc.add_paragraph(); r = p.add_run("pip install openpyxl pypdf"); r.font.name = "Consolas"
para("Then run (paths in quotes if they contain spaces):", None)
p = doc.add_paragraph(); r = p.add_run('python src/generate_readiness_reports.py "Responses.xlsx" "01_AI_Governance_Risk_Compliance/1.3_AI_Readiness_Assessment_Tool/GT_AI_Readiness_Assessment_Tool.xlsx" "Output" --gt-contact "Name, Grant Thornton Bahrain"')
r.font.name = "Consolas"; r.font.size = Pt(8.5)
para("For each response row the script creates a folder Output/<Organisation>/ containing the PDF report, the filled engine (internal) and <Organisation>_email_draft.eml. "
     "The console shows the headline result, e.g. 'Grant Thornton Bahrain: readiness 60% (Informed), tier Medium, exposure Moderate, 7 report pages'. "
     "A WARNING line lists any unanswered questions - the report is then marked provisional; follow up with the client before issuing.", None)
h("3.4 Quality review (mandatory before sending)", 2)
bullets(["Organisation name, sector, scope and date on page 1 are correct.",
         "Headline results are plausible against what we know of the client (e.g., a regulated bank with no AI policy should not be 'Repeatable').",
         "Inherent risk tier: check overrides (AI influencing decisions about individuals cannot be Low).",
         "AI adoption maturity curve: the adoption stage follows the client's AI-use answer (A1) and the governance stage follows overall readiness - check both are plausible and that any 'adoption ahead of governance' message is reflected in the debrief.",
         "Read the narrative, the top-5 priorities and the roadmap; amend in the engine copy and re-export if a professional judgement differs - record the change in the engagement file.",
         "Regulatory section: confirm applicability of EU AI Act and CBB items to this client.",
         "Second-person review for Plus engagements or where exposure is High / Critical."])
h("3.5 Send the report", 2)
bullets(["Double-click <Organisation>_email_draft.eml. Classic Outlook for Windows opens it as an unsent draft with the PDF attached and the client's email (from the form) in To.",
         "Review the text, add CCs, check the attachment, and send from your GT mailbox.",
         "New Outlook / Outlook for Mac may open the .eml as a received message: use Forward, or create a new email and attach the PDF; the text can be copied from the .eml.",
         "Never attach the engine workbook."], "List Number")

h("4. Data handling")
bullets(["Responses and outputs are stored only in the client engagement folder in GT's Microsoft 365 environment.",
         "Do not paste client answers or reports into public AI tools.",
         "Personal data is limited to respondent contact details (Bahrain PDPL); retain per the engagement letter and delete on request.",
         "Reports are marked 'Confidential' and addressed only to the named sponsor."])

h("5. Troubleshooting")
tbl([["Message / issue", "Cause", "Fix"],
     ["answer not recognised: '...'", "Option text was edited in the form or typed by hand", "Use the exact option text or its letter A-E in the responses file"],
     ["LibreOffice (soffice) not found", "LibreOffice not installed / not on the default path", "Install LibreOffice; on Windows the default install path is detected automatically"],
     ["A question is reported as unanswered", "Column header lost its [tag] in Forms", "Restore the tag in the form title, or rename the column header in the export to include it"],
     ["Report has more / fewer pages than usual", "Long free-text answers or many findings", "Normal - pagination is automatic. Check no section is cut off"],
     ["Client asks for the Excel tool", "-", "Explain the methodology is proprietary; offer the debrief and, under the Plus option, the detailed evidence review"]], [5, 5, 6.5])

doc.save(OUT)
print("saved", OUT)
