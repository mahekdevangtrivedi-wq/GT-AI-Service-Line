# -*- coding: utf-8 -*-
"""GT AI Strategy & Roadmap - proposal deck (approach & methodology), built on the GT AI Service Lines template.

Usage: build_strategy_proposal.py TEMPLATE.pptx OUT.pptx IMGDIR
IMGDIR holds rendered pages of the Opportunity Discovery Toolkit retail-bank worked example
(tk_position.png, tk_report.png, tk_canvas.png, tk_library_crop.png).
"""
import sys, os, collections
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.dml.color import RGBColor
sys.path.insert(0, os.path.dirname(__file__))
import deck_common as dc
from deck_common import txt, box, table, content, divider, duplicate, find, set_text, GT, GT2, LAV, LAV2, TEAL, CORAL, DARK, GREY, WHITE, BLACK, AMBER, GREEN, L, R, CW, TOP
from strategy_data import READY7, READY7_OVERALL, READY7_TARGET_2027, APPS
from opportunity_data import ARCHETYPES, FAMILIES, VERDICTS, LEVELS, STANDARDS
from opportunity_library import LIBRARY, SECTOR_PROFILES, FUNCTIONS as LFUNC
from tool_data import DOMAIN_MAP, PILLARS as DOMAINS

TEMPLATE, OUT, IMG = sys.argv[1], sys.argv[2], sys.argv[3]
dc.init(TEMPLATE)
add = dc.NEW.append
VCOL = {"AI fit now": RGBColor(0xA9, 0xD0, 0x8E), "Automate - AI not needed": RGBColor(0x9B, 0xC2, 0xE6), "AI later - close data gap first": RGBColor(0xFF, 0xE6, 0x99),
        "Stretch - partner or later stage": RGBColor(0xF4, 0xB1, 0x83), "Redesign / standardise first": RGBColor(0xD9, 0xD9, 0xD9)}
RCOL = {"Low": RGBColor(0xC6, 0xE0, 0xB4), "Medium": RGBColor(0xFF, 0xE6, 0x99), "High": RGBColor(0xF4, 0xB1, 0x83)}
NL = len(LIBRARY); NSEC = sum(1 for x in LIBRARY if x[1] != "Cross-industry"); NXI = NL - NSEC
SECCOUNT = collections.Counter(x[1] for x in LIBRARY)


def pic(s, name, x, y, w=None, h=None):
    p = s.shapes.add_picture(os.path.join(IMG, name), Emu(x), Emu(y), width=Emu(w) if w else None, height=Emu(h) if h else None)
    p.line.color.rgb = LAV2; p.line.width = Pt(0.75)
    return p


def chevrons(s, y, items, h=620000, size=12):
    n = len(items); gap = -60000; w = (CW - gap * (n - 1)) // n
    for i, t in enumerate(items):
        box(s, L + i * (w + gap), y, w, h, fill=[GT, GT2, TEAL, DARK, CORAL, GT][i], text=t, size=size, bold=True, color=WHITE,
            shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
    return w


def agenda(items):
    ag = duplicate(dc.S_AGENDA)
    for old, new in zip(["Market Opportunities", "AI Service Lines for", "System Setup"], items):
        set_text(find(ag, old), new)
    for sh in ag.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() in items:
            sh.left = Emu(6795700); sh.width = Emu(4485523); sh.text_frame.word_wrap = True
            for p_ in sh.text_frame.paragraphs: p_.alignment = PP_ALIGN.LEFT
            for r_ in [r for p_ in sh.text_frame.paragraphs for r in p_.runs]: r_.font.size = Pt(18)
        if sh.has_text_frame and "Grant Thornton Bahrain" in sh.text_frame.text and "©" in sh.text_frame.text:
            set_text(sh, "© 2026 Grant Thornton Bahrain. All rights reserved")
    return ag


# GT Ready7 (Sept 2026 results report) consolidated into the five GT readiness domains
cur, tgt = {}, {}
for (p, v), t in zip(READY7, READY7_TARGET_2027):
    d = DOMAIN_MAP[p]; cur.setdefault(d, []).append(v); tgt.setdefault(d, []).append(t)
CUR = [int(sum(cur[d]) / len(cur[d]) + 0.5) for d in DOMAINS]; TGT = [int(sum(tgt[d]) / len(tgt[d]) + 0.5) for d in DOMAINS]
SRC = {d: " + ".join(p for p, _ in READY7 if DOMAIN_MAP[p] == d) for d in DOMAINS}

# =========================================================== 1 cover / 2 agenda
cov = duplicate(dc.S_COVER)
set_text(find(cov, "Artificial Intelligence"), "AI Strategy &\nRoadmap")
set_text(find(cov, "We go beyond"), "Proposal  |  October 2026")
add(cov)
add(agenda(["Context & Our Starting Point", "Approach & Methodology", "Engagement Plan & Deliverables"]))

# =========================================================== 3 questions
s = content("The Questions Every Leadership Team Is Asking", "An AI strategy is only useful if it says precisely where AI creates value in this organisation - and where it does not.")
qs = [("Where exactly does AI fit?", "Which processes, functions and decisions would genuinely benefit - beyond generic 'use ChatGPT' advice?"),
      ("Do we have the data?", "Is our data digital, joined-up, reliable and deep enough for AI - or do we need foundations first?"),
      ("Is automation enough?", "For many organisations - especially smaller ones - workflow automation, RPA and dashboards deliver most of the value at lower cost and risk."),
      ("Build, buy or partner?", "What can we realistically run with our scale, budget and technology team?"),
      ("What are the risks?", "Which uses affect customers, employees or regulators (PDPL, CBB, NHRA, TRA, EU AI Act) and what controls are required?"),
      ("What do we do first?", "A prioritised roadmap with quick wins, measurable KPIs and clear ownership.")]
cw = (CW - 2 * 150000) // 3; chh = 1900000
for i, (q, a) in enumerate(qs):
    x = L + (i % 3) * (cw + 150000); y = TOP + (i // 3) * (chh + 150000)
    box(s, x, y, cw, chh, fill=LAV); box(s, x, y, cw, 90000, fill=[GT, GT2, TEAL, CORAL, DARK, GREEN][i])
    txt(s, x + 100000, y + 200000, cw - 200000, 500000, [[(q, True)]], size=16, color=GT)
    txt(s, x + 100000, y + 750000, cw - 200000, 1100000, a, size=12.5)
box(s, L, TOP + 2 * chh + 400000, CW, 520000, fill=DARK, text=[[("Our approach answers these with evidence: ", True),
    (f"an AI position screen, a {NL}-entry AI use-case library, a process-by-process AI fit diagnostic and a risk-tiered roadmap - delivered within 5 working days.", False)]], size=12, color=WHITE)
add(s)

# =========================================================== 4 sector focus
s = content("Where the Value Is - Sector Focus", "AI and automation potential is highest where volumes, data and regulation meet: financial services, healthcare and telecom lead - and our library goes deepest there.")
rows = [["Sector", "Value at stake by 2030", "Potential", "Where the value is", "Library entries"]]
fills = {}
for i, (sec, gdp, pot, pools, regs) in enumerate(SECTOR_PROFILES, 1):
    rows.append([sec, gdp.split(" in the 2030")[0] if gdp.startswith("Not sized") else gdp.replace("  |  ", " | "), pot, pools, str(SECCOUNT.get(sec, 0))])
    fills[(i, 2)] = RGBColor(0xA9, 0xD0, 0x8E) if pot == "Very high" else RGBColor(0xC6, 0xE0, 0xB4) if pot == "High" else RGBColor(0xFF, 0xE6, 0x99)
table(s, L, TOP - 250000, CW, rows, [2.4, 2.6, 1.0, 6.2, 1.0], size=8, rowh=300000, fills=fills)
txt(s, L, TOP + 4280000, CW, 400000, f"Plus {NXI} cross-industry entries (finance, HR, customer service, procurement, risk, IT ...). "
    "Source (value at stake, GDP uplift by 2030): QuantumBlack & PwC analysis (Predictions for 2030); insurance and capital markets share the financial-services figure.", size=8.5, color=GREY)
add(s)

# =========================================================== 5 our starting point - Ready7
s = content("Our Starting Point - GT Bahrain AI Ready7 Position", "GT Bahrain's AI Ready7 results (Sept 2026), consolidated into the five GT readiness domains, against our proposed end-2027 targets.")
cd = CategoryChartData(); cd.categories = DOMAINS
cd.add_series("AI Ready7 result (Sept 2026)", CUR); cd.add_series("Target (end-2027, proposed)", TGT)
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(L), Emu(TOP - 50000), Emu(int(CW * 0.58)), Emu(4300000), cd)
ch = gf.chart; ch.has_title = False; ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(10)
pl = ch.plots[0]; pl.gap_width = 70; pl.overlap = -10; pl.has_data_labels = True
pl.data_labels.number_format = '0"%"'; pl.data_labels.number_format_is_linked = False; pl.data_labels.font.size = Pt(10); pl.data_labels.font.bold = True
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
for ser, col in zip(pl.series, [GT, LAV2]):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
ch.value_axis.maximum_scale = 100; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 25; ch.value_axis.tick_labels.font.size = Pt(9)
ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
ch.category_axis.tick_labels.font.size = Pt(10)
x2 = L + int(CW * 0.60); w2 = R - x2
rows = [["Domain", "Ready7", "Target", "Gap"]] + [[d, f"{c}%", f"{t}%", f"{t - c:+d}"] for d, c, t in zip(DOMAINS, CUR, TGT)] + [["Overall readiness", f"{READY7_OVERALL:.1f}%", ">= 75%", ""]]
fills = {(i, 1): RGBColor(0xF8, 0xCB, 0xAD) if c < 60 else RGBColor(0xFF, 0xE6, 0x99) for i, c in enumerate(CUR, 1)}
table(s, x2, TOP - 50000, w2, rows, [3.3, 0.9, 1, 0.8], size=10, rowh=360000, fills=fills)
lo = min(zip(CUR, DOMAINS)); hi = max(zip(CUR, DOMAINS))
txt(s, x2, TOP + 2550000, w2, 1700000, [[("What it means: ", True), (f"GT Bahrain is 'Informed' ({READY7_OVERALL:.1f}%) - strongest in {hi[1]} ({hi[0]}%), weakest in {lo[1]} ({lo[0]}%). "
     "Knowing where AI fits and governing it is our own priority too.", False)],
     [("Readiness vs fit: ", True), ("Ready7 tells us HOW READY an organisation is; the Opportunity Discovery method tells us WHERE AI FITS. A complete strategy needs both.", False)]], size=10.5, spacing=8)
txt(s, L, TOP + 4300000, CW, 300000, f"Source: GT Bahrain AI Ready7 Result Report (Sept 2026). Domain = average of its Ready7 pillars: Data & Technology = {SRC['Data & Technology']}; "
    f"AI Risk & Security = {SRC['AI Risk & Security']}. Maturity: <50% Partial | 50-70% Informed | 70-90% Repeatable | 90%+ Adaptive.", size=8, color=GREY)
add(s)

# =========================================================== 6 divider / 7 methodology
add(divider("Approach &\nMethodology", 2))
s = content("Our Methodology at a Glance", "Five phases, report within 5 working days. Phase 3 - AI Opportunity Discovery - is what makes the strategy specific to the organisation.")
w = chevrons(s, TOP - 50000, ["1  Mobilise", "2  Situational analysis", "3  AI opportunity discovery", "4  Strategy formulation", "5  Roadmap & report"], h=650000, size=12.5)
detail = [(["Scope, sponsor, stakeholders", "Pre-work: sector, process list and documents", "Kick-off"], "Day 0-1"),
          (["External scan (regulation, market, sector, technology)", "Internal scan: strategy, systems, data, current AI use", "Leadership interviews", "SWOT -> strategic options"], "Day 1"),
          (["AI position screen (data x scale)", "Function workshops with the use-case library", "AI fit diagnostic per process", "Value, feasibility & risk scoring"], "Day 1-3"),
          (["AI vision & ambition", "Objectives & KPIs", "Priority portfolio", "Build / buy / partner", "Risk appetite & guardrails", "Operating model"], "Day 3-4"),
          (["H1 / H2 / H3 roadmap", "Enablers: data, platforms, people, governance", "Use-case canvases", "Report & presentation"], "Day 4-5")]
for i, (items, when) in enumerate(detail):
    x = L + i * (w - 60000) + 30000; ww = w - 120000
    box(s, x, TOP + 700000, ww, 3150000, fill=LAV)
    txt(s, x + 30000, TOP + 780000, ww - 60000, 3050000, items, size=10.5, bullet=True, spacing=7)
    box(s, x, TOP + 3900000, ww, 450000, fill=WHITE, line=GT, text=when, size=11, bold=True, color=GT)
add(s)

# =========================================================== 8 principles
s = content("Principles Behind Our Approach", "What makes a GT AI strategy different from a technology wish-list.")
pr = [("Problem first, not technology first", "We start from processes, pain points and decisions - then choose the technology."),
      ("Automation before AI - where it is enough", "If rules, workflow or a dashboard solve it, we say so. AI is recommended only where it adds value that automation cannot."),
      ("Data decides", "Every AI recommendation passes a data gate: is the data digital, available, reliable and (for prediction) labelled?"),
      ("Right-sized to the organisation", "A solution ceiling based on scale, budget and technology capacity: an SME and a bank get different answers."),
      ("Proportionate risk", "Each opportunity is risk-tiered (NSW AIAF, EU AI Act Annex III, PDPL, sector regulators) and carries minimum controls before go-live."),
      ("Standards-based and vendor-agnostic", "ISO/IEC 42001, 22989, 5259, 42005, NIST AI RMF, OECD, APQC - and no vendor commissions.")]
pw = (CW - 150000) // 2; phh = 1050000
for i, (h, d) in enumerate(pr):
    x = L + (i % 2) * (pw + 150000); y = TOP - 50000 + (i // 2) * (phh + 120000)
    box(s, x, y, 650000, phh, fill=[GT, GT2, TEAL, CORAL, DARK, GREEN][i], text=str(i + 1), size=24, bold=True, color=WHITE)
    box(s, x + 650000, y, pw - 650000, phh, fill=LAV)
    txt(s, x + 750000, y + 90000, pw - 820000, 380000, [[(h, True)]], size=13.5, color=GT)
    txt(s, x + 750000, y + 470000, pw - 820000, 560000, d, size=11)
add(s)

# =========================================================== 9 phases 1-2
s = content("Phases 1-2: Mobilise and Situational Analysis", "Understand the organisation, its sector and its starting point before discussing solutions.")
cols = [("ACTIVITIES", GT, ["Kick-off with sponsor; confirm scope, objectives and success criteria", "Review strategy, IT landscape, data assets and current AI use (incl. shadow AI)",
                            "Leadership interviews (strategy, operations, IT, risk)", "External scan: Bahrain Vision 2030, iGA AI Policy, PDPL, sector regulator (CBB, NHRA, TRA ...), EU AI Act reach, peer AI use",
                            "SWOT -> strategic options"]),
        ("PRE-WORK & TOOLS", GT2, ["Client pre-work before Day 1: sector, organisation chart, top processes and pain points, key documents",
                                   "Sector profile and value-at-stake data", "PESTLE and SWOT / TOWS templates", "AI use inventory template (incl. embedded AI in SaaS)",
                                   "Optional baseline: GT AI Readiness & Risk Assessment (20 questions, 5 domains)"]),
        ("OUTPUTS", TEAL, ["Situational analysis and SWOT", "Strategic options", "Current AI use inventory", "Readiness baseline (if selected)"])]
cw = (CW - 2 * 150000) // 3
for i, (h, col, items) in enumerate(cols):
    x = L + i * (cw + 150000)
    box(s, x, TOP - 50000, cw, 420000, fill=col, text=h, size=13, bold=True, color=WHITE)
    box(s, x, TOP + 370000, cw, 3900000, fill=LAV)
    txt(s, x + 70000, TOP + 470000, cw - 140000, 3800000, items, size=12, bullet=True, spacing=9)
add(s)

# =========================================================== 10 toolkit overview
s = content("Phase 3: The GT AI Opportunity Discovery Toolkit", "A structured, facilitated method that answers 'where exactly does AI fit?' - process by process - and tells the organisation where automation is enough.")
steps = [("1", "Position", "Is the organisation positioned for AI? 12 questions on DATA (digitisation, history, quality, documents, integration, outcomes) and SCALE & CAPACITY (size, volumes, IT, budget, digital channels, change capacity).",
          "Archetype + solution ceiling"),
         ("2", "Discover", f"Function-by-function workshops using the GT AI Use-Case Library: {NL} entries - {NXI} cross-industry plus {NSEC} for 14 sectors, filtered to the client's sector.",
          "Process & pain-point inventory"),
         ("3", "Diagnose", "For every process: task nature -> solution family (10 families, 5 levels), data gate, organisational fit and risk tier - plus the library's GT recommendation.",
          "Verdict + recommendation"),
         ("4", "Prioritise", "Value (volume, effort, importance, customer impact) x feasibility (data, standardisation, simplicity, fit), risk-adjusted.",
          "Ranked opportunities"),
         ("5", "Plan", "Opportunity report, H1 / H2 / H3 roadmap lanes with the enabling foundations, and a one-page canvas per priority use case.",
          "Report, roadmap, canvases")]
sw_ = (CW - 4 * 100000) // 5
for i, (n, h, d, o) in enumerate(steps):
    x = L + i * (sw_ + 100000)
    box(s, x, TOP - 50000, sw_, 700000, fill=[GT, GT2, TEAL, DARK, CORAL][i], text=[[(n + "  " + h, True)]], size=16, color=WHITE)
    box(s, x, TOP + 650000, sw_, 2900000, fill=LAV, text=d, size=11, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    box(s, x, TOP + 3600000, sw_, 520000, fill=WHITE, line=GT, text=o, size=10, bold=True, color=GT)
box(s, L, TOP + 4250000, CW, 520000, fill=DARK, text=[[("Not a readiness test: ", True), ("the toolkit locates opportunities and the right type of solution. Governance maturity is measured separately by the GT AI Readiness & Risk Assessment.", False)]],
    size=11, color=WHITE)
add(s)

# =========================================================== 11 step 1 position
s = content("Step 1: Organisational AI Position - Is Automation Enough?", "Two scores place the organisation in one of four archetypes. Each archetype sets a ceiling on the complexity of solution it can realistically adopt today.")
mx = L + 1200000; my = TOP + 250000; qw = 2350000; qh = 1650000
for (rr, cc), ai in {(0, 0): 2, (0, 1): 3, (1, 0): 0, (1, 1): 1}.items():
    code, name, ceil, head, desc = ARCHETYPES[ai]
    col = [RGBColor(0x9B, 0xC2, 0xE6), RGBColor(0xFF, 0xE6, 0x99), RGBColor(0xC6, 0xE0, 0xB4), RGBColor(0xA9, 0xD0, 0x8E)][ai]
    box(s, mx + cc * (qw + 60000), my + rr * (qh + 60000), qw, qh, fill=col,
        text=[[(f"{code}. {name}", True)], [(head, False)], [(f"Ceiling: level {ceil} - {LEVELS[ceil - 1][1]}", True)]], size=10.5, color=BLACK)
txt(s, L, my + 300000, 1150000, qh, [[("DATA", True)], [("HIGH", True)]], size=11, color=GT, align=PP_ALIGN.CENTER)
txt(s, L, my + qh + 360000, 1150000, qh, [[("DATA", True)], [("LOW", True)]], size=11, color=GT, align=PP_ALIGN.CENTER)
txt(s, mx, my + 2 * qh + 150000, qw, 300000, [[("SCALE & CAPACITY LOW", True)]], size=11, color=GT, align=PP_ALIGN.CENTER)
txt(s, mx + qw + 60000, my + 2 * qh + 150000, qw, 300000, [[("SCALE & CAPACITY HIGH", True)]], size=11, color=GT, align=PP_ALIGN.CENTER)
x2 = mx + 2 * qw + 300000; w2 = R - x2
box(s, x2, TOP - 50000, w2, 380000, fill=GT, text="12 position questions", size=12, bold=True, color=WHITE)
txt(s, x2, TOP + 380000, w2, 2000000, [[("Data position: ", True), ("how processes run (paper -> integrated systems); depth of digital history; data quality & ownership; where documents live; integration; recorded outcomes ('labels').", False)],
                                       [("Scale & capacity: ", True), ("employees; transaction volumes; in-house technology skills; budget; digital interactions; capacity for change.", False)]], size=10.5, spacing=8)
box(s, x2, TOP + 2350000, w2, 380000, fill=TEAL, text="Solution levels", size=12, bold=True, color=WHITE)
txt(s, x2, TOP + 2750000, w2, 1700000, [f"{n}. {t}" for n, t in LEVELS], size=11, spacing=4)
txt(s, L, TOP + 4450000, CW, 300000, "Basis: OECD Framework for the Classification of AI Systems (Economic Context; Data & Input), ISO/IEC 42001 cl. 4.1-4.2, NIST AI RMF MAP 1.3-1.6.", size=8.5, color=GREY)
add(s)

# =========================================================== 12 library
s = content(f"Step 2: The GT AI Use-Case Library - {NL} Entries", "Recommendations come from a curated, standards-based library - not from a blank page. Deepest coverage where automation potential is highest.")
order = [sp[0] for sp in SECTOR_PROFILES]
cd = CategoryChartData(); cd.categories = ["Cross-industry"] + order
cd.add_series("Entries", [SECCOUNT["Cross-industry"]] + [SECCOUNT.get(x, 0) for x in order])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Emu(L), Emu(TOP - 150000), Emu(int(CW * 0.42)), Emu(4650000), cd)
ch = gf.chart; ch.has_legend = False; ch.has_title = False
pl = ch.plots[0]; pl.gap_width = 40; pl.has_data_labels = True; pl.data_labels.font.size = Pt(8.5); pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = GT2
for i, nm in enumerate(["Cross-industry"] + order):
    if nm in ("Banking & payments", "Insurance", "Healthcare", "Telecommunications", "Capital markets & wealth"):
        pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = GT
ch.category_axis.reverse_order = True; ch.category_axis.tick_labels.font.size = Pt(8.5); ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
x2 = L + int(CW * 0.44); w2 = R - x2
box(s, x2, TOP - 150000, w2, 380000, fill=GT, text="Every entry carries", size=12, bold=True, color=WHITE)
anat = [("Sector & function", "APQC Process Classification Framework level-1 functions"), ("Process & typical pain point", ""), ("Task nature -> solution family", "ISO/IEC 22989 / OECD-based, 10 families"),
        ("GT recommendation", "what to implement and how - specific to the use case"), ("Data needed & KPIs", "what the data gate checks; how value is measured"),
        ("Regulatory & risk notes", "CBB, NHRA, TRA, NBR, PDPL, EU AI Act, model-risk, sector standards"), ("Minimum risk tier", "EU AI Act Annex III high-risk uses (credit, life & health insurance, employment, education, public assistance, emergency triage) floored at High")]
y = TOP + 260000
for k, v in anat:
    txt(s, x2, y, w2, 300000, [[(k + (": " if v else ""), True), (v, False)]], size=9.5, spacing=0); y += 290000 if len(v) < 90 else 400000
ex = next(x for x in LIBRARY if x[0] == "BNK-05")
box(s, x2, y + 60000, w2, 300000, fill=GT2, text=f"Example entry  {ex[0]}  |  {ex[1]}  |  {ex[4]}", size=10, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
txt(s, x2, y + 380000, w2, 1500000, [[("Recommendation: ", True), (ex[6], False)], [("Data: ", True), (ex[7], False)], [("KPIs: ", True), (ex[8], False)],
                                     [("Notes: ", True), (ex[9], False)], [("Minimum risk tier: ", True), (ex[10], False)]], size=8.5, spacing=2)
add(s)

# =========================================================== 13 steps 2-3 diagnose
s = content("Step 3: Diagnose - From Task to Solution", "The nature of the work - not the hype - determines the solution. Each process gets one of five verdicts and the library's recommendation.")
rows = [["Nature of the task", "Solution family", "Level", "AI?", "Min. data"]]
for f in FAMILIES:
    rows.append([f[0], f[1], str(f[2]), f[3], ["None", "Partial", "Digital history", "With outcomes"][f[4]]])
fills = {(i, 3): RGBColor(0x9B, 0xC2, 0xE6) if f[3] == "No" else RGBColor(0xC6, 0xE0, 0xB4) if f[3] == "Light" else RGBColor(0xA9, 0xD0, 0x8E) for i, f in enumerate(FAMILIES, 1)}
table(s, L, TOP - 100000, int(CW * 0.62), rows, [4.2, 3.1, 0.6, 0.6, 1.3], size=8.5, rowh=355000, fills=fills)
x2 = L + int(CW * 0.62) + 150000; w2 = R - x2
box(s, x2, TOP - 100000, w2, 360000, fill=GT, text="Verdict logic (in order)", size=12, bold=True, color=WHITE)
flow = [("Process varies by person?", "Redesign / standardise first"), ("Family needs no AI?", "Automate - AI not needed"), ("Data below the family's minimum?", "AI later - close data gap first"),
        ("Level above the organisation's ceiling?", "Stretch - partner or later stage"), ("Otherwise", "AI fit now")]
for i, (q, v) in enumerate(flow):
    y = TOP + 330000 + i * 640000
    box(s, x2, y, int(w2 * 0.48), 560000, fill=LAV, text=q, size=9.5, bold=True, color=GT)
    box(s, x2 + int(w2 * 0.5), y, int(w2 * 0.5), 560000, fill=VCOL[v], text=v, size=9.5, bold=True)
txt(s, x2, TOP + 3600000, w2, 900000, "Inputs per process: volume, effort (hours / month), standardisation, data availability, strategic importance, customer / quality impact, effect on "
    "individuals' rights, personal data, intended autonomy. Basis: ISO/IEC 22989, 23053, OECD AI classification, ISO/IEC 5259, ISO/IEC 42001 Annex A.7.", size=9, color=GREY)
add(s)

# =========================================================== 14 step 4 prioritise
s = content("Step 4: Prioritise - Value, Feasibility and Risk", "Transparent scoring that leadership can challenge - and that separates quick wins from strategic bets.")
cards = [("VALUE (1-5)", GT, ["30%  Volume of the process", "25%  Effort today (staff hours / month)", "25%  Strategic importance", "20%  Customer / quality impact"]),
         ("FEASIBILITY (1-5)", GT2, ["30%  Data fit (vs. family minimum)", "25%  Process standardisation", "25%  Solution simplicity (level 1-5)", "20%  Organisational fit (vs. ceiling)"]),
         ("INHERENT RISK TIER", CORAL, ["+2  Affects individuals' rights or access", "+1  Personal or confidential data", "+0-2  Autonomy (assist / approve / automated)", "+1-2  Impact of errors",
                                         "Never below the library minimum tier"])]
cw = (CW - 2 * 150000) // 3
for i, (h, col, items) in enumerate(cards):
    x = L + i * (cw + 150000)
    box(s, x, TOP - 50000, cw, 420000, fill=col, text=h, size=13, bold=True, color=WHITE)
    box(s, x, TOP + 370000, cw, 1850000, fill=LAV)
    txt(s, x + 80000, TOP + 470000, cw - 160000, 1750000, items, size=12, spacing=9)
qx = L; qy = TOP + 2400000; qw2 = 1800000; qh2 = 700000
for i, (lab, col, d) in enumerate([("Strategic bets", GT2, "High value, harder"), ("Quick wins", GREEN, "High value, feasible"), ("Deprioritise", GREY, "Low value, harder"), ("Fill-ins", TEAL, "Low value, feasible")]):
    box(s, qx + (i % 2) * (qw2 + 50000), qy + (i // 2) * (qh2 + 50000), qw2, qh2, fill=col, text=[[(lab, True)], [(d, False)]], size=11, color=WHITE)
x2 = qx + 2 * qw2 + 250000; w2 = R - x2
txt(s, x2, qy, w2, 1600000, [[("Priority ", True), ("= value x feasibility, reduced for High risk (-15%) and where redesign is needed first (-20%).", False)],
                             [("Horizons: ", True), ("H1 (0-6 months) quick wins and automation;  H2 (6-12 months) strategic bets, stretch items, redesigned processes and every High-risk or agentic use case "
                                                      "(after impact assessment and controls);  H3 (12-24 months) items waiting on data foundations.", False)],
                             [("Minimum controls by tier ", True), ("(GT AI Risk Assessment Framework, ISO/IEC 42005, PDPL) on every canvas.", False)]], size=11, spacing=8)
txt(s, L, TOP + 4450000, CW, 300000, "Basis: NIST AI RMF MAP 3.1-3.2, ISO/IEC 42001 cl. 6.2, NSW AI Assessment Framework, EU AI Act Annex III, ISACA Securing AI Agents (2026).", size=8.5, color=GREY)
add(s)

# =========================================================== 15 illustrative output - retail bank
s = content("Illustrative Output - Bahrain Retail Bank", "Worked example included with the toolkit (illustrative inputs): 25 processes selected from the banking and cross-industry library entries.")
kp = [("Build & scale AI", "AI position (archetype D)\nData 61%  |  Scale & capacity 78%"), ("76%", "of processes: AI fit now\n(IDP, ML alert triage, virtual agents)"),
      ("8%", "automation is enough\n(payment repairs, CBB returns)"), ("16%", "need data or process\nfoundations first")]
kw = (CW - 3 * 120000) // 4
for i, (big, small) in enumerate(kp):
    x = L + i * (kw + 120000)
    box(s, x, TOP - 150000, kw, 1150000, fill=[GT, RGBColor(0x5B, 0xA8, 0x5B), RGBColor(0x2E, 0x75, 0xB6), CORAL][i])
    txt(s, x + 60000, TOP - 90000, kw - 120000, 520000, [[(big, True)]], size=22 if i else 18, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, x + 60000, TOP + 450000, kw - 120000, 600000, small.split("\n"), size=10.5, color=WHITE, align=PP_ALIGN.CENTER, spacing=0)
top = [["#", "Process / use case", "Solution family", "Verdict", "Risk", "Horizon"],
       ["1", "Payment exceptions & repairs", "Rules engine / workflow", "Automate - AI not needed", "Low", "H1"],
       ["2", "Sanctions & PEP screening alert review", "Predictive ML", "AI fit now", "Medium", "H2"],
       ["3", "Retail customer onboarding (eKYC)", "Intelligent document processing", "AI fit now", "High", "H2"],
       ["4", "Banking virtual assistant", "Conversational AI", "AI fit now", "Medium", "H1"],
       ["5", "Retail credit decisioning", "Predictive ML", "AI fit now", "High", "H2"],
       ["6", "Trade-finance document checking (LCs)", "Intelligent document processing", "AI fit now", "Medium", "H1"],
       ["7", "IT service desk", "Conversational AI", "AI fit now", "Low", "H1"],
       ["8", "Card & payment fraud detection", "Predictive ML", "AI fit now", "High", "H2"]]
fills = {}
for i, r_ in enumerate(top[1:], 1): fills[(i, 3)] = VCOL[r_[3]]; fills[(i, 4)] = RCOL[r_[4]]
table(s, L, TOP + 1120000, int(CW * 0.66), top, [0.3, 3.4, 2.4, 2.1, 0.8, 0.7], size=9, rowh=320000, fills=fills)
x2 = L + int(CW * 0.68); w2 = R - x2
txt(s, x2, TOP + 1120000, w2, 3200000, [[("What the client learns", True)],
     "Highest-ranked item needs no AI at all - rules-based payment repair.", "Credit decisioning and eKYC fit, but are High-risk (EU AI Act Annex III / biometrics): scheduled for H2 after impact assessment and controls.",
     "Collections, card disputes and IFRS 9 staging fail the data gate: capture outcomes first.", "Complaints handling varies by team: standardise before automating."], size=10.5, spacing=6, color=BLACK)
add(s)

# =========================================================== 16 what the client receives
s = content("What the Client Receives from Phase 3", "Illustrative pages from the retail-bank worked example.")
ih = 4500000
pic(s, "tk_position.png", L, TOP - 100000, h=ih)
pic(s, "tk_report.png", L + 3350000, TOP - 100000, h=ih)
pic(s, "tk_canvas.png", L + 6700000, TOP - 100000, h=ih)
for x, t in [(L, "Organisational AI position"), (L + 3350000, "AI opportunity report"), (L + 6700000, "Use-case canvas (one per priority)")]:
    txt(s, x, TOP + ih - 50000, 3200000, 300000, [[(t, True)]], size=10.5, color=GT)
x2 = L + 9900000
txt(s, x2, TOP - 100000, R - x2, ih, [[("Also delivered:", True)], "Top-10 opportunities with GT recommendations", "Full AI fit diagnostic (up to 50 processes)", "Roadmap lanes with enabling foundations",
                                       "Sector lens and regulators", "Standards appendix"], size=10.5, spacing=8)
add(s)

# =========================================================== 17 phase 4
s = content("Phase 4: AI Strategy Formulation", "The opportunity evidence becomes a strategy leadership can approve, fund and govern.")
blocks = [("Vision & ambition", "Where AI will - and will not - be used; ambition consistent with the AI position archetype."),
          ("Objectives & KPIs", "3-5 measurable objectives (ISO/IEC 42001 cl. 6.2) using the library KPIs: hours released, cycle time, quality, customer outcomes."),
          ("Priority portfolio", "The ranked opportunities grouped into themes; automation programme and AI programme distinguished."),
          ("Build / buy / partner", "Decided per family against the solution ceiling; vendor-selection criteria (security, data residency, no training on client data)."),
          ("Risk appetite & guardrails", "AI risk appetite statement, acceptable-use policy, approval by risk tier, human oversight."),
          ("Operating model", "Executive sponsor, AI owner / CoE, decision rights, intake and stage-gates; links to existing risk and IT governance."),
          ("People & change", "AI literacy for all users, role-based skills, change and communications plan."),
          ("Data & platforms", "Data foundations required by H2 / H3 opportunities; platform choices (automation, GenAI, data).")]
bw = (CW - 3 * 120000) // 4; bh = 1950000
for i, (h, d) in enumerate(blocks):
    x = L + (i % 4) * (bw + 120000); y = TOP - 50000 + (i // 4) * (bh + 150000)
    box(s, x, y, bw, 450000, fill=[GT, GT2, TEAL, DARK][i % 4], text=h, size=12.5, bold=True, color=WHITE)
    box(s, x, y + 450000, bw, bh - 450000, fill=LAV, text=d, size=11, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
add(s)

# =========================================================== 18 phase 5
s = content("Phase 5: Roadmap and Report", "A sequenced plan the organisation can execute - with the foundations each opportunity depends on.")
lanes = [("H1  Now  (0-6 months)", GT, "Quick wins and fill-ins: automation programme, off-the-shelf GenAI under policy, first configured AI pilots."),
         ("H2  Next  (6-12 months)", GT2, "Strategic bets, stretch items with partners, redesigned processes, High-risk and agentic use cases once controls are in place."),
         ("H3  Later  (12-24 months)", TEAL, "AI use cases unlocked by data foundations; scaling of proven pilots.")]
for i, (h, col, d) in enumerate(lanes):
    y = TOP - 50000 + i * 760000
    box(s, L, y, 2600000, 680000, fill=col, text=h, size=13, bold=True, color=WHITE)
    box(s, L + 2600000, y, CW - 2600000, 680000, fill=LAV, text=d, size=12, align=PP_ALIGN.LEFT)
box(s, L, TOP + 2300000, CW, 380000, fill=DARK, text="ENABLERS ACROSS ALL HORIZONS", size=11.5, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
en = [("Governance", "AI policy, AI register, tier-based risk assessment and controls (GT AI Risk Assessment Framework)"), ("Data", "owners, quality, integration, capture of outcomes"),
      ("Platforms", "one automation platform; enterprise GenAI; data platform as needed"), ("People", "AI literacy, champions, change management"),
      ("Value tracking", "KPI dashboard, benefits realisation, quarterly portfolio review")]
ew = (CW - 4 * 100000) // 5
for i, (h, d) in enumerate(en):
    box(s, L + i * (ew + 100000), TOP + 2760000, ew, 1250000, fill=LAV, text=[[(h, True)], [(d, False)]], size=10.5)
txt(s, L, TOP + 4100000, CW, 600000, [[("Report delivered on Day 5: ", True), ("AI strategy and roadmap report, opportunity report, top-10 use-case canvases and a leadership presentation.", False)]], size=11.5)
add(s)

# =========================================================== 19 standards
s = content("Standards and Frameworks Behind the Method", "Every step is traceable to recognised international guidance - giving boards and regulators confidence in the result.")
std = [["Step", "Standard or framework", "How it is used"]] + [[a, b, c] for a, b, c in STANDARDS]
table(s, L, TOP - 300000, CW, std, [1.0, 4.2, 6.7], size=7, rowh=212000)
add(s)

# =========================================================== 20 divider / 21 plan
add(divider("Engagement Plan &\nDeliverables", 3))
s = content("Engagement Plan - Report Within 5 Working Days", "A focused, workshop-based engagement. Client pre-work before Day 1 keeps the timeline to five days.")
days = [("Day 0", "Pre-work", ["Kick-off call (30 min)", "Client shares sector, top processes, pain points and key documents", "Workshop invitations"]),
        ("Day 1", "Position & context", ["Leadership kick-off", "AI position screen (30-45 min)", "Situational scan & interviews", "First function workshops"]),
        ("Day 2", "Discovery", ["Function workshops with the use-case library (2-3 h each)", "Process inventory completed"]),
        ("Day 3", "Diagnose & prioritise", ["AI fit diagnostic", "Value / feasibility / risk scoring", "Validation call with sponsor"]),
        ("Day 4", "Strategy & roadmap", ["Strategy formulation", "Roadmap lanes & enablers", "Use-case canvases", "GT quality review"]),
        ("Day 5", "Report & presentation", ["Strategy & roadmap report issued", "Leadership presentation", "Agree next steps"])]
dw = (CW - 5 * 80000) // 6
for i, (d, h, items) in enumerate(days):
    x = L + i * (dw + 80000)
    box(s, x, TOP - 50000, dw, 420000, fill=[GREY, GT, GT2, TEAL, DARK, CORAL][i], text=d, size=14, bold=True, color=WHITE)
    box(s, x, TOP + 370000, dw, 380000, fill=LAV2, text=h, size=11, bold=True, color=GT)
    box(s, x, TOP + 750000, dw, 2700000, fill=LAV)
    txt(s, x + 30000, TOP + 830000, dw - 60000, 2600000, items, size=10.5, bullet=True, spacing=7)
txt(s, L, TOP + 3600000, CW, 800000, [[("Client time: ", True), ("sponsor 30 min / day; one 2-3 hour workshop per function in scope; position screen with leadership, IT and finance; Day-5 presentation.  ", False),
                                      ("Scope: ", True), ("standard scope covers up to 6 functions and 50 processes; larger groups run additional workshop days in parallel.", False)]], size=11)
add(s)

# =========================================================== 22 deliverables & options
s = content("Deliverables and Engagement Options")
rows = [["Deliverable", "Format"], ["AI strategy & roadmap report (situational analysis, strategy, roadmap)", "Report (PDF) + leadership deck"],
        ["Organisational AI position and opportunity report", "Toolkit output (PDF)"], ["AI fit diagnostic of in-scope processes with GT recommendations", "Excel (GT AI Opportunity Discovery Toolkit)"],
        ["Use-case canvases for the top 10 opportunities", "One-page canvases (PDF)"], ["Roadmap lanes, enablers and KPIs", "Report section + Excel"]]
table(s, L, TOP - 250000, int(CW * 0.55), rows, [3.6, 2.4], size=10.5, rowh=520000)
x2 = L + int(CW * 0.55) + 200000; w2 = R - x2
opts = [("AUTOMATION & AI SCAN", TEAL, "report in 3 working days - SMEs", "Position screen, one discovery workshop, diagnostic of up to 15 processes, opportunity report and roadmap lanes. Answers: is automation enough?"),
        ("AI STRATEGY & ROADMAP", GT, "report in 5 working days - standard", "All five phases, up to 6 functions / 50 processes, strategy, roadmap and top-10 canvases."),
        ("STRATEGY + READINESS", DARK, "both reports in 5 working days", "Standard scope plus the GT AI Readiness & Risk Assessment (questionnaire completed as pre-work) - where AI fits AND how ready the organisation is.")]
for i, (h, col, sub, d) in enumerate(opts):
    y = TOP - 250000 + i * 1450000
    box(s, x2, y, w2, 420000, fill=col, text=[[(h + "   ", True), (sub, False)]], size=11, color=WHITE)
    box(s, x2, y + 420000, w2, 950000, fill=LAV, text=d, size=10.5, align=PP_ALIGN.LEFT)
txt(s, L, TOP + 3050000, int(CW * 0.55), 500000, "Fees are proposed per option once scope (entities, functions, locations) is confirmed. Implementation support is scoped separately.", size=10, color=GREY)
add(s)

# =========================================================== 23 why GT
s = content("Why Grant Thornton Bahrain", "Strategy, governance, risk, cyber and assurance under one roof - with proprietary tools and the GT global network.")
why = [("Practise what we advise", "GT Bahrain assessed its own AI readiness with AI Ready7 (Sept 2026) and uses the results to set its own priorities."),
       ("Trusted and regulation-ready", "Audit, GRC and cyber heritage; Bahrain PDPL, CBB, iGA AI Policy and EU AI Act mapped into the GT AI Risk Assessment Framework."),
       ("Honest about automation", "We recommend AI only where it beats automation - protecting budgets and reducing risk."),
       ("Sector depth", f"A {NL}-entry use-case library with the deepest coverage in banking, insurance, healthcare and telecom."),
       ("End-to-end AI service lines", "AI Strategy & Roadmap | AI Governance, Risk & Compliance | AI-Driven Automation & Apps | Training & Awareness.")]
half = (CW - 150000) // 2
for i, (h, d) in enumerate(why):
    y = TOP - 50000 + i * 640000
    box(s, L, y, half, 580000, fill=LAV, text=[[(h, True)], [(d, False)]], size=10, align=PP_ALIGN.LEFT)
x2 = L + half + 150000
box(s, x2, TOP - 50000, half, 380000, fill=GT, text="GT AI tools and apps", size=12, bold=True, color=WHITE)
rows = [["Tool / app", "Status"], ["GT AI Opportunity Discovery Toolkit (240-entry library)", "Available"], ["GT AI Readiness & Risk Assessment (20 questions, 5 domains)", "Available"]] + \
       [[a[0], a[3]] for a in APPS if a[0] != "GT AI Readiness & Risk Assessment Tool"]
fills = {(i, 1): RGBColor(0xFF, 0xE6, 0x99) for i, r_ in enumerate(rows) if i and r_[1] in ("In preparation", "Planned")}
table(s, x2, TOP + 330000, half, rows, [4.2, 1.6], size=9.5, rowh=380000, fills=fills)
txt(s, x2, TOP + 3400000, half, 500000, "The AI-powered BCP tool and AI-powered Financial Statements tool are in preparation and are not yet available to clients.", size=9, color=GREY)
add(s)

ty = duplicate(dc.S_THANKS)
t = find(ty, "Grant Thornton Bahrain")
if t: set_text(t, "© 2026 Grant Thornton Bahrain. All rights reserved")
add(ty)
dc.finish(OUT)
